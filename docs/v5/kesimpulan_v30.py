import sys; sys.path.insert(0, '/tmp/pptwork/v5')
from single import make_subset
make_subset('/tmp/pptwork/v5/out7.pptx', [74], '/tmp/pptwork/v5/_base_ks.pptx')
exec(open('/tmp/pptwork/helpers_v4.py').read())
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
p = Presentation('/tmp/pptwork/v5/_base_ks.pptx')
GREEN, PINK, NAVY = 'C9F0D6', 'FDE2DE', '0B5E6E'
US = '/tmp/pptwork/v5/us/'


def find_title(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.top is not None and 0.6*E <= sh.top <= 1.3*E and sh.width > 9*E and sh.text_frame.text.strip():
            return sh


def prep(s, title, resume):
    t = find_title(s)
    for sh in list(s.shapes):
        keep = (sh._element is t._element) or sh.top >= 10.3*E or (sh.left >= 18.0*E and sh.top < 1.0*E) or sh.top < 0.45*E
        if not keep:
            sh._element.getparent().remove(sh._element)
    for sh in list(s.shapes):
        if sh.top > 10.3*E and 16.5*E < sh.left < 17.9*E:
            sh._element.getparent().remove(sh._element)
    for sh in [x for x in s.shapes if x.has_text_frame and x.left > 17.9*E and x.top > 10.4*E and x.text_frame.text.strip().isdigit()]:
        set_text(sh, '')
    ps = t.text_frame.paragraphs
    ps[0].runs[0].text = title
    ps[1].runs[0].text = resume


def intinya(s, text, y=9.05, h=0.85):
    return box(s, 0.9, y, 18.2, h, fill=YEL_L, line=YEL, anchor=MSO_ANCHOR.MIDDLE,
               paras=[[('Intinya  ', 14, True, TEAL), (text, 14, False, DARK)]])


def head(s, x, y, w, text, sub=None):
    box(s, x, y, w, 0.5, fill=TEAL, line=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.15,
        paras=[[(text, 13.5, True, WHITE)] + ([('   ' + sub, 11, False, 'D7EEF2')] if sub else [])])


def qa(s, x, y, w, h, rows, size=12.5):
    """rows: list of (label, text[, color])"""
    paras = []
    for r in rows:
        paras.append([(r[0] + '  ', size, True, r[2] if len(r) > 2 else TEAL), (r[1], size, False, DARK)])
    return box(s, x, y, w, h, fill=LIGHT, line=LINE, paras=paras)


def chip(s, x, y, w, h, text, fill=LIGHT, line=LINE, size=10.5, color=DARK, bold=False):
    return box(s, x, y, w, h, fill=fill, line=line, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, margin=0.06,
               paras=[(text, size, bold, color)])


def outline(s, x, y, w, h, color=RED):
    r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.background(); r.line.color.rgb = rgb(color); r.line.width = Pt(2.5); r.shadow.inherit = False
    return r


s = p.slides[0]
prep(s, 'Kesimpulan dan Saran', 'Ketiga tujuan penelitian tercapai, dengan beberapa hal yang bisa dikembangkan')
head(s, 0.9, 1.95, 10.6, 'Kesimpulan')
K = [('Tujuan 1', 'Variabel penyusun model', 'Model disusun dari 5 stok, 10 aliran, 11 variabel bantu, dan 35 parameter, dengan 6 feedback loop. Diadaptasi dari Mai & Smith (2018) sesuai data resmi Indonesia.'),
     ('Tujuan 2', 'Model dan simulasi skenario', 'Struktur model valid dan perilakunya cukup dekat dengan data. Tidak ada skenario yang terbaik di semua hal, dan urutan ini tetap sama walaupun asumsi digeser.'),
     ('Tujuan 3', 'Aplikasi web', 'Semua fungsi berjalan (9/9) dan aplikasi dinilai mudah dipakai (skor SUS 88,50).')]
for k, (a, b, c) in enumerate(K):
    y = 2.6 + k * 1.95
    box(s, 0.9, y, 3.1, 1.8, fill=TEAL, line=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.2, radius=0.05,
        paras=[[(a, 11, True, YEL)], [(b, 13, True, WHITE)]])
    box(s, 4.1, y, 7.4, 1.8, fill=LIGHT, line=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.25, radius=0.05,
        paras=[[(c, 13.5, False, DARK)]])
head(s, 11.9, 1.95, 7.2, 'Saran')
SR = [('Perkuat asumsi dengan data', 'Survei langsung untuk laju penurunan wisatawan dan batas kepadatan.'),
      ('Kembangkan model', 'Tambahkan kejadian mendadak (mis. pandemi) dan pisahkan sumber dana pembangunan objek wisata.'),
      ('Libatkan pemangku kepentingan', 'Tentukan nilai kebijakan bersama Pemda dan Dinas Pariwisata.'),
      ('Kembangkan aplikasi', 'Tambahkan pembaruan data tahunan dan ekspor laporan.')]
for k, (a, b) in enumerate(SR):
    y = 2.6 + k * 1.45
    num(s, 11.95, y + 0.35, k + 1, d=0.55, fill=CYAN, size=14)
    box(s, 12.7, y, 6.4, 1.3, fill=WHITE, line=LINE, anchor=MSO_ANCHOR.MIDDLE, margin=0.2, radius=0.05,
        paras=[[(a, 13, True, TEAL)], [(b, 12, False, DARK)]])
intinya(s, 'Model yang sudah diuji dapat dipakai untuk membandingkan skenario kebijakan pariwisata DIY, dan aplikasinya membuat hasil itu mudah diakses.', y=8.6, h=0.95)
source(s, 'Sumber: Buku Bab V.', y=10.12)
s.notes_slide.notes_text_frame.text = (
    'Sebagai kesimpulan, ketiga tujuan tercapai. Pertama, model disusun dari lima stok, sepuluh aliran, sebelas variabel bantu, dan 35 parameter, dengan enam feedback loop. '
    'Kedua, struktur model valid dan perilakunya cukup dekat dengan data; dari simulasi, tidak ada skenario yang terbaik di semua hal, dan urutannya tetap sama walaupun asumsi digeser. '
    'Ketiga, aplikasi web berjalan sesuai rancangan dan dinilai mudah dipakai. Untuk ke depan, saya menyarankan empat hal: memperkuat asumsi dengan survei, mengembangkan model, '
    'melibatkan pemangku kepentingan dalam menentukan nilai kebijakan, dan menambah fitur aplikasi. Terima kasih.')
p.save('/tmp/pptwork/v5/kesimpulan_1slide.pptx')
print('ok')
