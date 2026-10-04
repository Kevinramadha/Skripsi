import sys; sys.path.insert(0, '/tmp/pptwork/v5')
from single import make_subset
make_subset('/tmp/pptwork/v5/out7.pptx', [57, 58, 60], '/tmp/pptwork/v5/_base_sk.pptx')
exec(open('/tmp/pptwork/helpers_v4.py').read())
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
p = Presentation('/tmp/pptwork/v5/_base_sk.pptx')
GREEN, PINK, NAVY = 'C9F0D6', 'FDE2DE', '0B5E6E'
US = '/tmp/pptwork/v5/us/'


def find_title(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.top is not None and 0.6*E <= sh.top <= 1.3*E and sh.width > 9*E and sh.text_frame.text.strip():
            return sh


def prep(s, title, resume):
    t = find_title(s)
    for sh in list(s.shapes):
        keep = (sh._element is t._element) or sh.top >= 10.3*E or (sh.left >= 18.0*E and sh.top < 1.0*E) or sh.top < 0.45*E
        if not keep:
            sh._element.getparent().remove(sh._element)
    for sh in list(s.shapes):
        if sh.top > 10.3*E and 16.5*E < sh.left < 17.9*E:
            sh._element.getparent().remove(sh._element)
    for sh in [x for x in s.shapes if x.has_text_frame and x.left > 17.9*E and x.top > 10.4*E and x.text_frame.text.strip().isdigit()]:
        set_text(sh, '')
    ps = t.text_frame.paragraphs
    ps[0].runs[0].text = title
    ps[1].runs[0].text = resume


def intinya(s, text, y=9.05, h=0.85):
    return box(s, 0.9, y, 18.2, h, fill=YEL_L, line=YEL, anchor=MSO_ANCHOR.MIDDLE,
               paras=[[('Intinya  ', 14, True, TEAL), (text, 14, False, DARK)]])


def head(s, x, y, w, text, sub=None):
    box(s, x, y, w, 0.5, fill=TEAL, line=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.15,
        paras=[[(text, 13.5, True, WHITE)] + ([('   ' + sub, 11, False, 'D7EEF2')] if sub else [])])


def qa(s, x, y, w, h, rows, size=12.5):
    """rows: list of (label, text[, color])"""
    paras = []
    for r in rows:
        paras.append([(r[0] + '  ', size, True, r[2] if len(r) > 2 else TEAL), (r[1], size, False, DARK)])
    return box(s, x, y, w, h, fill=LIGHT, line=LINE, paras=paras)


def chip(s, x, y, w, h, text, fill=LIGHT, line=LINE, size=10.5, color=DARK, bold=False):
    return box(s, x, y, w, h, fill=fill, line=line, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, margin=0.06,
               paras=[(text, size, bold, color)])


def outline(s, x, y, w, h, color=RED):
    r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.background(); r.line.color.rgb = rgb(color); r.line.width = Pt(2.5); r.shadow.inherit = False
    return r


S = list(p.slides)
GRY, CY, YL = '8A979E', CYAN, YEL

# ======================= 1. rancangan
s = S[0]
prep(s, 'Skenario Kebijakan', 'Tiga skenario dibuat dari dua kebijakan yang bisa diatur pemerintah daerah')
tb(s, 0.9, 1.95, 4.0, 0.5, [[('Kenapa dua kebijakan ini?', 12.5, True, TEAL)]])
for k, t_ in enumerate(['Berpengaruh pada hasil model', 'Punya dasar Perda DIY', 'Bisa dikendalikan pemerintah']):
    chip(s, 4.9 + k * 4.75, 1.95, 4.55, 0.5, '✓  ' + t_, fill=LIGHT, size=11.5, bold=True, color=TEAL)
for k, (a, b) in enumerate([('Insentif Kebijakan', 'Tambahan dorongan investasi pariwisata. 0 berarti tanpa insentif. Dasar: Perda DIY 1/2012.'),
                            ('Konservasi Lahan', 'Seberapa besar kebutuhan lahan fasilitas wisata dipenuhi tanpa membuka lahan baru. 1 berarti sepenuhnya. Dasar: Perda DIY 6/2021.')]):
    box(s, 0.9 + k * 9.2, 2.7, 9.0, 1.1, fill=WHITE, line=LINE, anchor=MSO_ANCHOR.MIDDLE, margin=0.22,
        paras=[[(a, 12.5, True, TEAL)], [(b, 11, False, DARK)]])
SC = [('Business-as-Usual', GRY, WHITE, '0', '0', 'Seperti sekarang: tanpa insentif tambahan, dan lahan untuk fasilitas wisata tetap dibuka seperti biasa.'),
      ('Sustainable', CY, WHITE, '0,1', '1,0', 'Insentif kecil, dan fasilitas wisata baru tidak boleh membuka lahan baru.'),
      ('Development Priority', YL, DARK, '0,3', '0', 'Insentif besar untuk mendorong investasi, tanpa pembatasan lahan.')]
w = 5.9
for k, (nm, col, tcol, ins, kon, desc) in enumerate(SC):
    x = 0.9 + k * (w + 0.25)
    box(s, x, 4.05, w, 4.3, fill=LIGHT, line=LINE, radius=0.05)
    box(s, x, 4.05, w, 0.75, fill=col, line=None, radius=0.1, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, paras=[[(nm, 15, True, tcol)]])
    for j, (v, lab) in enumerate([(ins, 'Insentif'), (kon, 'Konservasi lahan')]):
        tb(s, x + 0.2 + j * (w - 0.4) / 2, 5.0, (w - 0.4) / 2, 1.2, [[(v, 34, True, TEAL)], [(lab, 11.5, False, GREY)]], align=PP_ALIGN.CENTER)
    tb(s, x + 0.3, 6.55, w - 0.6, 1.7, [[(desc, 13.5, False, DARK)]], align=PP_ALIGN.CENTER)
intinya(s, 'Semua skenario dibandingkan dengan BAU, dan dinilai per dimensi (ekonomi, lingkungan, sosial), tanpa skor gabungan.', y=8.6, h=0.95)
source(s, 'Sumber: Buku Subbab 3.8 dan 4.8 (Tabel 58–59).', y=10.12)
s.notes_slide.notes_text_frame.text = (
    'Dari analisis sensitivitas dan kajian kebijakan, ada dua kebijakan yang bisa diatur pemerintah daerah: insentif investasi pariwisata dan konservasi lahan. '
    'Keduanya berpengaruh pada model, punya dasar Perda DIY, dan memang bisa dikendalikan pemerintah. Dari dua kebijakan ini saya susun tiga skenario. '
    'Business-as-Usual adalah kondisi seperti sekarang. Sustainable memberi insentif kecil dan melarang fasilitas wisata baru membuka lahan baru. '
    'Development Priority memberi insentif besar tanpa pembatasan lahan. Ketiganya dibandingkan dengan BAU, per dimensi, tanpa digabung menjadi satu skor.')

# ======================= 2. hasil
s = S[1]
prep(s, 'Hasil Perbandingan Skenario', 'Tidak ada skenario yang paling baik di semua hal')
B1 = 'C9EEF3'
rows = [['Dimensi', 'Indikator tahun 2050', 'BAU', 'Sustainable', 'Development Priority'],
        ['Ekonomi', 'Lapangan kerja pariwisata (jiwa)', '650.355', '667.458', '697.021'],
        ['', 'PDRB pariwisata (miliar Rp)', '40.061', '41.194', '43.152'],
        ['Lingkungan', 'Lahan terbangun (ha)', '122.208', '121.094', '122.377'],
        ['', 'Lahan baru dibuka untuk pariwisata (ha)', '947,0', '0,0', '1.095,1'],
        ['', 'Rasio daya dukung lahan (aman bila > 0,331)', '0,6145', '0,6180', '0,6140'],
        ['Sosial', 'Indeks kepadatan (aman bila < 2,0)', '2,3638', '2,4307', '2,5462'],
        ['', 'Daya tarik destinasi', '0,8132', '0,8205', '0,8322'],
        ['Skala', 'Jumlah wisatawan (juta kunjungan)', '96,20', '98,92', '103,62']]
hl = {(1, 4): B1, (2, 4): B1, (3, 3): B1, (4, 3): B1, (5, 3): B1, (6, 2): B1, (7, 4): B1}
gf = table(s, 0.9, 1.95, 11.6, [1.5, 4.6, 1.6, 1.8, 2.1], rows, size=10.5, rowh=0.62, hl=hl)
T = gf.table
for r in range(len(rows)):
    for c in (2, 3, 4):
        T.cell(r, c).text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        if (r, c) in hl: T.cell(r, c).text_frame.paragraphs[0].runs[0].font.bold = True
    if r and rows[r][0]:
        T.cell(r, 0).text_frame.paragraphs[0].runs[0].font.bold = True; T.cell(r, 0).text_frame.paragraphs[0].runs[0].font.color.rgb = rgb(TEAL)
for a, b in ((1, 2), (3, 5), (6, 7)):
    T.cell(a, 0).merge(T.cell(b, 0))
tb(s, 0.9, 7.6, 11.6, 0.35, [[('Kotak biru = paling baik pada indikator tersebut. Jumlah wisatawan hanya dilaporkan, tidak dinilai.', 9.5, False, GREY, True)]])
K3 = [('Development Priority', 'unggul di ekonomi', 'Lapangan kerja +7,18% dan PDRB +7,71% dibanding BAU.', YL),
      ('Sustainable', 'unggul di lingkungan', 'Tidak ada lahan baru dibuka untuk pariwisata (0 ha, BAU 947 ha).', CY),
      ('Business-as-Usual', 'paling tidak padat', 'Indeks kepadatan terendah (2,36), karena wisatawan tumbuh paling lambat.', GRY)]
for k, (a, a2, b, col) in enumerate(K3):
    y = 1.95 + k * 1.9
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(12.9), Inches(y), Inches(0.14), Inches(1.7))
    bar.fill.solid(); bar.fill.fore_color.rgb = rgb(col); bar.line.fill.background()
    box(s, 13.04, y, 6.06, 1.7, fill=LIGHT, line=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.22, radius=0.03,
        paras=[[(a + ' ', 13, True, TEAL), (a2, 13, True, DARK)], [(b, 11.5, False, DARK)]])
box(s, 0.9, 8.1, 18.2, 1.4, fill=TEAL, line=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.3, paras=[
    [('Jadi, pilihan skenario bergantung pada prioritas pemerintah.', 14, True, YL)],
    [('Urutan ini tetap sama walaupun asumsi yang paling berpengaruh digeser ke 7 kondisi berbeda, jadi perbandingannya bisa dipegang.', 12, False, WHITE)]])
source(s, 'Sumber: Buku Subbab 4.8 (Tabel 60) dan uji kekokohan peringkat skenario (Tabel 63).', y=10.12)
s.notes_slide.notes_text_frame.text = (
    'Hasilnya, tidak ada skenario yang paling baik di semua hal. Development Priority unggul di ekonomi: lapangan kerja naik 7,18 persen dan PDRB 7,71 persen dibanding BAU. '
    'Sustainable unggul di lingkungan: tidak ada lahan baru yang dibuka untuk pariwisata. BAU justru paling tidak padat, karena wisatawannya tumbuh paling lambat. '
    'Jadi pilihan skenario bergantung pada prioritas pemerintah. Ini sekaligus menjawab pertanyaan dari analisis sensitivitas: walaupun asumsi yang paling berpengaruh digeser ke tujuh kondisi, urutan ini tidak berubah.')

# ======================= 3. kenapa
s = S[2]
prep(s, 'Batas Kepadatan Tetap Terlewati', 'Kedua kebijakan hanya menunda kepadatan berlebih, bukan mencegahnya')
pc = pic(s, MED + 'image64.png', 0.9, 2.0, w=11.0); pc.line.color.rgb = rgb(LINE); pc.line.width = Pt(1)
tb(s, 0.9, 2.0 + pc.height / E + 0.08, 11.0, 0.4, [[('Garis hijau putus-putus = batas kepadatan 2,0 (dua kali kepadatan 2025; batas peringatan dari peneliti, bukan standar resmi).', 9.5, False, GREY, True)]])
tb(s, 12.3, 1.95, 6.8, 0.45, [[('Tahun batas 2,0 terlewati', 12, True, TEAL)]])
for k, (yr, nm, col, tcol) in enumerate([('2041', 'Development Priority', YL, DARK), ('2042', 'Sustainable', CY, WHITE), ('2043', 'BAU', GRY, WHITE)]):
    box(s, 12.3 + k * 2.3, 2.45, 2.15, 1.35, fill=col, line=None, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, margin=0.05,
        paras=[[(yr, 24, True, tcol)], [(nm, 10, True, tcol)]])
P = [('Kenapa hanya bergeser 1–2 tahun?', 'Kedua kebijakan mengatur investasi dan lahan, sedangkan kepadatan ditentukan oleh jumlah wisatawan.'),
     ('Kenapa daya tarik DP paling tinggi?', 'Karena objek wisatanya bertambah (+0,0297), lebih besar dari turunnya nilai kepadatan (−0,0106).'),
     ('Bagaimana dengan lahan?', 'Masih aman: rasio daya dukung lahan sekitar 0,614, jauh di atas batas 0,331 di ketiga skenario.')]
for k, (a, b) in enumerate(P):
    box(s, 12.3, 4.0 + k * 1.5, 6.8, 1.35, fill=LIGHT, line=LINE, anchor=MSO_ANCHOR.MIDDLE, margin=0.22,
        paras=[[(a, 12, True, TEAL)], [(b, 11, False, DARK)]])
intinya(s, 'Untuk mengendalikan kepadatan wisatawan, dibutuhkan kebijakan lain di luar dua kebijakan dalam model ini.', y=8.75, h=0.95)
source(s, 'Sumber: Buku Subbab 4.8 (Tabel 60–61, Gambar 61).', y=10.12)
s.notes_slide.notes_text_frame.text = (
    'Temuan yang paling penting: di ketiga skenario, batas kepadatan 2,0 tetap terlewati, yaitu pada 2041, 2042, dan 2043. Selisihnya hanya dua tahun. '
    'Ini karena kedua kebijakan bekerja di sisi investasi dan lahan, sedangkan kepadatan ditentukan oleh jumlah wisatawan. Jadi kebijakan hanya menunda, bukan mencegah. '
    'Daya tarik Development Priority paling tinggi karena objek wisatanya bertambah, dan pengaruhnya lebih besar dari turunnya nilai kepadatan. '
    'Untuk lahan, kondisinya masih aman sampai 2050. Artinya, untuk mengendalikan kepadatan dibutuhkan kebijakan lain di luar model ini.')
p.save('/tmp/pptwork/v5/skenario_3slide.pptx')
print('ok')
