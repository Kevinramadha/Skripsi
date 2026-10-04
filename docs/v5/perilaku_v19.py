import sys; sys.path.insert(0, '/tmp/pptwork/v5')
from single import make_single
make_single('/tmp/pptwork/v5/out7.pptx', 45, '/tmp/pptwork/v5/_base_perilaku.pptx')
exec(open('/tmp/pptwork/helpers_v4.py').read())
from pptx import Presentation
p = Presentation('/tmp/pptwork/v5/_base_perilaku.pptx')
s = p.slides[0]
GREEN = 'C9F0D6'
def find_title(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.top is not None and 0.6*E <= sh.top <= 1.3*E and sh.width > 9*E and sh.text_frame.text.strip():
            return sh
t = find_title(s)
for sh in list(s.shapes):
    keep = (sh._element is t._element) or sh.top >= 10.3*E or (sh.left >= 18.0*E and sh.top < 1.0*E) or sh.top < 0.45*E
    if not keep:
        sh._element.getparent().remove(sh._element)
for sh in list(s.shapes):
    if sh.top > 10.3*E and 16.5*E < sh.left < 17.9*E:
        sh._element.getparent().remove(sh._element)
for sh in [x for x in s.shapes if x.has_text_frame and x.left > 17.9*E and x.top > 10.4*E and x.text_frame.text.strip().isdigit()]:
    set_text(sh, '')
ps = t.text_frame.paragraphs
ps[0].runs[0].text = 'Rancangan Uji Perilaku'
ps[1].runs[0].text = 'Hasil simulasi dibandingkan dengan data aktual melalui dua bentuk uji, lalu dinilai berurutan mengikuti Barlas (1989)'

def head(x, y, w, letter, text, sub):
    box(s, x, y, w, 0.5, fill=TEAL, line=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.15,
        paras=[[(letter + '  ', 13.5, True, YEL), (text, 13.5, True, WHITE), ('   ' + sub, 10.5, False, 'D7EEF2')]])

# ---------------- kiri: bentuk uji
L, LW = 0.9, 8.5
head(L, 1.95, LW, 'A', 'Uji parsial · 2016–2025', 'apakah struktur tiap subsistem sudah benar?')
tb(s, L, 2.5, LW, 0.75, [[('Tiap subsistem diuji terpisah: variabel dari subsistem lain diganti data aktual, sehingga kesalahan subsistem lain tidak ikut terbawa (partial-model testing; Homer, 2012).', 10.5, False, DARK)]])
rows = [['Kode', 'Subsistem', 'Diganti data aktual', 'Variabel dinilai'],
        ['P1', 'Akomodasi', 'Total malam menginap, investasi', 'Jumlah hotel, TPK'],
        ['P1b', 'Akomodasi (uji kekokohan P1)', 'Seperti P1 + rata-rata kamar per unit aktual', 'Jumlah hotel, TPK'],
        ['P2', 'Objek daya tarik wisata', 'Investasi', 'Jumlah ODTW'],
        ['P3', 'Tenaga kerja', 'PDRB sektor pariwisata', 'Tenaga kerja pariwisata'],
        ['P4', 'Lahan', 'Total malam menginap, investasi', 'Lahan terbangun'],
        ['P5', 'Wisatawan', 'Jumlah ODTW, lahan terbangun', 'Jumlah wisatawan*']]
gf = table(s, L, 3.2, LW, [0.7, 2.1, 3.1, 2.1], rows, size=9.5, rowh=0.36, hl={(2, c): YEL_L for c in range(4)})
for r in range(len(rows)):
    gf.table.cell(r, 0).text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    if r: gf.table.cell(r, 0).text_frame.paragraphs[0].runs[0].font.bold = True
tb(s, L, 6.05, LW, 0.35, [[('Angka utama: tanpa 2020–2021 (masih tersisa 8 dari 10 tahun data).', 10, True, TEAL)]])

head(L, 6.5, LW, 'B', 'Uji penuh · 2019–2025', 'apakah kesalahan merambat antarsubsistem?')
tb(s, L, 7.05, LW, 1.2, [
    [('Model dijalankan tanpa penggantian data; semua feedback antarsubsistem aktif. 9 variabel dinilai: ', 10.5, False, DARK),
     ('wisatawan, hotel, TPK, ODTW, tenaga kerja, lahan terbangun, PDRB pariwisata, investasi, total malam menginap.', 10.5, False, GREY)],
    [('Angka utama: termasuk 2020–2021 (tanpa itu hanya tersisa 5 tahun data).', 10, True, TEAL)]])

box(s, L, 8.35, LW, 1.6, fill=YEL_L, line=YEL, paras=[
    [('* Penyesuaian data wisnus (untuk P5)', 11, True, TEAL)],
    [('Metode pencatatan wisnus berubah 2018→2019 (8,0 → 20,5 juta). Data 2015–2018 disambung dengan faktor:', 10, False, DARK)],
    [('Faktor = (W₂₀₁₉ / W₂₀₁₈) / (1 + ḡ) = 2,2997', 11, True, DARK, True), ('   (rentang 2,22–2,38)', 10, False, GREY)],
    [('ḡ = rata-rata pertumbuhan tahunan sebelum (2015–2018) dan sesudah (2021–2025) patahan', 9.5, False, GREY, True)]])

# ---------------- kanan: ukuran penilaian
R, RW = 9.7, 9.4
box(s, R, 1.95, RW, 0.5, fill=TEAL, line=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.15,
    paras=[[('Ukuran penilaian', 13.5, True, WHITE), ('   dibaca berurutan 1 → 4, lalu dua pelengkap', 10.5, False, 'D7EEF2')]])
M = [['Ukuran', 'Rumus', 'Pedoman'],
     ['1 · Uji tren\n(saringan awal)', 't = |bₛ − bₐ| / √(Var(bₛ) + Var(bₐ))\ndb = 2n − 4', 'α = 0,05: tidak berbeda nyata / berbeda nyata / berlawanan arah. Jika berbeda, E1–E2 tidak dibaca'],
     ['2 · E1\n(selisih rata-rata)', 'E1 = |S̄ − Ā| / Ā', '< 0,05 (pedoman empiris)'],
     ['3 · E2\n(selisih variasi)', 'E2 = |σₛ − σₐ| / σₐ', '< 0,30 (pedoman empiris)'],
     ['4 · Discrepancy coefficient (DC)', 'DC = SD(e) / (SD(A) + SD(S))\neᵢ = Aᵢ − Sᵢ', '0,4–0,7 = rata-rata sampai baik (ringkasan, bukan uji)'],
     ['Pelengkap · MAPE', 'MAPE = (100% / n) · Σ |Sᵢ − Aᵢ| / Aᵢ', '< 10% sangat baik · 10–20% baik · 20–50% layak · > 50% buruk'],
     ['Pelengkap · Dekomposisi error (Sterman, 1984)', 'U1 = (S̄ − Ā)² / MSE\nU2 = (σₛ − σₐ)² / MSE\nU3 = 2(1 − r)σₛσₐ / MSE', 'U1 = beda rata-rata · U2 = beda naik-turun · U3 = beda tanpa pola. U1/U2 besar → simulasi meleset ke satu arah secara konsisten, perlu ditelusuri; U3 besar → selisih tanpa pola, wajar']]
rh = [0.45, 1.1, 0.75, 0.75, 1.0, 0.95, 1.35]
gs = s.shapes.add_table(len(M), 3, Inches(R), Inches(2.5), Inches(RW), Inches(sum(rh)))
T = gs.table
for j, cw in enumerate([2.35, 3.55, 3.5]):
    T.columns[j].width = Inches(cw)
for i, row in enumerate(M):
    T.rows[i].height = Inches(rh[i])
    for j, val in enumerate(row):
        c = T.cell(i, j)
        c.margin_left = c.margin_right = Inches(0.1); c.margin_top = c.margin_bottom = Inches(0.04)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        c.fill.solid(); c.fill.fore_color.rgb = rgb(TEAL if i == 0 else ('F3F7F8' if i >= 5 else (WHITE if i % 2 else LIGHT)))
        tf = c.text_frame; tf.word_wrap = True
        for k, line in enumerate(val.split('\n')):
            pp = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            r = pp.add_run(); r.text = line
            if i == 0: font(r, 10.5, True, WHITE)
            elif j == 0: font(r, 10, True, TEAL if i < 5 else GREY)
            elif j == 1: font(r, 10, False, DARK, True)
            else: font(r, 9.5, False, DARK)
tb(s, R, 2.5 + sum(rh) + 0.08, RW, 0.6, [[('Keterangan: ', 9.5, True, TEAL),
    ('S = simulasi · A = data aktual · S̄, Ā = rata-rata · σ = simpangan baku · b = kemiringan regresi · n = jumlah tahun · r = korelasi S dan A · MSE = rata-rata kuadrat selisih', 9.5, False, GREY)]])
source(s, 'Sumber: Buku Subbab 3.7.2 dan 3.7.4 (Tabel 18), Subbab 4.5.1–4.5.2 (Tabel 45–46).', y=10.12)
s.notes_slide.notes_text_frame.text = (
    'Uji perilaku membandingkan hasil simulasi dengan data aktual dalam dua bentuk. Uji parsial menguji tiap subsistem secara terpisah pada 2016 sampai 2025, '
    'dengan mengganti variabel dari subsistem lain menggunakan data aktual. Ada P1 sampai P5, ditambah P1b, yaitu variasi P1 yang memakai rata-rata kamar per unit aktual sebagai uji kekokohan. '
    'Uji penuh menjalankan model tanpa penggantian data pada 2019 sampai 2025 dan menilai sembilan variabel, untuk melihat apakah kesalahan merambat antarsubsistem. '
    'Penilaian dibaca berurutan mengikuti Barlas: uji tren dulu sebagai saringan, lalu E1, E2, dan terakhir discrepancy coefficient. MAPE dan dekomposisi error dipakai sebagai pelengkap, '
    'terutama untuk melihat apakah simulasi meleset ke satu arah secara konsisten, atau selisihnya naik-turun tanpa pola. Untuk P5, data wisnus 2015–2018 disambung dengan faktor 2,2997 karena perubahan metode pencatatan pada 2019.')
p.save('/tmp/pptwork/v5/perilaku_1slide.pptx')
print('ok')
