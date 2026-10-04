import sys; sys.path.insert(0, '/tmp/pptwork/v5')
from single import make_subset
make_subset('/tmp/pptwork/v5/out7.pptx', [27, 28, 29], '/tmp/pptwork/v5/_base3.pptx')
exec(open('/tmp/pptwork/helpers_v4.py').read())
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
p = Presentation('/tmp/pptwork/v5/_base3.pptx')
NAVY = '0B5E6E'
def find_title(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.top is not None and 0.6*E <= sh.top <= 1.3*E and sh.width > 9*E and sh.text_frame.text.strip():
            return sh
def prep(s, title, resume):
    t = find_title(s)
    for sh in list(s.shapes):
        keep = (sh._element is t._element) or sh.top >= 10.3*E or (sh.left >= 18.0*E and sh.top < 1.0*E) or sh.top < 0.45*E
        if not keep:
            sh._element.getparent().remove(sh._element)
    for sh in list(s.shapes):
        if sh.top > 10.3*E and 16.5*E < sh.left < 17.9*E:
            sh._element.getparent().remove(sh._element)
    ps = t.text_frame.paragraphs
    ps[0].runs[0].text = title
    for r in ps[0].runs[1:]: r._r.getparent().remove(r._r)
    ps[1].runs[0].text = resume
    for r in ps[1].runs[1:]: r._r.getparent().remove(r._r)
    for k, sh in enumerate([x for x in s.shapes if x.has_text_frame and x.left > 17.9*E and x.top > 10.4*E and x.text_frame.text.strip().isdigit()]):
        set_text(sh, '')
def center_cols(gf, cols, start=0):
    tbl = gf.table
    for r in range(start, len(tbl.rows)):
        for k in cols:
            tbl.cell(r, k).text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

S = list(p.slides)
# ================= Slide 1: metrik evaluasi
s = S[0]
prep(s, 'Estimasi Jumlah ODTW — Metrik Evaluasi', 'Estimasi ODTW dari NTL memenuhi kedua kriteria kelayakan')
box(s, 0.9, 1.95, 18.2, 0.62, fill=NAVY, line=None, paras=[[('Model yang dievaluasi  ', 14, True, YEL),
    ('ODTW = 121,61 + 34,59 × NTL   ·   regresi linear sederhana   ·   n = 7 tahun (data BPS 2018–2024)', 14, True, WHITE)]], anchor=MSO_ANCHOR.MIDDLE)
rows = [['Metrik', 'Rumus', 'Kriteria', 'Hasil', 'Keputusan'],
 ['Korelasi pada nilai level', 'r = Σ(xᵢ − x̄)(yᵢ − ȳ) / √[Σ(xᵢ − x̄)² · Σ(yᵢ − ȳ)²]', 'Hanya pembanding (bisa tinggi karena tren yang sama)', 'R² = 0,785', 'Pembanding'],
 ['Korelasi pada selisih tahunan', 'Δxₜ = xₜ − xₜ₋₁ ;  Δyₜ = yₜ − yₜ₋₁ ;  r dihitung pada Δx dan Δy', 'R² ≥ 0,30', 'R² = 0,405', 'Layak'],
 ['LOOCV (RMSE relatif)', 'RMSE = √[(1/k) · Σ(ŷ₋ᵢ − yᵢ)²] ;  RMSE relatif = RMSE / ȳ × 100%', '≤ 10%', 'RMSE 11,29 unit → 5,99%', 'Layak'],
 ['Selang prediksi 95%', 'ŷ₀ ± t₀,₉₇₅;ₙ₋₂ · s · √[1 + 1/n + (x₀ − x̄)² / Sₓₓ]', 'Tidak ada batas; menunjukkan ketidakpastian estimasi', '2025: 176,7–224,8 · 2016 terlebar (≈58 unit)', 'Dasar rentang uji sensitivitas (189–218)']]
hl = {(2, 4): 'C9F0D6', (3, 4): 'C9F0D6', (1, 4): LIGHT, (4, 4): YEL_L}
gf = table(s, 0.9, 2.75, 18.2, [2.7, 6.6, 3.4, 2.9, 2.6], rows, size=12.5, rowh=0.9, hl=hl, bold_first_col=True)
center_cols(gf, [3, 4], 1)
for r in (2, 3):
    gf.table.cell(r, 4).text_frame.paragraphs[0].runs[0].font.bold = True
tb(s, 0.9, 7.35, 18.2, 0.4, [('x = NTL · y = ODTW BPS · ŷ₋ᵢ = prediksi tahun i dari model tanpa tahun i · k = 7 iterasi · s = standard error residual · Sₓₓ = jumlah kuadrat simpangan NTL', 11, False, GREY, True)])
rows2 = [['Error LOOCV per tahun', '2018', '2019', '2020', '2021', '2022', '2023', '2024'],
         ['Selisih prediksi vs aktual (%)', '+1,6', '+1,2', '+3,2', '+4,0', '−8,2', '+7,0', '−8,9']]
gf = table(s, 0.9, 7.85, 12.0, [3.6, 1.2, 1.2, 1.2, 1.2, 1.2, 1.2, 1.2], rows2, size=12, rowh=0.5, bold_first_col=True,
           hl={(1, 5): YEL_L, (1, 7): YEL_L})
center_cols(gf, range(1, 8))
box(s, 13.2, 7.85, 5.9, 1.0, fill=YEL_L, line=YEL, paras=[('Error terbesar pada 2022 dan 2024, tetapi semuanya masih di bawah 10%.', 12, False, DARK)], anchor=MSO_ANCHOR.MIDDLE)
source(s, 'Sumber: Buku Subbab 3.5.1 dan 4.2.1, Tabel 29–30; repository Pengujian_NTL_ODTW.csv dan LOOCV_NTL_ODTW.csv.', y=10.12)
s.notes_slide.notes_text_frame.text = ('Kelayakan estimasi ODTW dari NTL dinilai dengan dua kriteria. Pertama, korelasi pada selisih tahunan, bukan pada nilai level, '
    'karena korelasi level bisa tinggi hanya karena kedua data sama-sama naik. Hasilnya R kuadrat 0,405, di atas batas 0,30. '
    'Kedua, LOOCV: satu tahun dikeluarkan bergantian lalu diprediksi; RMSE relatifnya 5,99 persen, di bawah batas 10 persen. '
    'Selang prediksi 95 persen menunjukkan ketidakpastian setiap estimasi, dan dipakai sebagai dasar rentang uji sensitivitas nilai awal ODTW.')

# ================= Slide 2: validasi temporal
s = S[1]
prep(s, 'Estimasi Jumlah ODTW — Validasi Temporal', 'NTL dan ODTW bergerak searah pada sebagian besar tahun 2018–2024')
yrs = [str(y) for y in range(2015, 2026)]
NTL = [1.262, 1.054, 1.438, 1.668, 2.005, 1.828, 1.541, 1.483, 2.512, 2.457, 2.288]
cd = CategoryChartData(); cd.categories = yrs; cd.add_series('NTL (nW/cm²/sr)', NTL)
ch = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.9), Inches(2.0), Inches(11.2), Inches(3.55), cd).chart
cd2 = CategoryChartData(); cd2.categories = yrs
cd2.add_series('ODTW BPS', [None, None, None, 177, 189, 180, 170, 183, 201, 218, None])
cd2.add_series('ODTW estimasi NTL', [165.2, 158.1, 171.4, None, None, None, None, None, None, None, 200.7])
ch2 = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.9), Inches(5.6), Inches(11.2), Inches(3.75), cd2).chart
for c, cols, fmt, lo, hi in [(ch, [CYAN], '0.00', 0.8, 2.8), (ch2, [NAVY, 'E0A800'], '0', 140, 230)]:
    c.has_title = False; c.font.size = Pt(11); c.font.name = F; c.font.color.rgb = rgb(DARK)
    c.has_legend = True; c.legend.position = XL_LEGEND_POSITION.TOP; c.legend.include_in_layout = False
    for ser, col in zip(c.plots[0].series, cols):
        ser.format.line.color.rgb = rgb(col); ser.format.line.width = Pt(2.75); ser.smooth = False
        ser.marker.size = 8; ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = rgb(col)
        ser.marker.format.line.color.rgb = rgb(col)
    c.plots[0].has_data_labels = True
    dl = c.plots[0].data_labels; dl.number_format = fmt; dl.number_format_is_linked = False
    dl.font.size = Pt(10); dl.position = XL_LABEL_POSITION.ABOVE
    c.value_axis.minimum_scale = lo; c.value_axis.maximum_scale = hi
    c.value_axis.has_major_gridlines = True; c.value_axis.major_gridlines.format.line.color.rgb = rgb('E3EEF1')
    c.value_axis.tick_labels.font.size = Pt(10); c.category_axis.tick_labels.font.size = Pt(10)
ch2.plots[0].series[1].format.line.dash_style = 4
rows = [['Periode', 'Arah NTL', 'Arah ODTW', 'Sesuai?'],
        ['2018→2019', 'Naik', 'Naik', 'Searah'], ['2019→2020', 'Turun', 'Turun', 'Searah'],
        ['2020→2021', 'Turun', 'Turun', 'Searah'], ['2021→2022', 'Turun', 'Naik', 'Tidak searah'],
        ['2022→2023', 'Naik', 'Naik', 'Searah'], ['2023→2024', 'Turun', 'Naik', 'Tidak searah']]
hl = {}
for r in range(1, 7):
    hl[(r, 3)] = 'C9F0D6' if rows[r][3] == 'Searah' else 'FDE2DE'
gf = table(s, 12.5, 2.0, 6.6, [1.8, 1.5, 1.5, 1.8], rows, size=12.5, rowh=0.5, hl=hl, bold_first_col=True)
center_cols(gf, [1, 2, 3])
box(s, 12.5, 5.75, 6.6, 1.65, fill=LIGHT, paras=[('4 dari 6 periode searah', 15, True, TEAL),
    ('Naik bersama pada 2019 dan melonjak bersama pada 2023.', 12, False, DARK)])
box(s, 12.5, 7.55, 6.6, 1.8, fill=YEL_L, line=YEL, paras=[('Tidak searah: 2021–2022 dan 2023–2024', 14, True, TEAL),
    ('Sama dengan tahun error LOOCV terbesar (2022 −8,2%; 2024 −8,9%). Tahun estimasi (kuning) bukan bukti karena nilainya berasal dari NTL.', 12, False, DARK)])
source(s, 'Sumber: Buku Subbab 3.5.1 dan 4.2.1, Gambar 11; repository NTL_DIY_2015_2025.csv dan ODTW_DIY_2015_2025_Final.csv. Arah dihitung dari selisih tahunan.', y=9.55)
s.notes_slide.notes_text_frame.text = ('Validasi temporal menyandingkan data NTL dan ODTW. Pada periode data BPS, 2018 sampai 2024, keduanya searah pada empat dari enam periode, '
    'termasuk naik bersama pada 2019 dan melonjak bersama pada 2023. Yang tidak searah adalah 2021 ke 2022 dan 2023 ke 2024, ketika NTL turun tetapi ODTW naik; '
    'ini sama dengan tahun error LOOCV terbesar. Tahun estimasi tidak dihitung sebagai bukti, karena nilainya memang berasal dari NTL.')

# ================= Slide 3: validasi spasial
s = S[2]
prep(s, 'Estimasi Jumlah ODTW — Validasi Spasial', 'NTL tinggi di kawasan padat dan rendah di kawasan minim aktivitas')
pic_fit(s, '/home/user/Skripsi/Estimasi ODTW menggunakan NTL/Gambar_Validasi_Spasial_NTL_DIY_2025.png', 0.9, 1.95, 9.3, 7.2)
rows = [['Titik sampel (radius 300 m)', 'Karakteristik', 'NTL 2025', 'Sesuai harapan?'],
        ['Malioboro–Keraton', 'Wisata, perkotaan padat', '37,376', 'Ya, tertinggi'],
        ['Kotagede', 'Permukiman padat', '12,631', 'Ya'],
        ['Pantai Baron', 'Wisata pantai, bangunan jarang', '1,692', 'Ya, rendah'],
        ['Kawasan Industri Sentolo', 'Industri tahap awal', '1,161', 'Kurang (lebih rendah dari Baron)'],
        ['TN Gunung Merapi', 'Hutan, minim aktivitas', '0,959', 'Ya, terendah']]
hl = {(r, 3): 'C9F0D6' for r in (1, 2, 3, 5)}; hl[(4, 3)] = YEL_L
gf = table(s, 10.5, 1.95, 8.6, [2.6, 2.6, 1.2, 2.2], rows, size=12, rowh=0.62, hl=hl, bold_first_col=True)
center_cols(gf, [2])
box(s, 10.5, 5.85, 8.6, 1.15, fill=LIGHT, paras=[('Cara', 13.5, True, TEAL),
    ('Nilai NTL 2025 di lima titik dibandingkan dengan kondisi lapangan pada citra dasar (basemap) Google Earth Engine.', 12, False, DARK)])
box(s, 10.5, 7.15, 8.6, 1.0, fill=LIGHT, paras=[('Hasil', 13.5, True, TEAL),
    ('Kawasan padat jauh lebih terang (37,4 dan 12,6) daripada kawasan hutan (0,96).', 12, False, DARK)])
box(s, 10.5, 8.3, 8.6, 1.0, fill=YEL_L, line=YEL, paras=[('Keterbatasan', 13.5, True, TEAL),
    ('NTL kurang tajam membedakan industri tahap awal dan wisata skala kecil (Sentolo 1,16 vs Baron 1,69).', 12, False, DARK)])
source(s, 'Sumber: Buku Subbab 3.5.1 dan 4.2.1, Gambar 12; repository Validasi_Spasial_NTL_DIY_2025.csv (VIIRS DNB, Google Earth Engine).', y=9.55)
s.notes_slide.notes_text_frame.text = ('Validasi spasial memeriksa apakah nilai NTL masuk akal di lapangan. Saya mengambil lima titik dengan karakter berbeda. '
    'Malioboro–Keraton paling terang, 37,4, lalu Kotagede 12,6, sedangkan TN Gunung Merapi paling gelap, 0,96. Ini sesuai harapan. '
    'Keterbatasannya, Kawasan Industri Sentolo justru sedikit lebih gelap dari Pantai Baron, karena kawasan industrinya masih tahap awal; '
    'artinya NTL kurang tajam untuk aktivitas skala kecil.')
p.save('/tmp/pptwork/v5/odtw_evaluasi_3slide.pptx')
print('ok', len(p.slides))
