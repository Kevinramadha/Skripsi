import sys; sys.path.insert(0, '/tmp/pptwork/v5')
from single import make_single
make_single('/tmp/pptwork/v5/out7.pptx', 47, '/tmp/pptwork/v5/_base_hasil.pptx')
exec(open('/tmp/pptwork/helpers_v4.py').read())
from pptx import Presentation
p = Presentation('/tmp/pptwork/v5/_base_hasil.pptx')
s = p.slides[0]
GREEN = 'C9F0D6'
def find_title(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.top is not None and 0.6*E <= sh.top <= 1.3*E and sh.width > 9*E and sh.text_frame.text.strip():
            return sh
t = find_title(s)
for sh in list(s.shapes):
    keep = (sh._element is t._element) or sh.top >= 10.3*E or (sh.left >= 18.0*E and sh.top < 1.0*E) or sh.top < 0.45*E
    if not keep:
        sh._element.getparent().remove(sh._element)
for sh in list(s.shapes):
    if sh.top > 10.3*E and 16.5*E < sh.left < 17.9*E:
        sh._element.getparent().remove(sh._element)
for sh in [x for x in s.shapes if x.has_text_frame and x.left > 17.9*E and x.top > 10.4*E and x.text_frame.text.strip().isdigit()]:
    set_text(sh, '')
ps = t.text_frame.paragraphs
ps[0].runs[0].text = 'Rancangan Uji Perilaku'
ps = t.text_frame.paragraphs
ps[0].runs[0].text = 'Hasil Uji Parsial dan Uji Penuh'
ps[1].runs[0].text = 'Tiap subsistem cukup baik saat diuji sendiri; saat dijalankan bersama, error dari subsistem wisatawan merambat ke subsistem lain'
PINK = 'FDE2DE'
X, W = 0.9, 12.3
# header 2 tingkat
box(s, X + 2.35, 1.95, 4.75, 0.42, fill=CYAN, line=None, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, margin=0.05,
    paras=[[('Uji parsial · 2016–2025 (tanpa 2020–2021)', 11, True, WHITE)]])
box(s, X + 7.15, 1.95, 5.15, 0.42, fill=TEAL, line=None, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, margin=0.05,
    paras=[[('Uji penuh · 2019–2025 (termasuk 2020–2021)', 11, True, WHITE)]])
D = '—'
rows = [['Variabel', 'Uji', 'Tren', 'DC', 'MAPE', 'Dom.', 'Tren', 'DC', 'MAPE', 'Dom.'],
        ['Jumlah wisatawan', 'P5', 'Berbeda nyata', '0,4057*', D, 'U1+U2', 'Berbeda nyata', '0,4482', '12,06%', 'U2'],
        ['Jumlah hotel', 'P1', 'Sama', '0,2011', '6,05%', 'U3', 'Berlawanan arah', '0,8228', '13,16%', 'U1'],
        ['TPK', 'P1', 'Sama', '0,6692', '17,37%', 'U2', 'Sama', '0,6224', '18,98%', 'U1'],
        ['Jumlah ODTW', 'P2', 'Sama', '0,3367', '8,48%', 'U1', 'Sama', '0,7855', '5,68%', 'U2'],
        ['Tenaga kerja', 'P3', 'Sama', '0,4945', '12,88%', 'U1', 'Berlawanan arah', '0,8860', '12,39%', 'U1'],
        ['Lahan terbangun', 'P4', 'Sama', '0,4813', '12,94%', 'U1, U3', 'Sama', '0,6619', '13,43%', 'U3'],
        ['PDRB pariwisata', D, D, D, D, D, 'Sama', '0,3062', '25,06%', 'U1'],
        ['Investasi pariwisata', D, D, D, D, D, 'Sama', '0,7356', '33,54%', 'U1'],
        ['Total malam menginap', D, D, D, D, D, 'Sama', '0,7041', '29,52%', 'U1']]
hl = {}
for r, row in enumerate(rows[1:], 1):
    for c in (2, 6):
        if row[c] in ('Berbeda nyata', 'Berlawanan arah'): hl[(r, c)] = PINK
    for c in (3, 7):
        try:
            if float(row[c].strip('*').replace(',', '.')) > 0.7: hl[(r, c)] = PINK
        except ValueError: pass
    for c in (4, 8):
        if row[c].endswith('%') and float(row[c][:-1].replace(',', '.')) > 20: hl[(r, c)] = YEL_L
colw = [2.35, 0.55, 1.35, 0.85, 0.95, 1.05, 1.5, 0.85, 0.95, 0.85]
gf = table(s, X, 2.42, W, colw, rows, size=9.5, rowh=0.5, hl=hl, bold_first_col=True)
T = gf.table
for r in range(len(rows)):
    for c in range(1, 10):
        T.cell(r, c).text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    for c in (1, 2, 3, 4, 5):
        if r == 0: T.cell(r, c).fill.fore_color.rgb = rgb('2A9DB0')
y = 2.42 + 0.5 * len(rows) + 0.08
box(s, X, y, W, 0.62, fill=YEL_L, line=YEL, anchor=MSO_ANCHOR.MIDDLE, margin=0.12, paras=[[
    ('P1b (uji kekokohan P1, rata-rata kamar per unit aktual): ', 10, True, TEAL),
    ('hotel DC 0,2319, MAPE 7,86% · TPK DC 0,7053, MAPE 19,56% → pola serupa P1', 10, False, DARK)]])
tb(s, X, y + 0.7, W, 0.75, [[('Sama = tren simulasi tidak berbeda nyata dari data. Merah muda = tren berbeda/berlawanan atau DC > 0,7. Kuning = MAPE layak (20–50%). '
    'Jika tren berbeda, E1–E2 tidak dibaca; nilai E1–E2 lengkap pada Tabel 46–47. *DC P5 dari notebook repository (tidak ada di tabel buku).', 9, False, GREY, True)]])

R, RW = 13.45, 5.65
def card(y, h, title, col, paras):
    box(s, R, y, RW, h, fill=LIGHT, line=LINE, paras=[[(title, 12, True, col)]] + paras)
card(1.95, 2.2, 'Uji parsial: tiap subsistem cukup baik', TEAL, [
    [('MAPE seluruhnya < 20%. Hanya wisatawan (P5) yang gagal uji tren: simulasi ≈6,4%/tahun vs data ≈11,7%/tahun; error sistematis (U1+U2 ≈ 0,96).', 11.5, False, DARK)]])
card(4.25, 2.5, 'Uji penuh: error merambat', RED, [
    [('DC naik saat semua subsistem aktif: hotel 0,20 → 0,82; ODTW 0,34 → 0,79; tenaga kerja 0,49 → 0,89.', 11.5, False, DARK)],
    [('Penyebab: wisatawan tumbuh lebih lambat dari data → pengeluaran, PDRB, dan investasi lebih rendah → hotel dan ODTW yang dibangun lebih sedikit.', 11.5, False, DARK)]])
card(6.85, 2.05, 'Sumbernya satu, bukan tersebar', TEAL, [
    [('MAPE uji penuh: 1 sangat baik, 5 baik, 3 layak, 0 buruk. U1 (bias level) dominan pada 6 dari 9 variabel → error berasal dari satu sumber bersama.', 11.5, False, DARK)],
    [('TPK justru membaik (DC 0,6692 → 0,6224).', 11.5, False, DARK)]])
box(s, X, 9.05, 18.2, 0.85, fill=YEL_L, line=YEL, anchor=MSO_ANCHOR.MIDDLE,
    paras=[[('Intinya  ', 14, True, TEAL), ('Struktur tiap subsistem memadai; kelemahan utama ada pada pertumbuhan wisatawan, dan inilah yang ditelusuri lewat kalibrasi.', 14, False, DARK)]])
source(s, 'Sumber: Buku Subbab 4.5.2–4.5.3, Tabel 46–47. DC = discrepancy coefficient (0,4–0,7 = rata-rata sampai baik; Barlas, 1989). Dom. = komponen error dominan.', y=10.12)
s.notes_slide.notes_text_frame.text = (
    'Slide ini membandingkan hasil uji parsial dan uji penuh. Pada uji parsial, saat input dari subsistem lain diganti data aktual, setiap subsistem bisa mengikuti pola datanya, dengan MAPE seluruhnya di bawah 20 persen. '
    'Satu-satunya yang gagal uji tren adalah subsistem wisatawan: simulasi tumbuh sekitar 6,4 persen per tahun, sedangkan data sekitar 11,7 persen, dan error-nya sistematis. '
    'P1b sebagai uji kekokohan P1 memberi pola yang serupa. Pada uji penuh, DC hotel, ODTW, dan tenaga kerja naik tajam. Ini bukan karena struktur subsistemnya salah, '
    'tetapi karena wisatawan yang tumbuh lebih lambat membuat pengeluaran, PDRB, dan investasi ikut rendah, sehingga hotel dan ODTW yang dibangun lebih sedikit. '
    'MAPE uji penuh tetap wajar, dan U1 dominan pada enam dari sembilan variabel, artinya error berasal dari satu sumber bersama. Sumber inilah yang ditelusuri lewat kalibrasi.')
p.save('/tmp/pptwork/v5/hasil_1slide.pptx')
print('ok')
