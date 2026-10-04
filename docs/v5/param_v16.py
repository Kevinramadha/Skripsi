import sys; sys.path.insert(0, '/tmp/pptwork/v5')
from single import make_single
make_single('/tmp/pptwork/v5/out7.pptx', 38, '/tmp/pptwork/v5/_base_param.pptx')
exec(open('/tmp/pptwork/helpers_v4.py').read())
from pptx import Presentation
p = Presentation('/tmp/pptwork/v5/_base_param.pptx')
s = p.slides[0]
def find_title(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.top is not None and 0.6*E <= sh.top <= 1.3*E and sh.width > 9*E and sh.text_frame.text.strip():
            return sh
t = find_title(s)
for sh in list(s.shapes):
    keep = (sh._element is t._element) or sh.top >= 10.3*E or (sh.left >= 18.0*E and sh.top < 1.0*E) or sh.top < 0.45*E
    if not keep:
        sh._element.getparent().remove(sh._element)
for sh in list(s.shapes):
    if sh.top > 10.3*E and 16.5*E < sh.left < 17.9*E:
        sh._element.getparent().remove(sh._element)
for sh in [x for x in s.shapes if x.has_text_frame and x.left > 17.9*E and x.top > 10.4*E and x.text_frame.text.strip().isdigit()]:
    set_text(sh, '')
ps = t.text_frame.paragraphs
ps[0].runs[0].text = 'Penetapan Parameter'
ps[1].runs[0].text = '35 parameter dari lima jenis penetapan; hanya parameter yang belum pasti yang dikalibrasi'
CAT = {'Data resmi': ('D7EEF2', NAVY := '0B5E6E'), 'Stok-aliran': ('CDEFF5', '0B5E6E'), 'Regulasi': ('E3F1D9', '2F6B1F'),
       'Asumsi': ('FFF1BF', '7A5A00'), 'Normalisasi': ('E6E9EC', '3E4A50'), 'Kebijakan': ('FDE2DE', '9B2C1F')}
S10, RP, S10RP, NO = '±10%', 'Rentang penuh', '±10% & rentang penuh', '—'
P = [
 ('Pengeluaran per kunjungan', 'Data resmi', NO, S10), ('Rasio nilai tambah pariwisata', 'Data resmi', NO, S10),
 ('Rasio investasi terhadap PDRB', 'Data resmi', NO, S10), ('Proporsi wisatawan menginap', 'Data resmi', NO, S10),
 ('Rata-rata lama menginap', 'Data resmi', NO, S10), ('Penghunian ganda kamar (TPG)', 'Data resmi', NO, S10),
 ('Rata-rata kamar per unit akomodasi', 'Data resmi', NO, S10), ('Malam tersedia per kamar', 'Data resmi', NO, S10),
 ('Luas lahan tersedia', 'Data resmi', NO, S10), ('Intensitas tenaga kerja awal', 'Data resmi', NO, S10),
 ('Laju kenaikan produktivitas', 'Data resmi', NO, S10RP),
 ('Laju pertumbuhan eksternal (LPE)', 'Stok-aliran', 'Tahap 3 (diagnostik)', S10), ('Laju penurunan dasar (LPD)', 'Stok-aliran', 'Tahap 3 (diagnostik)', S10RP + '*'),
 ('Sensitivitas konstruksi thd TPK', 'Stok-aliran', 'Ikut Tahap 1', S10), ('Batas TPK pemicu konstruksi', 'Stok-aliran', 'Tahap 1', S10RP),
 ('Sensitivitas ODTW thd investasi', 'Stok-aliran', NO, S10), ('Pembangunan ODTW non-investasi', 'Stok-aliran', NO, S10),
 ('Laju penutupan dasar ODTW', 'Stok-aliran', NO, S10RP), ('Laju konversi dasar lahan', 'Stok-aliran', 'Tahap 2', S10RP + '*'),
 ('Laju demolisi dasar', 'Regulasi', 'Tahap 1', S10RP), ('Laju keluar dasar tenaga kerja', 'Regulasi', NO, S10RP),
 ('Bobot ODTW', 'Asumsi', NO, S10RP + '*'), ('Bobot kepadatan', 'Asumsi', NO, S10RP + '*'), ('Bobot daya dukung lahan', 'Asumsi', NO, S10RP + '*'),
 ('Elastisitas daya tarik ODTW', 'Asumsi', NO, S10RP), ('Batas maksimum efek ODTW', 'Asumsi', NO, S10RP),
 ('Lahan per hotel', 'Asumsi', NO, S10RP), ('Lahan per ODTW', 'Asumsi', NO, S10RP), ('Waktu penyesuaian tenaga kerja', 'Asumsi', NO, S10RP),
 ('ODTW referensi', 'Normalisasi', NO, S10RP), ('Kepadatan referensi', 'Normalisasi', NO, S10), ('RDDL referensi', 'Normalisasi', NO, S10),
 ('Tahun dasar intensitas tenaga kerja', 'Normalisasi', NO, NO),
 ('Insentif Kebijakan', 'Kebijakan', 'Diatur per skenario', RP), ('Kebijakan Konservasi Lahan', 'Kebijakan', 'Diatur per skenario', RP)]
assert len(P) == 35
hdr = ['No', 'Parameter', 'Jenis penetapan', 'Dikalibrasi?', 'Uji sensitivitas']
left, right = P[:18], P[18:]
for (chunk, x, start) in [(left, 0.9, 1), (right, 10.15, 19)]:
    rows = [hdr] + [[str(start + i), a, b, c, d] for i, (a, b, c, d) in enumerate(chunk)]
    hl = {}
    for r, (a, b, c, d) in enumerate(chunk, 1):
        hl[(r, 2)] = CAT[b][0]
        if c not in (NO, 'Diatur per skenario'): hl[(r, 3)] = 'C9F0D6'
    gf = table(s, x, 1.95, 8.95, [0.45, 2.9, 1.45, 1.85, 2.3], rows, size=9.5, rowh=0.37, hl=hl)
    T = gf.table
    for r in range(len(rows)):
        for k in (0, 3, 4):
            T.cell(r, k).text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        if r:
            run = T.cell(r, 2).text_frame.paragraphs[0].runs[0]
            run.font.bold = True; run.font.color.rgb = rgb(CAT[chunk[r - 1][1]][1])
tb(s, 0.9, 9.5, 18.2, 0.55, [[('Keterangan  ', 10.5, True, TEAL),
    ('±10% = dinaikkan/diturunkan 10% satu per satu (32 parameter) · Rentang penuh = diuji pada seluruh rentang nilainya (17 kelompok) · '
     '* juga diuji pada robustness peringkat skenario (C0–C6) · Tahap = tahap kalibrasi (hasilnya tidak ada nilai yang diganti) · Ikut Tahap 1 = dihitung ulang mengikuti batas TPK', 10.5, False, DARK)]])
source(s, 'Sumber: Buku Subbab 3.4.2, 3.7.3–3.7.5, Tabel 12; Subbab 4.6–4.7 (Tabel 56–57); Lampiran 2 (nilai parameter).', y=10.12)
s.notes_slide.notes_text_frame.text = ('Tiga puluh lima parameter saya kelompokkan menurut cara penetapannya: data resmi, persamaan stok-aliran, regulasi, asumsi pemodelan, dan normalisasi, ditambah dua variabel kebijakan. '
    'Jenis penetapan ini menentukan perlakuannya. Hanya parameter yang belum pasti yang dikalibrasi: batas TPK dan laju demolisi di Tahap 1, laju konversi di Tahap 2, serta LPE dan LPD diperiksa di Tahap 3. '
    'Hasilnya tidak ada nilai yang diganti. Hampir semua parameter diuji sensitivitas plus-minus 10 persen, dan parameter asumsi diuji pada seluruh rentang nilainya. '
    'Nilai setiap parameter ada di Lampiran 2.')
p.save('/tmp/pptwork/v5/param_1slide.pptx')
print('ok')
