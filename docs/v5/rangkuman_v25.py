import sys; sys.path.insert(0, '/tmp/pptwork/v5')
from single import make_subset
make_subset('/tmp/pptwork/v5/out7.pptx', [53], '/tmp/pptwork/v5/_base_rk.pptx')
exec(open('/tmp/pptwork/helpers_v4.py').read())
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
p = Presentation('/tmp/pptwork/v5/_base_rk.pptx')
GREEN, PINK, NAVY = 'C9F0D6', 'FDE2DE', '0B5E6E'
US = '/tmp/pptwork/v5/us/'


def find_title(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.top is not None and 0.6*E <= sh.top <= 1.3*E and sh.width > 9*E and sh.text_frame.text.strip():
            return sh


def prep(s, title, resume):
    t = find_title(s)
    for sh in list(s.shapes):
        keep = (sh._element is t._element) or sh.top >= 10.3*E or (sh.left >= 18.0*E and sh.top < 1.0*E) or sh.top < 0.45*E
        if not keep:
            sh._element.getparent().remove(sh._element)
    for sh in list(s.shapes):
        if sh.top > 10.3*E and 16.5*E < sh.left < 17.9*E:
            sh._element.getparent().remove(sh._element)
    for sh in [x for x in s.shapes if x.has_text_frame and x.left > 17.9*E and x.top > 10.4*E and x.text_frame.text.strip().isdigit()]:
        set_text(sh, '')
    ps = t.text_frame.paragraphs
    ps[0].runs[0].text = title
    ps[1].runs[0].text = resume


def intinya(s, text, y=9.05, h=0.85):
    return box(s, 0.9, y, 18.2, h, fill=YEL_L, line=YEL, anchor=MSO_ANCHOR.MIDDLE,
               paras=[[('Intinya  ', 14, True, TEAL), (text, 14, False, DARK)]])


def head(s, x, y, w, text, sub=None):
    box(s, x, y, w, 0.5, fill=TEAL, line=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.15,
        paras=[[(text, 13.5, True, WHITE)] + ([('   ' + sub, 11, False, 'D7EEF2')] if sub else [])])


def qa(s, x, y, w, h, rows, size=12.5):
    """rows: list of (label, text[, color])"""
    paras = []
    for r in rows:
        paras.append([(r[0] + '  ', size, True, r[2] if len(r) > 2 else TEAL), (r[1], size, False, DARK)])
    return box(s, x, y, w, h, fill=LIGHT, line=LINE, paras=paras)


def chip(s, x, y, w, h, text, fill=LIGHT, line=LINE, size=10.5, color=DARK, bold=False):
    return box(s, x, y, w, h, fill=fill, line=line, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, margin=0.06,
               paras=[(text, size, bold, color)])


def outline(s, x, y, w, h, color=RED):
    r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.background(); r.line.color.rgb = rgb(color); r.line.width = Pt(2.5); r.shadow.inherit = False
    return r


s = p.slides[0]
prep(s, 'Rangkuman Kelayakan Model', 'Model layak dipakai untuk simulasi skenario, dengan batas pemakaian yang jelas')
CARDS = [('1', 'Uji struktur', 'LOLOS', TEAL, ['Satuan seluruh persamaan konsisten', 'Jumlah tiap stok cocok, tidak ada selisih', '6 loop bekerja sesuai rancangan', '17 uji kondisi ekstrem: 0 gagal']),
         ('2', 'Uji perilaku', 'CUKUP BAIK', '2A9DB0', ['Diuji terpisah: semua MAPE < 20%', 'Dijalankan bersama: tidak ada MAPE buruk', 'Selisih berasal dari satu sumber: pertumbuhan wisatawan']),
         ('3', 'Kalibrasi', 'NILAI DARI DATA DIPAKAI', TEAL, ['3 parameter dikalibrasi, 0 diganti', 'Selisih wisatawan karena lonjakan setelah pandemi, di luar cakupan model'])]
w, g = 5.9, 0.25
for k, (n, title, status, col, pts) in enumerate(CARDS):
    x = 0.9 + k * (w + g)
    box(s, x, 1.95, w, 3.25, fill=LIGHT, line=LINE, radius=0.05)
    box(s, x, 1.95, w, 1.0, fill=col, line=None, radius=0.1, anchor=MSO_ANCHOR.MIDDLE, margin=0.2,
        paras=[[(n + '  ' + title, 12, True, 'D7EEF2')], [(status, 17, True, WHITE)]])
    tb(s, x + 0.25, 3.12, w - 0.45, 2.0, [[('✓  ', 13.5, True, '1E7A43'), (t_, 13.5, False, DARK)] for t_ in pts])
    if k < 2:
        arrow(s, x + w + 0.02, 3.55, g - 0.04, 0.3)
head(s, 0.9, 5.45, 18.2, 'Batas pemakaian', 'hal yang perlu diingat saat membaca hasil model')
B = [('Bukan alat ramal angka pasti', 'Hasil dibaca sebagai arah dan perbandingan antarskenario, bukan angka tepat untuk tahun tertentu.'),
     ('Belum diuji dengan data baru', 'Kecocokan setelah kalibrasi dihitung dengan data yang sama, jadi bukan validasi dengan data terpisah.'),
     ('Tanpa kejadian mendadak', 'Pandemi, bencana, dan guncangan sejenis tidak dimodelkan.'),
     ('LPE dan LPD tidak bisa dipisahkan', 'Data hanya bisa menentukan selisih laju pertumbuhan dan penurunan wisatawan.')]
for k, (a, b) in enumerate(B):
    x = 0.9 + (k % 2) * 9.2; y = 6.05 + (k // 2) * 1.15
    box(s, x, y, 9.0, 1.02, fill=YEL_L, line=YEL, anchor=MSO_ANCHOR.MIDDLE, margin=0.18,
        paras=[[(a, 12.5, True, '9B2C1F')], [(b, 11.5, False, DARK)]])
box(s, 0.9, 8.55, 18.2, 1.25, fill=TEAL, line=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.3, paras=[
    [('Keputusan  ', 14, True, YEL), ('Model layak dipakai untuk membandingkan arah dan besar perbedaan antarskenario kebijakan 2025–2050.', 14, True, WHITE)]])
source(s, 'Sumber: Buku Subbab 4.4–4.6.', y=10.12)
s.notes_slide.notes_text_frame.text = (
    'Sebagai rangkuman, ada tiga bukti kelayakan model. Pertama, uji struktur lolos seluruhnya: satuan konsisten, stok cocok, enam loop bekerja, dan 17 uji ekstrem tidak ada yang gagal. '
    'Kedua, uji perilaku cukup baik: saat diuji terpisah semua MAPE di bawah 20 persen, saat dijalankan bersama tidak ada MAPE buruk, dan selisihnya berasal dari satu sumber, yaitu pertumbuhan wisatawan. '
    'Ketiga, kalibrasi tidak mengganti satu pun nilai parameter. Ada empat batas pemakaian: model bukan alat ramal angka pasti, belum diuji dengan data baru, tidak memuat kejadian mendadak, '
    'dan laju pertumbuhan serta penurunan wisatawan tidak bisa dipisahkan oleh data. Jadi model layak dipakai untuk membandingkan arah dan besar perbedaan antarskenario kebijakan 2025 sampai 2050.')
p.save('/tmp/pptwork/v5/rangkuman_1slide.pptx')
print('ok')
