exec(open('/tmp/pptwork/helpers_v4.py').read())
from pptx import Presentation
from pptx.oxml.ns import qn
p = Presentation('/tmp/pptwork/v5/s7.pptx')
S = list(p.slides)
s = S[17]  # slide baru (posisi 18)
def find_title(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.top is not None and 0.6*E <= sh.top <= 1.3*E and sh.width > 9*E and sh.text_frame.text.strip():
            return sh
t = find_title(s)
for sh in list(s.shapes):
    keep = (sh._element is t._element) or sh.top >= 10.3*E or (sh.left >= 18.0*E and sh.top < 1.0*E) or sh.top < 0.45*E
    if not keep:
        sh._element.getparent().remove(sh._element)
ps = t.text_frame.paragraphs
ps[0].runs[0].text = 'Kandidat Variabel Struktur Konseptual Pariwisata'
for r in ps[0].runs[1:]: r._r.getparent().remove(r._r)
ps[1].runs[0].text = 'Kandidat variabel dari enam subsistem beserta dasar konteks, peran, dan literatur pendukungnya'
for r in ps[1].runs[1:]: r._r.getparent().remove(r._r)
rows = [['Subsistem', 'Kandidat variabel', 'Dasar konteks DIY', 'Peran dalam sistem', 'Literatur pendukung'],
 ['Permintaan dan daya tarik destinasi', 'Jumlah Wisatawan; Laju Kedatangan Wisatawan; Daya Tarik Destinasi Wisata; Kepadatan Wisatawan', 'RIPPARDA DIY; Renja Dispar DIY; UU 10/2009; IPKN 2024', 'Menggambarkan permintaan, kemampuan destinasi menarik kunjungan, dan kepadatan sebagai pembatas perkembangan destinasi', 'Mai & Smith (2015, 2018); Gazoni & da Silva (2022); Hu & Ritchie (1993); Saveriades (2000)'],
 ['Atraksi/ODTW', 'Jumlah Objek Daya Tarik Wisata', 'UU 10/2009; RIPPARDA DIY; Renja Dispar DIY', 'Ketersediaan atraksi sebagai pembentuk daya tarik dan pendorong kunjungan', 'Leiper (1990); Hu & Ritchie (1993); Gazoni & da Silva (2022); Fatina dkk. (2023)'],
 ['Ekonomi dan investasi', 'Pengeluaran Wisatawan; PDRB Sektor Pariwisata; Investasi Sektor Pariwisata', 'RIPPARDA DIY; Renja Dispar DIY; Renstra Kemenpar 2025–2029; IPKN 2024', 'Jalur dari kunjungan ke konsumsi wisatawan, nilai tambah ekonomi, dan pengembangan kapasitas lewat investasi', 'Frechtling (2010); Munjal (2013); Mai & Smith (2018); Fatina dkk. (2023)'],
 ['Akomodasi', 'Jumlah Hotel dan Akomodasi; Rasio Permintaan terhadap Kapasitas Kamar; Laju Konstruksi Hotel dan Akomodasi', 'UU 10/2009; RIPPARDA DIY; IPKN 2024', 'Kapasitas akomodasi dan penyesuaian pasokan terhadap permintaan kamar', 'Wheaton & Rossoff (1998); Fatina dkk. (2023)'],
 ['Lahan dan daya dukung', 'Rasio Daya Dukung Lahan', 'RIPPARDA DIY; IPKN; dokumen Sustainable Tourism Development', 'Keterbatasan fisik destinasi dan tekanan pembangunan terhadap lahan sebagai pembatas pertumbuhan', 'Saveriades (2000); Mai & Smith (2015, 2018); Fatina dkk. (2023)'],
 ['Tenaga kerja pariwisata', 'Tenaga Kerja Pariwisata; Selisih Tenaga Kerja Dibutuhkan; Laju Penyerapan Tenaga Kerja', 'RIPPARDA DIY; Renja Dispar DIY; Renstra Kemenpar 2025–2029', 'Dampak aktivitas ekonomi pariwisata terhadap kebutuhan dan penyesuaian tenaga kerja', 'Munjal (2013); Fatina dkk. (2023); Sánchez López (2023)']]
table(s, 0.9, 2.05, 18.2, [2.6, 4.0, 3.5, 4.6, 3.5], rows, size=11.5, rowh=0.98, bold_first_col=True)
box(s, 0.9, 9.0, 18.2, 0.95, fill=YEL_L, line=YEL, paras=[[('Catatan  ', 13, True, TEAL),
    ('15 kandidat ini belum otomatis masuk CLD; semuanya dievaluasi dulu (relevansi, definisi, mekanisme sebab-akibat, peran dalam feedback loop, dukungan teori). Rujukan menjadi dasar keberadaan variabel, bukan dasar nilai parameter.', 12.5, False, DARK)]], anchor=MSO_ANCHOR.MIDDLE)
source(s, 'Sumber: Buku Subbab 3.7.2.1, Tabel 10.', y=10.12)
s.notes_slide.notes_text_frame.text = ('Ini kandidat variabel struktur konseptual. Ada 15 kandidat dari enam subsistem: permintaan dan daya tarik, atraksi, ekonomi dan investasi, '
    'akomodasi, lahan dan daya dukung, serta tenaga kerja. Setiap kandidat punya dasar konteks DIY dari dokumen seperti RIPPARDA, Renja Dinas Pariwisata, '
    'dan IPKN, serta literatur pendukung. Kandidat ini belum otomatis masuk CLD; semuanya dievaluasi dulu, dan hasilnya ada di slide berikutnya.')
# nomor halaman
def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6: yield from walk(sh.shapes)
        else: yield sh
for k, sl in enumerate(p.slides, 1):
    for sh in walk(sl.shapes):
        if sh.has_text_frame and sh.left is not None and sh.left > 17.9*E and sh.top > 10.4*E and sh.text_frame.text.strip().isdigit():
            set_text(sh, str(k))
p.save('/tmp/pptwork/v5/out7.pptx')
# versi satu slide
import copy
print('ok', len(p.slides))
