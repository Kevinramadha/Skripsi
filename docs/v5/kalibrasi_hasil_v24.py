import sys; sys.path.insert(0, '/tmp/pptwork/v5')
from single import make_subset
make_subset('/tmp/pptwork/v5/out7.pptx', [51], '/tmp/pptwork/v5/_base_kh.pptx')
exec(open('/tmp/pptwork/helpers_v4.py').read())
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
p = Presentation('/tmp/pptwork/v5/_base_kh.pptx')
GREEN, PINK, NAVY = 'C9F0D6', 'FDE2DE', '0B5E6E'
US = '/tmp/pptwork/v5/us/'


def find_title(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.top is not None and 0.6*E <= sh.top <= 1.3*E and sh.width > 9*E and sh.text_frame.text.strip():
            return sh


def prep(s, title, resume):
    t = find_title(s)
    for sh in list(s.shapes):
        keep = (sh._element is t._element) or sh.top >= 10.3*E or (sh.left >= 18.0*E and sh.top < 1.0*E) or sh.top < 0.45*E
        if not keep:
            sh._element.getparent().remove(sh._element)
    for sh in list(s.shapes):
        if sh.top > 10.3*E and 16.5*E < sh.left < 17.9*E:
            sh._element.getparent().remove(sh._element)
    for sh in [x for x in s.shapes if x.has_text_frame and x.left > 17.9*E and x.top > 10.4*E and x.text_frame.text.strip().isdigit()]:
        set_text(sh, '')
    ps = t.text_frame.paragraphs
    ps[0].runs[0].text = title
    ps[1].runs[0].text = resume


def intinya(s, text, y=9.05, h=0.85):
    return box(s, 0.9, y, 18.2, h, fill=YEL_L, line=YEL, anchor=MSO_ANCHOR.MIDDLE,
               paras=[[('Intinya  ', 14, True, TEAL), (text, 14, False, DARK)]])


def head(s, x, y, w, text, sub=None):
    box(s, x, y, w, 0.5, fill=TEAL, line=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.15,
        paras=[[(text, 13.5, True, WHITE)] + ([('   ' + sub, 11, False, 'D7EEF2')] if sub else [])])


def qa(s, x, y, w, h, rows, size=12.5):
    """rows: list of (label, text[, color])"""
    paras = []
    for r in rows:
        paras.append([(r[0] + '  ', size, True, r[2] if len(r) > 2 else TEAL), (r[1], size, False, DARK)])
    return box(s, x, y, w, h, fill=LIGHT, line=LINE, paras=paras)


def chip(s, x, y, w, h, text, fill=LIGHT, line=LINE, size=10.5, color=DARK, bold=False):
    return box(s, x, y, w, h, fill=fill, line=line, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, margin=0.06,
               paras=[(text, size, bold, color)])


def outline(s, x, y, w, h, color=RED):
    r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.background(); r.line.color.rgb = rgb(color); r.line.width = Pt(2.5); r.shadow.inherit = False
    return r


s = p.slides[0]
prep(s, 'Kalibrasi: Hasil Tiap Tahap',
     'Lima tahap dijalankan; dari tiga parameter yang dikalibrasi, tidak ada yang diganti')
C = [
 dict(tag='TAHAP 0', title='Membetulkan cara hitung', hfill=GREY,
      par='LPE dan LPD (laju pertumbuhan dan penurunan wisatawan)', did='Rumus dicek ulang terhadap persamaan model',
      big=None, rows=[('LPE', '0,26864 → 0,27726'), ('LPD', '0,20148 → 0,207945')], blab='sebelum → sesudah dibetulkan',
      why='Rumus awal menganggap daya tarik 2024 = 1, padahal 0,99223. Dampaknya kecil: wisatawan 2050 naik sekitar 1,6%.',
      badge='Dibetulkan, bukan kalibrasi', bfill=YEL, bcol=DARK),
 dict(tag='TAHAP 1', title='Subsistem akomodasi', hfill=TEAL,
      par='Batas TPK pemicu pembangunan hotel dan laju demolisi', did='775 pasangan nilai dicoba · uji P1',
      big='11,64%', bcolr=TEAL, blab='lebih cocok pada nilai terbaik', sub='→ tinggal 5,74% bila data 2019 dibuang',
      why='Nilai terbaik jatuh di batas atas yang dicoba, dan keunggulannya hanya muncul karena satu tahun data yang janggal.',
      badge='Nilai dari data tetap dipakai\n(0,275 dan 0,05)', bfill=TEAL, bcol=WHITE),
 dict(tag='TAHAP 2', title='Subsistem lahan', hfill=TEAL,
      par='Laju konversi dasar (lahan menjadi lahan terbangun)', did='Nilai 0,0034–0,0851 dicoba · uji P4',
      big='14,76%', bcolr=TEAL, blab='lebih cocok pada nilai 0,0639', sub='selisih rata-rata (U1): 0,4250 → 0,0755',
      why='Letaknya di tengah, bukan di tepi, sehingga memenuhi syarat (R4). Namun masih harus diuji pada model penuh.',
      badge='Diterima sementara\n(0,0639 → diuji di Tahap 4)', bfill=YEL, bcol=DARK),
 dict(tag='TAHAP 3', title='Subsistem wisatawan', hfill=TEAL,
      par='LPE dan LPD (hanya diperiksa, tidak dikalibrasi)', did='Berapa LPE yang diminta data?',
      big='0,41–0,49', bcolr=TEAL, blab='LPE yang diminta data', sub='nilai dari data: 0,27726',
      why='Banyak pasangan LPE–LPD sama cocoknya, jadi data hanya bisa menentukan selisih keduanya. Penyebab: lonjakan wisatawan setelah pandemi (puncak 24,86% pada 2024).',
      badge='Nilai dari data tetap dipakai\n(0,27726 dan 0,207945)', bfill=TEAL, bcol=WHITE),
 dict(tag='TAHAP 4', title='Uji pada model penuh', hfill=RED,
      par='Laju konversi 0,0639 dari Tahap 2', did='Dicoba pada 13 variabel yang dinilai',
      big='32,4%', bcolr=RED, blab='lahan terbangun (uji penuh) jadi kurang cocok', sub='batas toleransi 5% · 12 variabel lain hampir tetap',
      why='Uji penuh mulai 2019 saat lahan sudah tinggi, sehingga laju yang lebih cepat membuat simulasi melampaui data.',
      badge='Ditolak\n(kembali ke 0,0436)', bfill=RED, bcol=WHITE)]
w, g = 3.5, 0.175
for k, c in enumerate(C):
    x = 0.9 + k * (w + g)
    box(s, x, 1.95, w, 6.5, fill=WHITE, line=LINE, radius=0.05)
    box(s, x, 1.95, w, 0.85, fill=c['hfill'], line=None, radius=0.12, anchor=MSO_ANCHOR.MIDDLE, margin=0.15,
        paras=[[(c['tag'], 9.5, True, YEL if c['hfill'] != YEL else DARK)], [(c['title'], 12.5, True, WHITE)]])
    tb(s, x + 0.15, 2.88, w - 0.3, 0.6, [[(c['par'], 9.5, True, DARK, True)]])
    box(s, x + 0.15, 3.5, w - 0.3, 0.42, fill=LIGHT, line=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.08,
        paras=[[('Dicoba: ', 9, True, TEAL), (c['did'], 9, False, DARK)]])
    if c['big']:
        tb(s, x + 0.15, 4.0, w - 0.3, 0.75, [[(c['big'], 28, True, c['bcolr'])]], align=PP_ALIGN.CENTER)
        tb(s, x + 0.15, 4.78, w - 0.3, 0.75, [[(c['blab'], 10, True, DARK)], [(c['sub'], 9.5, False, GREY)]], align=PP_ALIGN.CENTER)
    else:
        for j, (lab, val) in enumerate(c['rows']):
            box(s, x + 0.15, 4.05 + j * 0.55, w - 0.3, 0.47, fill=LIGHT, line=None, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, margin=0.05,
                paras=[[(lab + '  ', 10.5, True, TEAL), (val, 11.5, True, DARK)]])
        tb(s, x + 0.15, 5.18, w - 0.3, 0.35, [[(c['blab'], 9.5, False, GREY)]], align=PP_ALIGN.CENTER)
    ln = s.shapes.add_connector(1, Inches(x + 0.3), Inches(5.6), Inches(x + w - 0.3), Inches(5.6))
    ln.line.color.rgb = rgb(LINE); ln.line.width = Pt(1)
    tb(s, x + 0.15, 5.68, w - 0.3, 1.75, [[('Kenapa? ', 10.5, True, TEAL), (c['why'], 10.5, False, DARK)]])
    box(s, x + 0.12, 7.52, w - 0.24, 0.8, fill=c['bfill'], line=None, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, margin=0.06,
        paras=[[(ln_, 10.5 if i == 0 else 9.5, i == 0, c['bcol'])] for i, ln_ in enumerate(c['badge'].split('\n'))])
    if k < 4:
        arrow(s, x + w - 0.02, 2.22, g + 0.04, 0.3)
box(s, 0.9, 8.65, 18.2, 1.3, fill=TEAL, line=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.25, paras=[
    [('Hasil akhir: dari 3 parameter yang dikalibrasi (batas TPK, laju demolisi, laju konversi), tidak ada yang diganti', 13.5, True, YEL)],
    [('Nilai yang dihitung dari data sudah layak dipakai. Selisih pada jumlah wisatawan bukan karena nilai parameter yang salah, tetapi karena lonjakan wisatawan setelah pandemi yang berada di luar cakupan model.', 11, False, WHITE)]])
source(s, 'Sumber: Buku Subbab 4.6 (Tabel 48–55). R4 dan M1 = aturan keputusan dan pemicu penolakan pada slide sebelumnya (Tabel 19–20).', y=10.12)
s.notes_slide.notes_text_frame.text = (
    'Kalibrasi dijalankan dalam lima tahap. Tahap 0 bukan kalibrasi: saya membetulkan rumus LPE dan LPD karena rumus awal menganggap daya tarik 2024 bernilai 1, padahal 0,99223. Dampaknya kecil, dan hasil uji ekstrem tidak berubah. '
    'Tahap 1 pada akomodasi: pasangan nilai terbaik memang 11,64 persen lebih cocok, tetapi letaknya di batas atas, dan tanpa data 2019 keunggulannya tinggal 5,74 persen. Jadi keunggulannya hanya karena satu tahun data yang janggal, dan nilai dari data tetap dipakai. '
    'Tahap 2 pada lahan: laju konversi 0,0639 lebih cocok 14,76 persen dan memenuhi syarat, sehingga diterima sementara. '
    'Tahap 3 pada wisatawan hanya pemeriksaan: data meminta LPE yang lebih tinggi, tetapi data hanya bisa menentukan selisih LPE dan LPD, dan penyebab selisihnya adalah lonjakan wisatawan setelah pandemi. '
    'Tahap 4 menguji hasil Tahap 2 pada model penuh: lahan terbangun justru 32,4 persen kurang cocok, karena uji penuh dimulai dari lahan yang sudah tinggi. Hasilnya ditolak. '
    'Jadi dari tiga parameter yang dikalibrasi, tidak ada yang diganti.')
p.save('/tmp/pptwork/v5/kalibrasi_hasil_1slide.pptx')
print('ok')
