import sys; sys.path.insert(0, '/tmp/pptwork/v5')
from single import make_single
make_single('/tmp/pptwork/v5/out7.pptx', 29, '/tmp/pptwork/v5/_base_odtw.pptx')
exec(open('/tmp/pptwork/helpers_v4.py').read())
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
p = Presentation('/tmp/pptwork/v5/_base_odtw.pptx')
s = p.slides[0]
def find_title(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.top is not None and 0.6*E <= sh.top <= 1.3*E and sh.width > 9*E and sh.text_frame.text.strip():
            return sh
t = find_title(s)
for sh in list(s.shapes):
    keep = (sh._element is t._element) or sh.top >= 10.3*E or (sh.left >= 18.0*E and sh.top < 1.0*E) or sh.top < 0.45*E
    if not keep:
        sh._element.getparent().remove(sh._element)
for sh in list(s.shapes):   # buang penanda W/S
    if sh.top > 10.3*E and sh.left > 16.5*E and sh.left < 17.9*E:
        sh._element.getparent().remove(sh._element)
ps = t.text_frame.paragraphs
ps[0].runs[0].text = 'Estimasi Jumlah ODTW — Hasil Estimasi dan Data Aktual'
ps[1].runs[0].text = 'Data ODTW 2015–2025: tahun berdata BPS memakai data aktual, tahun kosong diisi estimasi dari NTL'

NTL = {2015: 1.261585, 2016: 1.054415, 2017: 1.438316, 2018: 1.668258, 2019: 2.004828, 2020: 1.827902,
       2021: 1.540893, 2022: 1.483312, 2023: 2.512341, 2024: 2.457036, 2025: 2.287813}
BPS = {2018: 177, 2019: 189, 2020: 180, 2021: 170, 2022: 183, 2023: 201, 2024: 218}
PI = {2015: '138,5 – 191,9', 2016: '128,9 – 187,3', 2017: '146,4 – 196,3', 2025: '176,7 – 224,8'}
EST = {2015: 165.2, 2016: 158.1, 2017: 171.4, 2025: 200.7}
used = {2015: 165, 2016: 158, 2017: 171, 2025: 201}; used.update(BPS)
c = lambda x, d=1: ('{:,.%df}' % d).format(x).replace(',', 'X').replace('.', ',').replace('X', '.')
rows = [['Tahun', 'NTL', 'ODTW aktual (BPS)', 'Hasil regresi NTL', 'Selang prediksi 95%', 'Dipakai di model', 'Status']]
for y in range(2015, 2026):
    fit = EST.get(y, 121.61 + 34.59 * NTL[y])
    rows.append([str(y), c(NTL[y], 3), str(BPS[y]) if y in BPS else '—', c(fit), PI.get(y, '—'), str(used[y]),
                 'Estimasi NTL' if y not in BPS else 'Data BPS'])
hl = {}
for r in range(1, 12):
    if rows[r][6] == 'Estimasi NTL':
        for k in range(7):
            hl[(r, k)] = YEL_L
gf = table(s, 0.9, 2.0, 11.6, [0.9, 1.0, 1.75, 1.75, 2.25, 1.65, 1.6], rows, size=12, rowh=0.52, hl=hl)
tbl = gf.table
for r in range(0, 12):
    for k in range(7):
        para = tbl.cell(r, k).text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER
        if r > 0 and k in (5, 6):
            para.runs[0].font.bold = True
        if r > 0 and rows[r][6] == 'Estimasi NTL' and k == 6:
            para.runs[0].font.color.rgb = rgb('9A6B00')

# grafik (bisa diedit): aktual vs dipakai di model
cd = CategoryChartData()
cd.categories = [str(y) for y in range(2015, 2026)]
cd.add_series('ODTW aktual (BPS)', [BPS.get(y) for y in range(2015, 2026)])
cd.add_series('Estimasi NTL', [EST.get(y) for y in range(2015, 2026)])
gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(12.8), Inches(2.0), Inches(6.3), Inches(5.6), cd)
ch = gf.chart
ch.has_title = False
ch.font.size = Pt(11); ch.font.name = F; ch.font.color.rgb = rgb(DARK)
ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.TOP; ch.legend.include_in_layout = False
for ser, col in zip(ch.plots[0].series, [TEAL, 'E0A800']):
    ser.format.line.color.rgb = rgb(col); ser.format.line.width = Pt(2.75); ser.smooth = False
    ser.marker.size = 9; ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = rgb(col)
    ser.marker.format.line.color.rgb = rgb(col)
ch.plots[0].series[1].format.line.dash_style = 4  # garis putus-putus
ch.plots[0].has_data_labels = True
dl = ch.plots[0].data_labels; dl.number_format = '0'; dl.number_format_is_linked = False
dl.font.size = Pt(10); dl.position = XL_LABEL_POSITION.ABOVE
ch.value_axis.minimum_scale = 140; ch.value_axis.maximum_scale = 230
ch.value_axis.has_major_gridlines = True
ch.value_axis.major_gridlines.format.line.color.rgb = rgb('E3EEF1')
ch.value_axis.tick_labels.font.size = Pt(10); ch.category_axis.tick_labels.font.size = Pt(10)

box(s, 12.8, 7.75, 6.3, 1.45, fill=LIGHT, paras=[('Cara membaca', 13, True, TEAL),
    ('Baris kuning = tahun tanpa data BPS (2015–2017 dan 2025), diisi dengan ODTW = 121,61 + 34,59 × NTL. Nilai 2025 (201 unit) menjadi stok awal model.', 11.5, False, DARK)])
box(s, 0.9, 8.4, 11.6, 0.8, fill=YEL_L, line=YEL, paras=[[('Catatan  ', 12.5, True, TEAL),
    ('Kolom "Hasil regresi NTL" pada tahun berdata BPS hanya pembanding; model tetap memakai data aktual BPS.', 12, False, DARK)]], anchor=MSO_ANCHOR.MIDDLE)
source(s, 'Sumber: BPS DIY (ODTW 2018–2024); VIIRS DNB via Google Earth Engine (NTL); Buku Subbab 4.2.1, Tabel 30; repository ODTW_DIY_2015_2025_Final.csv.', y=10.12)
s.notes_slide.notes_text_frame.text = ('Tabel ini menunjukkan data ODTW 2015 sampai 2025. Tahun 2018 sampai 2024 memakai data aktual BPS. '
    'Tahun 2015, 2016, 2017, dan 2025 tidak ada data BPS, sehingga diisi dengan estimasi dari cahaya malam memakai regresi ODTW = 121,61 + 34,59 × NTL, '
    'lengkap dengan selang prediksi 95 persen. Nilai 2025 sebesar 201 unit menjadi stok awal model. '
    'Kolom hasil regresi pada tahun berdata BPS hanya pembanding untuk melihat seberapa dekat model dengan data aktual.')
p.save('/tmp/pptwork/v5/odtw_1slide.pptx')
print('ok')
