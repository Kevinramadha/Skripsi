exec(open('/tmp/pptwork/helpers_v4.py').read())
from pptx import Presentation
p = Presentation('/tmp/pptwork/v5/kandidat_1slide.pptx')
s = p.slides[0]
def find_title(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.top is not None and 0.6*E <= sh.top <= 1.3*E and sh.width > 9*E and sh.text_frame.text.strip():
            return sh
t = find_title(s)
for sh in list(s.shapes):
    keep = (sh._element is t._element) or sh.top >= 10.3*E or (sh.left >= 18.0*E and sh.top < 1.0*E) or sh.top < 0.45*E
    if not keep:
        sh._element.getparent().remove(sh._element)
ps = t.text_frame.paragraphs
ps[0].runs[0].text = 'Matriks Hubungan Kausal Penyusun CLD'
ps[1].runs[0].text = '18 hubungan sebab-akibat: arah hubungan, cara kerjanya dalam model, dan dasarnya'
rows = [['Loop', 'Hubungan kausal', '+/−', 'Cara kerja dalam model', 'Dasar / rujukan'],
 ['R1', 'Jumlah Wisatawan → Laju Kedatangan', '+', 'Makin banyak wisatawan, makin besar basis yang tumbuh, sehingga kedatangan bertambah', 'Struktur pertumbuhan; Sterman (2000); Mai & Smith (2018)'],
 ['R1, R2, B1, B2', 'Daya Tarik Destinasi → Laju Kedatangan', '+', 'Daya tarik yang lebih tinggi membuat destinasi lebih mampu menarik kunjungan', 'Hu & Ritchie (1993); Leiper (1990)'],
 ['R2', 'Jumlah Wisatawan → Pengeluaran Wisatawan', '+', 'Dengan pengeluaran per kunjungan tetap, kunjungan lebih banyak berarti total pengeluaran lebih besar', 'Definisi penjumlahan; Frechtling (2010); Dispar DIY (2024)'],
 ['R2', 'Pengeluaran Wisatawan → PDRB Sektor Pariwisata', '+', 'Pengeluaran wisatawan menjadi nilai tambah melalui rasio nilai tambah', 'Frechtling (2010); Munjal (2013)'],
 ['R2', 'PDRB Sektor Pariwisata → Investasi', '+', 'Investasi dihitung sebagai proporsi PDRB dari data historis (formulasi data, bukan klaim umum)', 'Formulasi berbasis data; RIPPARDA DIY (2012)'],
 ['R2', 'Investasi → Laju Pembangunan ODTW', '+', 'Investasi lebih besar menambah pembangunan ODTW', 'Fatina dkk. (2023); RIPPARDA DIY (2012)'],
 ['R2', 'Jumlah ODTW → Daya Tarik Destinasi', '+', 'Ketersediaan atraksi adalah salah satu pembentuk daya tarik', 'Leiper (1990); Hu & Ritchie (1993); Gazoni & da Silva (2022)'],
 ['B1', 'Jumlah Wisatawan → Kepadatan Wisatawan', '+', 'Luas wilayah tetap, sehingga kunjungan lebih banyak berarti lebih padat', 'Definisi model; IPKN 2024'],
 ['B1', 'Kepadatan Wisatawan → Daya Tarik Destinasi', '−', 'Destinasi yang terlalu padat menurunkan kenyamanan berwisata', 'Saveriades (2000); Mai & Smith (2018)'],
 ['B2', 'Konstruksi Hotel & Pembangunan ODTW → Konversi Lahan Pariwisata', '+', 'Fasilitas baru butuh lahan, sehingga konversi lahan pariwisata bertambah', 'Formulasi model; Mai & Smith (2018); Fatina dkk. (2023)'],
 ['B2', 'Konversi Lahan Pariwisata → Lahan Terbangun', '+', 'Konversi lahan adalah aliran yang menambah lahan terbangun', 'Persamaan stok-aliran'],
 ['B2', 'Lahan Terbangun → Rasio Daya Dukung Lahan', '−', 'Makin banyak lahan terbangun, makin kecil sisa lahan', 'Definisi model; IPKN 2024'],
 ['B2', 'Rasio Daya Dukung Lahan → Daya Tarik Destinasi', '+', 'Lingkungan yang lebih terjaga menaikkan komponen lingkungan pada daya tarik', 'Saveriades (2000); Mai & Smith (2018)'],
 ['B3', 'Jumlah Hotel → Rasio Permintaan thd Kapasitas Kamar', '−', 'Dengan permintaan yang sama, kamar lebih banyak berarti tekanan permintaan lebih kecil', 'Definisi kapasitas; Wheaton & Rossoff (1998)'],
 ['B3', 'Rasio Permintaan thd Kapasitas Kamar → Laju Konstruksi Hotel', '+', 'Jika permintaan kamar di atas batas, pembangunan akomodasi baru terdorong', 'Wheaton & Rossoff (1998); formulasi model'],
 ['B4', 'PDRB Sektor Pariwisata → Tenaga Kerja Dibutuhkan', '+', 'Aktivitas ekonomi pariwisata lebih besar membutuhkan tenaga kerja lebih banyak', 'Munjal (2013); Sánchez López (2023)'],
 ['B4', 'Tenaga Kerja Pariwisata → Selisih TK Dibutuhkan', '−', 'Tenaga kerja yang sudah ada memperkecil kekurangan tenaga kerja', 'Definisi penyesuaian stok'],
 ['B4', 'Selisih TK Dibutuhkan → Laju Penyerapan TK', '+', 'Kekurangan tenaga kerja yang lebih besar mendorong penyerapan tenaga kerja', 'Struktur goal-seeking; Sterman (2000)']]
hl = {}
for r in range(1, len(rows)):
    hl[(r, 2)] = 'FDE2DE' if rows[r][2] == '−' else 'D7F2F6'
table(s, 0.9, 1.95, 18.2, [1.05, 4.8, 0.5, 7.6, 4.25], rows, size=9.5, rowh=0.39, hl=hl, bold_first_col=True)
for r in range(1, len(rows)):
    c = s.shapes[-1].table.cell(r, 2)
    c.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    c.text_frame.paragraphs[0].runs[0].font.bold = True
source(s, 'Sumber: Buku Subbab 4.1.3, Tabel 25. Hubungan Tenaga Kerja Pariwisata → pembangunan ODTW tidak dimasukkan karena dasar sebab-akibatnya tidak cukup.', y=10.27)
s.notes_slide.notes_text_frame.text = ('Ini matriks lengkap 18 hubungan kausal penyusun CLD. Setiap hubungan diberi arah, plus atau minus, '
    'cara kerjanya di dalam model, dan dasarnya. Dasarnya dibedakan: sebagian didukung literatur, misalnya daya tarik terhadap kedatangan dan kepadatan terhadap daya tarik; '
    'sebagian lagi berupa definisi atau persamaan dalam model, misalnya lahan terbangun terhadap rasio daya dukung lahan. '
    'Hubungan tenaga kerja ke pembangunan ODTW tidak dimasukkan karena dasar sebab-akibatnya tidak cukup.')
for sh in list(s.shapes):
    if sh.has_text_frame and sh.text_frame.text.strip() in ('W', 'S') and sh.top > 10.3*E:
        sh._element.getparent().remove(sh._element)
    elif sh.shape_type == 6 and sh.top > 10.3*E and sh.left > 16.5*E and ''.join(x.text_frame.text for x in sh.shapes if x.has_text_frame).strip() in ('W', 'S'):
        sh._element.getparent().remove(sh._element)
p.save('/tmp/pptwork/v5/matriks_1slide.pptx')
print('ok')
