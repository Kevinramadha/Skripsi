import sys; sys.path.insert(0, '/tmp/pptwork/v5')
from single import make_subset
make_subset('/tmp/pptwork/v5/out7.pptx', [50, 51], '/tmp/pptwork/v5/_base_kal.pptx')
exec(open('/tmp/pptwork/helpers_v4.py').read())
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
p = Presentation('/tmp/pptwork/v5/_base_kal.pptx')
GREEN, PINK, NAVY = 'C9F0D6', 'FDE2DE', '0B5E6E'
US = '/tmp/pptwork/v5/us/'


def find_title(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.top is not None and 0.6*E <= sh.top <= 1.3*E and sh.width > 9*E and sh.text_frame.text.strip():
            return sh


def prep(s, title, resume):
    t = find_title(s)
    for sh in list(s.shapes):
        keep = (sh._element is t._element) or sh.top >= 10.3*E or (sh.left >= 18.0*E and sh.top < 1.0*E) or sh.top < 0.45*E
        if not keep:
            sh._element.getparent().remove(sh._element)
    for sh in list(s.shapes):
        if sh.top > 10.3*E and 16.5*E < sh.left < 17.9*E:
            sh._element.getparent().remove(sh._element)
    for sh in [x for x in s.shapes if x.has_text_frame and x.left > 17.9*E and x.top > 10.4*E and x.text_frame.text.strip().isdigit()]:
        set_text(sh, '')
    ps = t.text_frame.paragraphs
    ps[0].runs[0].text = title
    ps[1].runs[0].text = resume


def intinya(s, text, y=9.05, h=0.85):
    return box(s, 0.9, y, 18.2, h, fill=YEL_L, line=YEL, anchor=MSO_ANCHOR.MIDDLE,
               paras=[[('Intinya  ', 14, True, TEAL), (text, 14, False, DARK)]])


def head(s, x, y, w, text, sub=None):
    box(s, x, y, w, 0.5, fill=TEAL, line=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.15,
        paras=[[(text, 13.5, True, WHITE)] + ([('   ' + sub, 11, False, 'D7EEF2')] if sub else [])])


def qa(s, x, y, w, h, rows, size=12.5):
    """rows: list of (label, text[, color])"""
    paras = []
    for r in rows:
        paras.append([(r[0] + '  ', size, True, r[2] if len(r) > 2 else TEAL), (r[1], size, False, DARK)])
    return box(s, x, y, w, h, fill=LIGHT, line=LINE, paras=paras)


def chip(s, x, y, w, h, text, fill=LIGHT, line=LINE, size=10.5, color=DARK, bold=False):
    return box(s, x, y, w, h, fill=fill, line=line, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, margin=0.06,
               paras=[(text, size, bold, color)])


def outline(s, x, y, w, h, color=RED):
    r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.background(); r.line.color.rgb = rgb(color); r.line.width = Pt(2.5); r.shadow.inherit = False
    return r


S = list(p.slides)

# ============================================================ 1/2 aturan
s = S[0]
prep(s, 'Kalibrasi (1/2): Aturan yang Ditetapkan Sejak Awal',
     'Kalibrasi dipakai untuk menguji nilai parameter, bukan sekadar mencari nilai yang paling cocok dengan data')
box(s, 0.9, 1.95, 8.9, 1.25, fill=YEL_L, line=YEL, anchor=MSO_ANCHOR.MIDDLE, paras=[
    [('Tujuan  ', 12, True, TEAL), ('Memeriksa apakah selisih yang berpola pada uji perilaku bisa dijelaskan oleh nilai parameter, '
      'atau justru berasal dari luar parameter (gangguan data atau hal di luar cakupan model). Kalibrasi diposisikan sebagai alat uji (Oliva, 2003).', 11, False, DARK)]])
head(s, 0.9, 3.35, 8.9, 'Lima ketentuan')
K = [('Parameter dibatasi', 'Hanya parameter yang nilainya belum pasti (asumsi/estimasi). Data resmi, hasil persamaan stok-aliran, dan nilai acuan tahun dasar tidak dikalibrasi.'),
     ('Nilai yang dicoba punya batas yang jelas', 'Batasnya diambil dari rujukan, atau dari selang kepercayaan 95% saat parameter dihitung dari data. Contoh: laju konversi lahan dicoba 0,0034–0,0851.'),
     ('Parameter yang saling terkait ikut disesuaikan', 'Contoh: setiap kali batas TPK diubah, sensitivitas konstruksi dihitung ulang agar model tetap sesuai data tahun dasar.'),
     ('Ukuran kecocokan: NRMSE', 'Selisih simulasi dan data, dibandingkan dengan rata-rata data; makin kecil makin cocok. MAPE sebagai pelengkap.'),
     ('Selisih kecil dianggap setara (batas 5%)', 'Nilai yang NRMSE-nya paling banyak 5% di atas nilai terbaik dianggap sama baiknya. Contoh: terbaik 0,100 → nilai dengan NRMSE ≤ 0,105 setara.')]
for k, (a, b) in enumerate(K):
    y = 3.95 + k * 0.98
    num(s, 0.95, y + 0.14, k + 1, d=0.5, size=13)
    box(s, 1.6, y, 8.2, 0.88, fill=LIGHT, line=LINE, anchor=MSO_ANCHOR.MIDDLE, margin=0.14,
        paras=[[(a, 11, True, TEAL)], [(b, 10, False, DARK)]])

head(s, 10.15, 1.95, 8.95, 'Aturan keputusan', 'dipakai pada setiap tahap')
rows = [['', 'Kondisi', 'Keputusan'],
        ['R1', 'Nilai dari data sudah termasuk yang sama baiknya (≤ 5% di atas nilai terbaik)', 'Tidak diganti, karena nilai baru tidak lebih baik secara berarti'],
        ['R2', 'Nilai yang sama baiknya tersebar dari ujung bawah sampai ujung atas batas', 'Data tidak bisa menunjuk satu nilai → nilai dari data tetap dipakai'],
        ['R3', 'Nilai terbaik jatuh tepat di batas yang dicoba', 'Dicek ulang dengan membuang satu tahun data secara bergantian: apakah nilai terbaik hanya muncul karena satu tahun yang janggal?'],
        ['R4', 'Tidak termasuk ketiga kondisi di atas', 'Nilai terbaik dipakai, tetapi harus lolos uji pada model penuh (Tahap 4)']]
gf = table(s, 10.15, 2.5, 8.95, [0.5, 3.6, 4.85], rows, size=10, rowh=0.78, hl={(4, 2): GREEN})
for r in range(1, 5):
    c = gf.table.cell(r, 0).text_frame.paragraphs[0]; c.alignment = PP_ALIGN.CENTER; c.runs[0].font.bold = True; c.runs[0].font.color.rgb = rgb(TEAL)
box(s, 10.15, 6.6, 8.95, 1.85, fill=WHITE, line=LINE, paras=[
    [('Uji kompatibilitas (Tahap 4): ', 11, True, TEAL), ('nilai hasil kalibrasi dicoba pada model penuh (semua subsistem aktif), dan ditolak bila salah satu terjadi:', 10.5, False, DARK)],
    [('•  kecocokan variabel yang dikalibrasi memburuk lebih dari 5% ', 10.5, False, DARK), ('(M1)', 10.5, True, TEAL)],
    [('•  rata-rata kecocokan 9 variabel memburuk lebih dari 5% ', 10.5, False, DARK), ('(M2)', 10.5, True, TEAL)],
    [('•  model gagal pada uji kondisi ekstrem ', 10.5, False, DARK), ('(S)', 10.5, True, TEAL), ('  •  simulasi jangka panjang tidak stabil ', 10.5, False, DARK), ('(T)', 10.5, True, TEAL)]])
tb(s, 10.15, 8.5, 8.95, 0.4, [[('M1, M2, S, T = kode pada Tabel 20 buku. Batas 5% juga dicek pada 2% dan 10%.', 9.5, False, GREY, True)]])
intinya(s, 'Nilai dari data baru diganti bila nilai baru jelas lebih cocok, tidak hanya karena satu tahun data, dan tetap cocok saat semua subsistem dijalankan bersama.')
source(s, 'Sumber: Buku Subbab 3.7.4 (Tabel 19–20). NRMSE = akar rata-rata kuadrat selisih dibagi rata-rata data.', y=10.12)
s.notes_slide.notes_text_frame.text = (
    'Kalibrasi pada penelitian ini dipakai sebagai alat uji, mengikuti Oliva 2003. Pertanyaannya: apakah selisih yang berpola pada uji perilaku bisa dijelaskan oleh nilai parameter, '
    'atau berasal dari luar parameter. Supaya tidak sekadar mencari nilai yang paling cocok, ada lima ketentuan yang ditetapkan sebelum pencarian: hanya parameter yang belum pasti yang dikalibrasi, '
    'rentangnya punya dasar, parameter berpasangan dihitung ulang, kecocokan diukur dengan NRMSE, dan nilai yang selisihnya paling banyak 5 persen dari nilai terbaik dianggap sama baiknya. '
    'Keputusannya mengikuti empat aturan. Kalau nilai dari data sudah termasuk kelompok sama baiknya, atau data tidak bisa menentukan nilainya, nilai dari data dipertahankan. '
    'Kalau nilai terbaik ada di batas rentang, diperiksa apakah hanya ditarik satu tahun. Nilai baru baru dipakai kalau juga lolos uji kompatibilitas pada model penuh.')

# ============================================================ 2/2 hasil per tahap
s = S[1]
prep(s, 'Kalibrasi (2/2): Hasil Tiap Tahap',
     'Dari tiga parameter yang dikalibrasi, tidak ada yang diganti: nilai yang dihitung dari data dipertahankan')
C = [('TAHAP 0', 'Koreksi penurunan', 'LPE dan LPD (wisatawan)', [
        ('Diperiksa', 'Apakah cara menghitung LPE dan LPD sudah sesuai persamaan model yang berjalan.'),
        ('Temuan', 'Perhitungan awal memakai daya tarik = 1, padahal model memakai daya tarik tahun 2024 (0,99223).'),
        ('Hasil', 'LPE 0,26864 → 0,27726; LPD 0,20148 → 0,207945. Wisatawan 2050 naik sekitar 1,6%; uji ekstrem tetap 15 lolos, 2 catatan, 0 gagal.')],
      'Koreksi, bukan kalibrasi', YEL, DARK),
     ('TAHAP 1', 'Subsistem akomodasi', 'Batas TPK dan laju demolisi', [
        ('Diperiksa', '775 kombinasi (batas TPK 0,20–0,35; demolisi 0,02–0,08) pada uji P1.'),
        ('Temuan', 'Nilai terbaik (0,350; 0,060) memperbaiki 11,64%, tetapi tepat di batas rentang (R3). Tanpa tahun 2019, perbaikan tinggal 5,74%; versi dengan COVID memburuk dan P1b hanya 4,89%.'),
        ('Keputusan', 'Perbaikan ditarik satu tahun anomali.')],
      'Tetap 0,275 dan 0,05', TEAL, WHITE),
     ('TAHAP 2', 'Subsistem lahan', 'Laju konversi dasar', [
        ('Diperiksa', 'Rentang selang kepercayaan 95% (0,0034–0,0851) pada uji P4.'),
        ('Temuan', 'Nilai terbaik 0,0639 memperbaiki 14,76%, berada di dalam rentang, dan bias (U1) turun dari 0,4250 ke 0,0755.'),
        ('Keputusan', 'Memenuhi R4, tetapi masih harus lolos Tahap 4.')],
      'Sementara diadopsi', YEL, DARK),
     ('TAHAP 3', 'Subsistem wisatawan', 'LPE dan LPD (diagnostik)', [
        ('Diperiksa', 'Seberapa besar LPE yang dituntut data. Tidak dicari nilai optimum karena LPE dan LPD tidak memenuhi syarat dikalibrasi.'),
        ('Temuan', 'Data menuntut LPE 0,410–0,490, tetapi banyak pasangan LPE–LPD yang sama baiknya (LPE 0,26–0,80): data hanya bisa menentukan selisih keduanya.'),
        ('Penyebab', 'Lonjakan pascapandemi 2022–2024 (puncak 24,86%) di luar cakupan model.')],
      'Tetap 0,27726 dan 0,207945', TEAL, WHITE),
     ('TAHAP 4', 'Uji kompatibilitas', 'Hasil Tahap 2 di model penuh', [
        ('Diperiksa', 'Laju konversi 0,0639 dicoba pada model penuh (13 baris variabel).'),
        ('Temuan', '12 baris praktis tidak berubah, tetapi lahan terbangun pada uji penuh memburuk 32,4% (rasio 1,32 > 1,05, pemicu M1).'),
        ('Penyebab', 'Uji P4 mulai 2016 dari lahan rendah; uji penuh mulai 2019 dari lahan yang sudah tinggi.')],
      'Ditolak → kembali 0,0436', RED, WHITE)]
w, g = 3.5, 0.175
for k, (tag, title, par, items, badge, bfill, bcol) in enumerate(C):
    x = 0.9 + k * (w + g)
    paras = [[(tag, 10, True, GREY)], [(title, 13, True, TEAL)], [(par, 10.5, True, DARK, True)]]
    for lab, txt in items:
        paras.append([(lab + '  ', 11, True, TEAL), (txt, 11, False, DARK)])
    box(s, x, 1.95, w, 6.25, fill=LIGHT, line=LINE, paras=paras, margin=0.15)
    box(s, x + 0.15, 7.5, w - 0.3, 0.55, fill=bfill, line=None, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, margin=0.05,
        paras=[[(badge, 10.5, True, bcol)]])
    if k < 4:
        arrow(s, x + w - 0.02, 4.9, g + 0.04, 0.3)
box(s, 0.9, 8.35, 18.2, 1.55, fill=TEAL, line=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.25, paras=[
    [('3 parameter dikalibrasi (batas TPK, laju demolisi, laju konversi) → 0 diganti', 14, True, YEL)],
    [('Nilai yang dihitung dari data sudah dapat dipertanggungjawabkan. Selisih yang tersisa pada wisatawan berasal dari lonjakan pascapandemi di luar cakupan model, bukan dari nilai parameter.', 11.5, False, WHITE)]])
source(s, 'Sumber: Buku Subbab 4.6 (Tabel 48–55). LPE = laju pertumbuhan eksternal; LPD = laju penurunan dasar.', y=10.12)
s.notes_slide.notes_text_frame.text = (
    'Kalibrasi dijalankan dalam lima tahap. Tahap 0 bukan kalibrasi, melainkan koreksi: cara menghitung LPE dan LPD disesuaikan dengan persamaan model, dan uji ekstrem tetap memberi hasil yang sama. '
    'Tahap 1 pada subsistem akomodasi: nilai terbaik memperbaiki 11,64 persen, tetapi berada di batas rentang dan ternyata ditarik oleh tahun 2019; tanpa tahun itu perbaikannya tinggal 5,74 persen. Jadi nilai dari data dipertahankan. '
    'Tahap 2 pada subsistem lahan: laju konversi 0,0639 memperbaiki 14,76 persen dan memenuhi syarat, sehingga sementara diadopsi. '
    'Tahap 3 pada subsistem wisatawan bersifat diagnostik: data menuntut LPE yang lebih tinggi, tetapi data hanya bisa menentukan selisih LPE dan LPD, dan penyebab selisihnya adalah lonjakan pascapandemi. '
    'Tahap 4 menguji hasil Tahap 2 pada model penuh: lahan terbangun justru memburuk 32,4 persen, sehingga hasilnya ditolak. Kesimpulannya, dari tiga parameter yang dikalibrasi, tidak ada yang diganti.')
p.save('/tmp/pptwork/v5/kalibrasi_2slide.pptx')
print('ok')
