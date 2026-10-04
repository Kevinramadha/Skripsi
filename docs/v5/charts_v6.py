exec(open('/tmp/pptwork/helpers_v4.py').read())
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION
from pptx.util import Pt, Inches
p = Presentation('/tmp/pptwork/v5/out.pptx')
S = list(p.slides)
NAVY = '0B5E6E'

def remove_ids(slide, ids):
    for sh in list(slide.shapes):
        if sh.shape_id in ids:
            sh._element.getparent().remove(sh._element)

def add_chart(slide, ctype, cats, series, x, y, w, h, fmt='0.0', legend=False, labels=True, font=12):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, vals)
    gf = slide.shapes.add_chart(ctype, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.font.size = Pt(font); ch.font.name = F; ch.font.color.rgb = rgb(DARK)
    ch.has_title = False
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(font)
    if labels:
        pl = ch.plots[0]
        pl.has_data_labels = True
        dl = pl.data_labels
        dl.number_format = fmt; dl.number_format_is_linked = False
        dl.font.size = Pt(font); dl.font.color.rgb = rgb(DARK)
    try:
        ch.value_axis.has_major_gridlines = True
        ch.value_axis.major_gridlines.format.line.color.rgb = rgb('E3EEF1')
        ch.value_axis.format.line.fill.background()
        ch.value_axis.tick_labels.font.size = Pt(font - 1)
        ch.category_axis.tick_labels.font.size = Pt(font)
        ch.category_axis.format.line.color.rgb = rgb('B8C7CC')
    except Exception:
        pass
    return ch

def color_series(ch, cols):
    for s, c in zip(ch.plots[0].series, cols):
        f = s.format.fill; f.solid(); f.fore_color.rgb = rgb(c)

def color_points(ch, cols):
    s = ch.plots[0].series[0]
    for k, c in enumerate(cols):
        pt = s.points[k]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(c)

# 1. slide 3: CAGR (urutan dibalik karena bar chart digambar dari bawah)
s = S[2]; remove_ids(s, {67})
cats = ['Pertanian', 'Kesehatan', 'Konstruksi', 'Perdagangan besar & eceran', 'Industri pengolahan', 'Jasa keuangan', 'Informasi & komunikasi', 'Perjalanan & pariwisata']
vals = [1.5, 1.9, 2.1, 2.4, 2.6, 3.4, 3.4, 3.6]
ch = add_chart(s, XL_CHART_TYPE.BAR_CLUSTERED, cats, [('CAGR (%)', vals)], 0.9, 2.7, 9.6, 6.7, fmt='0.0', font=13)
ch.plots[0].gap_width = 45
color_points(ch, [CYAN] * 7 + [YEL])
ch.value_axis.minimum_scale = 0; ch.value_axis.maximum_scale = 4

# 2. slide 4: kontribusi PDB (garis) dan kunjungan (kolom)
s = S[3]; remove_ids(s, {71, 73, 78, 79, 80, 81, 82, 83})
yrs = ['2020', '2021', '2022', '2023', '2024', '2025']
ch = add_chart(s, XL_CHART_TYPE.LINE_MARKERS, yrs, [('Kontribusi terhadap PDB (%)', [2.2, 2.3, 3.7, 3.9, 4.0, 4.9])], 0.9, 2.6, 8.9, 6.8, fmt='0.0"%"', font=13)
ser = ch.plots[0].series[0]
ser.format.line.color.rgb = rgb(YEL); ser.format.line.width = Pt(3.5); ser.smooth = False
ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = rgb(YEL); ser.marker.size = 9
ser.marker.format.line.color.rgb = rgb(YEL)
ch.plots[0].data_labels.position = XL_LABEL_POSITION.ABOVE
ch.value_axis.minimum_scale = 1.5; ch.value_axis.maximum_scale = 5.5
ch.value_axis.tick_labels.number_format = '0.0'; ch.value_axis.tick_labels.number_format_is_linked = False
ch = add_chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, yrs,
               [('Wisatawan mancanegara (juta kunjungan)', [4.1, 1.6, 5.9, 11.7, 13.9, 15.4]),
                ('Wisatawan nusantara (juta perjalanan)', [524.6, 613.3, 734.9, 839.7, 1021.1, 1200.3])],
               10.23, 2.6, 9.0, 6.8, fmt='#,##0.0', legend=True, font=12)
color_series(ch, [YEL, CYAN]); ch.plots[0].gap_width = 60; ch.plots[0].overlap = 0
ch.plots[0].data_labels.position = XL_LABEL_POSITION.OUTSIDE_END

# 3. slide 42: uji feedback loop
s = S[41]; remove_ids(s, {98})
cats = ['Semua balancing loop dimatikan (hanya R1)', 'Goal-seeking tenaga kerja dimatikan', 'Okupansi akomodasi dimatikan',
        'R2 ekonomi–atraksi dimatikan', 'Daya dukung lahan dimatikan', 'B1 kepadatan dimatikan', 'Simulasi dasar (semua loop aktif)']
vals = [217.37, 96.20, 96.23, 80.24, 111.66, 250.32, 96.20]
ch = add_chart(s, XL_CHART_TYPE.BAR_CLUSTERED, cats, [('Jumlah wisatawan 2050 (juta)', vals)], 0.9, 3.1, 11.4, 7.0, fmt='0.0', font=12)
ch.plots[0].gap_width = 45
color_points(ch, [CYAN, CYAN, CYAN, CYAN, CYAN, RED, GREY])
ch.value_axis.minimum_scale = 0; ch.value_axis.maximum_scale = 300

# 4. slide 45: data wisnus mentah vs tersambung
s = S[44]; remove_ids(s, {98})
ch = add_chart(s, XL_CHART_TYPE.LINE_MARKERS, ['2016', '2017', '2018', '2019'],
               [('Data mentah BPS', [6.551294, 6.644412, 7.996959, 20.520104]),
                ('Data tersambung (×2,2997)', [15.066309, 15.280457, 18.390971, 20.520104])],
               0.9, 3.1, 11.0, 6.9, fmt='0.0', legend=True, font=13)
for ser, c in zip(ch.plots[0].series, [GREY, CYAN]):
    ser.format.line.color.rgb = rgb(c); ser.format.line.width = Pt(3.5); ser.smooth = False
    ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = rgb(c); ser.marker.size = 9
    ser.marker.format.line.color.rgb = rgb(c)
ch.plots[0].data_labels.position = XL_LABEL_POSITION.ABOVE
ch.value_axis.minimum_scale = 0; ch.value_axis.maximum_scale = 25

# 5. slide 47: DC uji parsial vs uji penuh
s = S[46]; remove_ids(s, {79})
ch = add_chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, ['Jumlah hotel', 'Jumlah ODTW', 'Tenaga kerja', 'Lahan terbangun', 'TPK'],
               [('Uji parsial', [0.2011, 0.3367, 0.4945, 0.4813, 0.6692]), ('Uji penuh', [0.8228, 0.7855, 0.8860, 0.6619, 0.6224])],
               0.9, 3.05, 11.1, 6.3, fmt='0.00', legend=True, font=13)
color_series(ch, [CYAN, NAVY]); ch.plots[0].gap_width = 70
ch.plots[0].data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
ch.value_axis.minimum_scale = 0; ch.value_axis.maximum_scale = 1

# 6. slide 54: pengaruh maksimum ±10%
s = S[53]
remove_ids(s, {67, 134} | set(range(68, 104)))
cats = ['Sensitivitas ODTW thd investasi', 'ODTW referensi', 'Laju pertumbuhan eksternal', 'RDDL referensi',
        'Bobot ODTW', 'Bobot kepadatan', 'Kepadatan referensi', 'Laju penurunan dasar']
vals = [2.57, 3.54, 5.06, 5.09, 7.27, 7.49, 8.61, 42.07]
GREY2, LCYAN = '8A9BA3', '9EDCE6'
ch = add_chart(s, XL_CHART_TYPE.BAR_CLUSTERED, cats, [('Pengaruh maksimum (%)', vals)], 0.9, 3.05, 11.3, 6.05, fmt='0.00"%"', font=12)
ch.plots[0].gap_width = 45
color_points(ch, [GREY2, LCYAN, NAVY, LCYAN, YEL, YEL, LCYAN, YEL])
ch.value_axis.minimum_scale = 0; ch.value_axis.maximum_scale = 50
for k, (lab, c) in enumerate([('Asumsi peneliti', YEL), ('Nilai acuan (normalisasi)', LCYAN), ('Dihitung dari data', NAVY), ('Persamaan stok-aliran', GREY2)]):
    box(s, 1.2 + k * 2.8, 9.3, 0.3, 0.3, fill=c, line=None, paras=None, radius=0.1)
    tb(s, 1.6 + k * 2.8, 9.27, 2.4, 0.36, [(lab, 11.5, False, DARK)], anchor=MSO_ANCHOR.MIDDLE)

# 7. slide 69: SUS per responden
s = S[68]; remove_ids(s, {126})
ch = add_chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, ['R%d' % k for k in range(1, 11)],
               [('Skor SUS', [97.5, 87.5, 100, 100, 85, 87.5, 97.5, 60, 70, 100])], 0.9, 3.05, 10.4, 5.0, fmt='0.0', font=12)
ch.plots[0].gap_width = 45
color_points(ch, [CYAN] * 7 + [RED] + [CYAN] * 2)
ch.plots[0].data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
ch.value_axis.minimum_scale = 0; ch.value_axis.maximum_scale = 110
ch.value_axis.major_unit = 20

p.save('/tmp/pptwork/v5/out6.pptx')
print('ok')
