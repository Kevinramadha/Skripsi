import sys; sys.path.insert(0, '/tmp/pptwork/v5')
from single import make_single
make_single('/tmp/pptwork/v5/out7.pptx', 31, '/tmp/pptwork/v5/_base_diag.pptx')
exec(open('/tmp/pptwork/helpers_v4.py').read())
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
p = Presentation('/tmp/pptwork/v5/_base_diag.pptx')
s = p.slides[0]
NAVY = '0B5E6E'; GREEN = 'C9F0D6'; PINK = 'FDE2DE'
def find_title(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.top is not None and 0.6*E <= sh.top <= 1.3*E and sh.width > 9*E and sh.text_frame.text.strip():
            return sh
t = find_title(s)
for sh in list(s.shapes):
    keep = (sh._element is t._element) or sh.top >= 10.3*E or (sh.left >= 18.0*E and sh.top < 1.0*E) or sh.top < 0.45*E
    if not keep:
        sh._element.getparent().remove(sh._element)
for sh in list(s.shapes):
    if sh.top > 10.3*E and 16.5*E < sh.left < 17.9*E:
        sh._element.getparent().remove(sh._element)
for sh in [x for x in s.shapes if x.has_text_frame and x.left > 17.9*E and x.top > 10.4*E and x.text_frame.text.strip().isdigit()]:
    set_text(sh, '')
ps = t.text_frame.paragraphs
ps[0].runs[0].text = 'Estimasi Lahan Terbangun — Diagnosis Ketidakstabilan Temporal'
ps[1].runs[0].text = 'Naik-turunnya luas Dynamic World mengikuti jumlah citra yang terekam, bukan perubahan lahan'

yrs = [str(y) for y in range(2016, 2026)]
lunak = [53500, 62527, 67508, 73595, 68309, 67656, 65024, 79881, 77437, 65019]
kor = [61642, 67827, 63122, 62572, 66751, 71390, 70269, 74887, 74115, 67880]
obs = [4.3, 7.0, 16.2, 22.4, 13.5, 8.5, 7.1, 16.7, 15.2, 9.3]
cd = CategoryChartData(); cd.categories = yrs
cd.add_series('Luas sebelum koreksi', lunak); cd.add_series('Luas setelah koreksi', kor)
ch = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.9), Inches(1.95), Inches(11.0), Inches(4.3), cd).chart
ch.has_title = False; ch.font.size = Pt(11); ch.font.name = F; ch.font.color.rgb = rgb(DARK)
ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.TOP; ch.legend.include_in_layout = False
for k, (ser, col) in enumerate(zip(ch.plots[0].series, [RED, NAVY])):
    ser.format.line.color.rgb = rgb(col); ser.format.line.width = Pt(2.75); ser.smooth = False
    ser.marker.size = 8; ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = rgb(col); ser.marker.format.line.color.rgb = rgb(col)
    if k == 1: ser.format.line.dash_style = 4
ch.value_axis.minimum_scale = 45000; ch.value_axis.maximum_scale = 85000
ch.value_axis.tick_labels.number_format = '#,##0'; ch.value_axis.tick_labels.number_format_is_linked = False
ch.value_axis.has_major_gridlines = True; ch.value_axis.major_gridlines.format.line.color.rgb = rgb('E3EEF1')
ch.value_axis.tick_labels.font.size = Pt(10); ch.category_axis.tick_labels.font.size = Pt(10)
cd2 = CategoryChartData(); cd2.categories = yrs; cd2.add_series('Jumlah observasi citra per piksel', obs)
ch2 = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.9), Inches(6.3), Inches(11.0), Inches(2.9), cd2).chart
ch2.has_title = False; ch2.font.size = Pt(11); ch2.font.name = F; ch2.font.color.rgb = rgb(DARK)
ch2.has_legend = True; ch2.legend.position = XL_LEGEND_POSITION.TOP; ch2.legend.include_in_layout = False
ch2.plots[0].gap_width = 60
ser = ch2.plots[0].series[0]; ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(CYAN)
for k in (3, 7):
    pt = ser.points[k]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb('E0A800')
ch2.plots[0].has_data_labels = True
dl = ch2.plots[0].data_labels; dl.number_format = '0.0'; dl.number_format_is_linked = False; dl.font.size = Pt(10); dl.position = XL_LABEL_POSITION.OUTSIDE_END
ch2.value_axis.minimum_scale = 0; ch2.value_axis.maximum_scale = 26; ch2.value_axis.has_major_gridlines = False
ch2.value_axis.tick_labels.font.size = Pt(10); ch2.category_axis.tick_labels.font.size = Pt(10)
tb(s, 0.9, 9.2, 11.0, 0.6, [('Tahun dengan citra terbanyak (kuning: 2019, 2023) juga tahun dengan luas terbesar. Grafik atas memakai luas probabilitas Dynamic World (dasar perhitungan diagnosis, ha), sehingga angkanya sedikit berbeda dari data model.', 10.5, False, GREY, True)])

box(s, 12.3, 1.95, 6.8, 1.45, fill=PINK, line=None, paras=[('Masalah: turun 27% dalam dua tahun', 13, True, RED),
    ('Data model: 75.624 ha (2023) → 55.030 ha (2025). Lahan terbangun tidak mungkin berkurang sebesar itu di lapangan.', 11.5, False, DARK)])
rows = [['Pemeriksaan', 'Hasil', 'Arti'],
        ['Luas vs jumlah observasi', 'r = 0,788 (p = 0,007)', 'Berhubungan kuat'],
        ['Jumlah observasi vs tahun', 'r = 0,174 (p = 0,631)', 'Tidak berhubungan → koreksi sah'],
        ['Setelah koreksi: luas vs observasi', 'r ≈ 0', 'Pengaruh observasi hilang'],
        ['Variasi perubahan tahunan', '12,3% → 6,6%', 'Data jadi lebih stabil']]
gf = table(s, 12.3, 3.55, 6.8, [2.6, 2.0, 2.2], rows, size=11, rowh=0.5, bold_first_col=True, hl={(1, 1): PINK, (2, 1): GREEN, (3, 1): GREEN, (4, 1): GREEN})
box(s, 12.3, 6.2, 6.8, 1.35, fill=LIGHT, paras=[('Cara koreksi (penyesuaian kovariat, mirip ANCOVA)', 12, True, TEAL),
    ('Luasₜ = a + b · Obsₜ  →  Luas terkoreksiₜ = Luasₜ − b · (Obsₜ − rata-rata Obs)', 11, False, DARK)])
box(s, 12.3, 7.7, 6.8, 1.5, fill=YEL_L, line=YEL, paras=[('Keputusan', 12.5, True, TEAL),
    ('Data model tetap memakai nilai asli, karena tahun 2025 sudah menjadi dasar semua parameter dan kalibrasi. Koreksi hanya untuk pemeriksaan; noise dicatat sebagai keterbatasan.', 11, False, DARK)])
source(s, 'Sumber: Buku Subbab 3.5.2 dan 4.2.2; notebook "Data Citra untuk Lahan Terbangun" (Lahan_Terbangun_Terkoreksi_DIY_2016_2025.csv). Diagnosis dihitung pada luas probabilitas Dynamic World.', y=10.0)
s.notes_slide.notes_text_frame.text = ('Data Dynamic World turun 27 persen dari 2023 ke 2025, yang tidak mungkin terjadi di lapangan. Saya periksa penyebabnya: luas ternyata berhubungan kuat dengan '
    'jumlah citra yang terekam per piksel setiap tahun, r 0,788, sedangkan jumlah citra tidak berhubungan dengan tahun, sehingga koreksi sah dilakukan. '
    'Setelah dikoreksi, pengaruh jumlah citra hilang dan data menjadi lebih stabil. Meski begitu, model tetap memakai nilai asli karena tahun 2025 sudah menjadi dasar semua parameter; '
    'koreksi hanya untuk pemeriksaan dan noise ini saya catat sebagai keterbatasan.')
p.save('/tmp/pptwork/v5/diagnosis_1slide.pptx')
print('ok')
