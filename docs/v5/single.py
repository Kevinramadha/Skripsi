from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
def make_single(src, idx, out):
    p = Presentation(src)
    keep = p.slides[idx - 1]
    # buang hyperlink antarslide pada slide yang disimpan
    slide_rids = [rId for rId, rel in list(keep.part.rels.items()) if rel.reltype == RT.SLIDE]
    for el in keep._element.iter():
        for attr in ('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id',):
            pass
    ns = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
    rid_attr = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
    for tag in ('hlinkClick', 'hlinkMouseOver'):
        for h in list(keep._element.iter(ns + tag)):
            if h.get(rid_attr) in slide_rids or h.get('action', '').startswith('ppaction://hlinksldjump'):
                h.getparent().remove(h)
    for rId in slide_rids:
        keep.part.drop_rel(rId)
    lst = p.slides._sldIdLst
    for k, el in reversed(list(enumerate(list(lst), 1))):
        if k != idx:
            p.part.drop_rel(el.rId); lst.remove(el)
    for rId, rel in list(p.part.rels.items()):
        if rel.reltype == RT.NOTES_SLIDE:
            p.part.drop_rel(rId)
    p.save(out)

def make_subset(src, idxs, out):
    p = Presentation(src)
    ns = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
    rid_attr = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
    for idx in idxs:
        keep = p.slides[idx - 1]
        slide_rids = [rId for rId, rel in list(keep.part.rels.items()) if rel.reltype == RT.SLIDE]
        for tag in ('hlinkClick', 'hlinkMouseOver'):
            for h in list(keep._element.iter(ns + tag)):
                if h.get(rid_attr) in slide_rids or h.get('action', '').startswith('ppaction://hlinksldjump'):
                    h.getparent().remove(h)
        for rId in slide_rids:
            keep.part.drop_rel(rId)
    lst = p.slides._sldIdLst
    for k, el in reversed(list(enumerate(list(lst), 1))):
        if k not in idxs:
            p.part.drop_rel(el.rId); lst.remove(el)
    for rId, rel in list(p.part.rels.items()):
        if rel.reltype == RT.NOTES_SLIDE:
            p.part.drop_rel(rId)
    p.save(out)
