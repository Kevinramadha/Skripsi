import sys; sys.path.insert(0, '/tmp/pptwork/v5')
from single import make_subset
make_subset('/tmp/pptwork/v5/out7.pptx', [66, 69, 70], '/tmp/pptwork/v5/_base_t3.pptx')
exec(open('/tmp/pptwork/helpers_v4.py').read())
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
p = Presentation('/tmp/pptwork/v5/_base_t3.pptx')
GREEN, PINK, NAVY = 'C9F0D6', 'FDE2DE', '0B5E6E'
US = '/tmp/pptwork/v5/us/'


def find_title(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.top is not None and 0.6*E <= sh.top <= 1.3*E and sh.width > 9*E and sh.text_frame.text.strip():
            return sh


def prep(s, title, resume):
    t = find_title(s)
    for sh in list(s.shapes):
        keep = (sh._element is t._element) or sh.top >= 10.3*E or (sh.left >= 18.0*E and sh.top < 1.0*E) or sh.top < 0.45*E
        if not keep:
            sh._element.getparent().remove(sh._element)
    for sh in list(s.shapes):
        if sh.top > 10.3*E and 16.5*E < sh.left < 17.9*E:
            sh._element.getparent().remove(sh._element)
    for sh in [x for x in s.shapes if x.has_text_frame and x.left > 17.9*E and x.top > 10.4*E and x.text_frame.text.strip().isdigit()]:
        set_text(sh, '')
    ps = t.text_frame.paragraphs
    ps[0].runs[0].text = title
    ps[1].runs[0].text = resume


def intinya(s, text, y=9.05, h=0.85):
    return box(s, 0.9, y, 18.2, h, fill=YEL_L, line=YEL, anchor=MSO_ANCHOR.MIDDLE,
               paras=[[('Intinya  ', 14, True, TEAL), (text, 14, False, DARK)]])


def head(s, x, y, w, text, sub=None):
    box(s, x, y, w, 0.5, fill=TEAL, line=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.15,
        paras=[[(text, 13.5, True, WHITE)] + ([('   ' + sub, 11, False, 'D7EEF2')] if sub else [])])


def qa(s, x, y, w, h, rows, size=12.5):
    """rows: list of (label, text[, color])"""
    paras = []
    for r in rows:
        paras.append([(r[0] + '  ', size, True, r[2] if len(r) > 2 else TEAL), (r[1], size, False, DARK)])
    return box(s, x, y, w, h, fill=LIGHT, line=LINE, paras=paras)


def chip(s, x, y, w, h, text, fill=LIGHT, line=LINE, size=10.5, color=DARK, bold=False):
    return box(s, x, y, w, h, fill=fill, line=line, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, margin=0.06,
               paras=[(text, size, bold, color)])


def outline(s, x, y, w, h, color=RED):
    r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.background(); r.line.color.rgb = rgb(color); r.line.width = Pt(2.5); r.shadow.inherit = False
    return r


from pptx.enum.chart import XL_CHART_TYPE
S = list(p.slides)
# ================= 1
s = S[0]
prep(s, 'Aplikasi Simulasi Berbasis Web', 'Model bisa dicoba langsung di browser, tanpa perlu software pemodelan')
box(s, 0.9, 1.95, 8.9, 1.55, fill=PINK, line=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.25, paras=[
    [('Masalahnya', 12.5, True, '9B2C1F')],
    [('Hasil model biasanya berhenti di laporan dan grafik. Untuk menjalankannya perlu Vensim dan keahlian pemodelan, jadi sulit dipakai pemerintah.', 12, False, DARK)]])
box(s, 0.9, 3.65, 8.9, 1.25, fill=GREEN, line=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.25, paras=[
    [('Jawabannya', 12.5, True, '1E7A43')],
    [('Aplikasi web: perhitungan tetap dari model, aplikasi hanya menerima masukan dan menampilkan hasil.', 12, False, DARK)]])
tb(s, 0.9, 5.15, 8.9, 0.45, [[('Cara kerjanya', 12.5, True, TEAL)]])
for k, (a, b) in enumerate([('Model Vensim', 'model yang sudah diuji'), ('SDEverywhere', 'diubah jadi kode web'), ('Aplikasi web', 'Next.js · TypeScript · Tailwind')]):
    x = 0.9 + k * 3.05
    box(s, x, 5.65, 2.65, 1.25, fill=TEAL if k == 2 else LIGHT, line=None if k == 2 else LINE, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, margin=0.08,
        paras=[[(a, 12.5, True, WHITE if k == 2 else TEAL)], [(b, 10, False, 'D7EEF2' if k == 2 else GREY)]])
    if k < 2: arrow(s, x + 2.7, 6.12, 0.3, 0.3)
rows = [['Kebutuhan pengguna', 'Fitur di aplikasi'],
        ['Memahami isi model dan seberapa valid', 'Halaman Model: struktur dan hasil pengujian'],
        ['Membandingkan pilihan kebijakan', 'Halaman Skenario: BAU, Sustainable, DP'],
        ['Mencoba nilai kebijakan sendiri', 'Halaman Eksplorasi: atur 2 kebijakan, jalankan 2025–2050'],
        ['Bisa dibuka tanpa instalasi', 'Aplikasi web, berjalan di browser']]
table(s, 10.2, 1.95, 8.9, [3.8, 5.1], rows, size=11, rowh=0.7)
s.shapes.add_picture('/tmp/pptwork/v5/us/qr_app.png', Inches(10.2), Inches(5.85), Inches(2.1), Inches(2.1))
box(s, 12.5, 5.85, 6.6, 2.1, fill=LIGHT, line=LINE, anchor=MSO_ANCHOR.MIDDLE, margin=0.25, paras=[
    [('Coba langsung', 12.5, True, TEAL)], [('sistemdinamispariwisata.vercel.app', 14, True, DARK)], [('Pindai kode QR di samping', 10.5, False, GREY)]])
intinya(s, 'Aplikasi menjembatani model dengan pengguna di Pemda DIY dan Dinas Pariwisata yang tidak punya latar pemodelan.', y=8.6, h=0.95)
source(s, 'Sumber: Buku Subbab 3.9 dan 4.9.1.', y=10.12)
s.notes_slide.notes_text_frame.text = (
    'Tujuan ketiga adalah membuat aplikasi. Masalahnya, hasil model sistem dinamis biasanya berhenti di laporan, dan untuk menjalankannya perlu Vensim serta keahlian pemodelan. '
    'Karena itu model saya ubah menjadi aplikasi web. Model Vensim dikonversi dengan SDEverywhere, lalu dijalankan di aplikasi yang dibangun dengan Next.js. Perhitungannya tetap dari model, aplikasi hanya menerima masukan dan menampilkan hasil. '
    'Setiap fitur menjawab kebutuhan pengguna, dan aplikasinya bisa dibuka langsung melalui alamat atau kode QR ini.')

# ================= 2
s = S[1]
prep(s, 'Halaman Aplikasi', 'Dari memilih kebijakan sampai melihat grafik hasil, semua di satu tempat')
P = [('app_beranda.png', 'Beranda', 'Gambaran umum dan pintu masuk ke setiap halaman'),
     ('app_model.png', 'Model', 'Struktur model (CLD, SFD) dan hasil pengujiannya'),
     ('app_skenario.png', 'Skenario', 'Menjalankan dan membandingkan BAU, Sustainable, DP'),
     ('app_eksplorasi.png', 'Eksplorasi Simulasi', 'Mengatur dua kebijakan sendiri dan memilih variabel yang ditampilkan')]
for k, (img, t1, t2) in enumerate(P):
    x = 0.9 + (k % 2) * 9.25; y = 1.95 + (k // 2) * 3.4
    pc = pic_fit(s, US + img, x, y, 4.6, 3.1); pc.line.color.rgb = rgb(LINE); pc.line.width = Pt(1)
    box(s, x + 4.8, y + 0.4, 4.15, 2.3, fill=LIGHT, line=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.22, paras=[
        [(str(k + 1) + '  ' + t1, 14, True, TEAL)], [(t2, 12, False, DARK)]])
box(s, 0.9, 8.8, 18.2, 0.95, fill=YEL_L, line=YEL, anchor=MSO_ANCHOR.MIDDLE, margin=0.3, paras=[[
    ('Alurnya:  ', 13, True, TEAL), ('pilih skenario atau atur insentif (0–0,5) dan konservasi lahan (0–1)  →  model menghitung 2025–2050 di browser  →  hasil tampil sebagai angka dan grafik', 12.5, False, DARK)]])
source(s, 'Sumber: Buku Subbab 4.9.1 (Gambar 62–66). Tampilan dipotong pada bagian atas tiap halaman.', y=10.12)
s.notes_slide.notes_text_frame.text = (
    'Aplikasinya punya empat halaman. Beranda sebagai gambaran umum. Halaman Model menampilkan struktur model dan hasil pengujiannya. '
    'Halaman Skenario untuk menjalankan dan membandingkan tiga skenario. Halaman Eksplorasi untuk mencoba nilai kebijakan sendiri. '
    'Alurnya sederhana: pilih skenario atau atur dua kebijakan, model menghitung di browser, lalu hasilnya tampil sebagai angka dan grafik. (Bila memungkinkan, tunjukkan demo singkat di sini.)')

# ================= 3
s = S[2]
prep(s, 'Hasil Evaluasi Aplikasi', 'Semua fungsi berjalan, dan pengguna menilai aplikasi mudah dipakai')
box(s, 0.9, 1.95, 8.6, 0.55, fill=TEAL, line=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.2, paras=[[('Fungsi aplikasi', 13.5, True, WHITE), ('   uji black-box', 10.5, False, 'D7EEF2')]])
tb(s, 0.9, 2.6, 2.6, 1.2, [[('9/9', 40, True, TEAL)], [('fungsi berjalan sesuai rancangan', 10.5, False, GREY)]])
FN = ['Navigasi antarhalaman', 'Ganti tampilan Model', 'Pilih skenario', 'Bandingkan skenario', 'Jalankan simulasi',
     'Pilih variabel keluaran', 'Ubah insentif kebijakan', 'Ubah konservasi lahan', 'Kembalikan ke BAU (reset)']
for k, f_ in enumerate(FN):
    y = 2.6 + k * 0.6
    box(s, 3.7, y, 5.8, 0.5, fill=LIGHT if k % 2 == 0 else WHITE, line=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.15,
        paras=[[('✓  ', 12, True, '1E7A43'), (f_, 11.5, False, DARK)]])
box(s, 10.0, 1.95, 9.1, 0.55, fill=TEAL, line=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.2, paras=[[('Kemudahan penggunaan', 13.5, True, WHITE), ('   System Usability Scale (SUS), 10 responden', 10.5, False, 'D7EEF2')]])
tb(s, 10.0, 2.6, 3.2, 1.3, [[('88,50', 40, True, TEAL)], [('rata-rata skor SUS (0–100)', 10.5, False, GREY)]])
# skala
x0, x1, y = 13.4, 18.4, 3.25
def sx(v): return x0 + (v - 50) / 50 * (x1 - x0)
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x0), Inches(y), Inches(x1 - x0), Inches(0.18)); bar.fill.solid(); bar.fill.fore_color.rgb = rgb('D7EEF2'); bar.line.fill.background()
for v, lab, mode in [(68, 'rata-rata normatif 68', 'c'), (85.5, 'Excellent 85,5', 'r'), (90.9, 'Best Imaginable 90,9', 'l')]:
    t_ = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(sx(v) - 0.015), Inches(y - 0.1), Inches(0.03), Inches(0.38)); t_.fill.solid(); t_.fill.fore_color.rgb = rgb(GREY); t_.line.fill.background()
    if mode == 'c': tb(s, sx(v) - 1.0, y - 0.45, 2.0, 0.35, [[(lab, 9, False, GREY)]], align=PP_ALIGN.CENTER)
    elif mode == 'r': tb(s, sx(v) - 1.85, y + 0.32, 1.8, 0.35, [[(lab, 9, False, GREY)]], align=PP_ALIGN.RIGHT)
    else: tb(s, sx(v) + 0.05, y + 0.32, 1.9, 0.35, [[(lab, 9, False, GREY)]])
mk = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(sx(88.5) - 0.16), Inches(y - 0.07), Inches(0.32), Inches(0.32)); mk.fill.solid(); mk.fill.fore_color.rgb = rgb(RED); mk.line.fill.background()
tb(s, sx(88.5) - 0.6, y - 0.5, 1.2, 0.35, [[('88,50', 10, True, RED)]], align=PP_ALIGN.CENTER)
vals = [97.5, 87.5, 100, 100, 85, 87.5, 97.5, 60, 70, 100]
cd = CategoryChartData(); cd.categories = ['R%d' % i for i in range(1, 11)]; cd.add_series('Skor SUS', vals)
ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(10.0), Inches(4.35), Inches(9.1), Inches(3.4), cd).chart
ch.has_title = False; ch.has_legend = False; ch.font.size = Pt(10); ch.font.name = F; ch.font.color.rgb = rgb(DARK)
pl = ch.plots[0]; pl.gap_width = 50; pl.has_data_labels = True
dl = pl.data_labels; dl.number_format = '0.0'; dl.number_format_is_linked = False; dl.show_value = True; dl.font.size = Pt(10); dl.position = XL_LABEL_POSITION.OUTSIDE_END
for i, v in enumerate(vals):
    pt = pl.series[0].points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(TEAL if v >= 85 else CYAN)
ch.value_axis.minimum_scale = 0; ch.value_axis.maximum_scale = 110; ch.value_axis.has_major_gridlines = True
ch.value_axis.major_gridlines.format.line.color.rgb = rgb('E3EEF1'); ch.value_axis.format.line.fill.background(); ch.value_axis.tick_labels.font.size = Pt(9)
tb(s, 10.0, 7.75, 9.1, 0.4, [[('8 dari 10 responden memberi skor 85 atau lebih (batang gelap).', 10, False, GREY, True)]])
intinya(s, 'Aplikasi berjalan sesuai rancangan dan dinilai mudah dipakai: skor 88,50 jauh di atas rata-rata normatif SUS (68).', y=8.6, h=0.95)
source(s, 'Sumber: Buku Subbab 4.9.2–4.9.3; Brooke (2013); Bangor et al. (2009).', y=10.12)
s.notes_slide.notes_text_frame.text = (
    'Aplikasi diuji dari dua sisi. Dari sisi fungsi, sembilan fungsi utama diuji dengan black-box testing dan semuanya berjalan sesuai rancangan. '
    'Dari sisi kemudahan penggunaan, sepuluh responden mengisi System Usability Scale, dengan rata-rata 88,50. '
    'Angka ini jauh di atas rata-rata normatif SUS, yaitu 68, dan berada di antara kategori Excellent dan Best Imaginable. Delapan dari sepuluh responden memberi skor 85 atau lebih.')
p.save('/tmp/pptwork/v5/tujuan3_3slide.pptx')
print('ok')
