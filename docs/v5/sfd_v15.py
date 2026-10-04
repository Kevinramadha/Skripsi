import sys; sys.path.insert(0, '/tmp/pptwork/v5')
from single import make_single
make_single('/tmp/pptwork/v5/out7.pptx', 35, '/tmp/pptwork/v5/_base_sfd.pptx')
exec(open('/tmp/pptwork/helpers_v4.py').read())
from pptx import Presentation
p = Presentation('/tmp/pptwork/v5/_base_sfd.pptx')
s = p.slides[0]
NAVY = '0B5E6E'
def find_title(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.top is not None and 0.6*E <= sh.top <= 1.3*E and sh.width > 9*E and sh.text_frame.text.strip():
            return sh
t = find_title(s)
for sh in list(s.shapes):
    keep = (sh._element is t._element) or sh.top >= 10.3*E or (sh.left >= 18.0*E and sh.top < 1.0*E) or sh.top < 0.45*E
    if not keep:
        sh._element.getparent().remove(sh._element)
for sh in list(s.shapes):
    if sh.top > 10.3*E and 16.5*E < sh.left < 17.9*E:
        sh._element.getparent().remove(sh._element)
for sh in [x for x in s.shapes if x.has_text_frame and x.left > 17.9*E and x.top > 10.4*E and x.text_frame.text.strip().isdigit()]:
    set_text(sh, '')
ps = t.text_frame.paragraphs
ps[0].runs[0].text = 'Stock-Flow Diagram (SFD)'
ps[1].runs[0].text = 'CLD diterjemahkan menjadi model kuantitatif: 61 variabel dengan lima stok utama'

pic_fit(s, '/tmp/pptwork/dx/word/media/image12.png', 0.9, 1.95, 12.5, 6.4)
tb(s, 0.9, 8.4, 12.5, 0.6, [[('Cara membaca: ', 11, True, TEAL),
    ('kotak = stock (jumlah yang menumpuk) · pipa berkatup = flow (aliran masuk/keluar stok) · nama tanpa kotak = auxiliary atau parameter · panah = pengaruh antarvariabel', 11, False, DARK)]])
# empat angka
for k, (n, lab) in enumerate([('5', 'stock'), ('10', 'flow'), ('11', 'auxiliary'), ('35', 'parameter')]):
    box(s, 13.7 + k * 1.37, 1.95, 1.27, 1.15, fill=LIGHT, line=None, paras=[(n, 24, True, NAVY), (lab, 11, False, DARK)],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.02)
tb(s, 13.7, 3.15, 5.4, 0.55, [('26 variabel dihitung oleh model + 35 parameter input = 61 variabel', 11, False, GREY, True)])
rows = [['Stok', 'Nilai awal 2025'], ['Jumlah wisatawan', '40.695.654 kunjungan'], ['Hotel dan akomodasi', '2.291 unit'],
        ['Objek daya tarik wisata', '201 unit'], ['Tenaga kerja pariwisata', '364.994 jiwa'], ['Lahan terbangun', '55.029,54 ha']]
table(s, 13.7, 3.75, 5.4, [2.8, 2.6], rows, size=11.5, rowh=0.44, bold_first_c=True) if False else table(s, 13.7, 3.75, 5.4, [2.8, 2.6], rows, size=11.5, rowh=0.44, bold_first_col=True)
box(s, 13.7, 6.55, 5.4, 1.85, fill=YEL_L, line=YEL, paras=[('Dua variabel kebijakan', 13, True, TEAL),
    ('Insentif Kebijakan → menambah investasi (loop R2)', 11.5, False, DARK),
    ('Konservasi Lahan → mengurangi konversi lahan untuk pariwisata (loop B2)', 11.5, False, DARK)])
box(s, 0.9, 9.15, 18.2, 0.8, fill=NAVY, line=None, paras=[[('Intinya  ', 13, True, YEL),
    ('Model disimulasikan 2025–2050 dengan langkah waktu 1 tahun (metode Euler). Persamaan lengkap ada di Lampiran 15.', 12.5, False, WHITE)]], anchor=MSO_ANCHOR.MIDDLE)
source(s, 'Sumber: Buku Subbab 4.1.4, Gambar 10, Tabel 27–28; model Vensim final.', y=10.08)
s.notes_slide.notes_text_frame.text = ('CLD kemudian saya terjemahkan menjadi stock-flow diagram agar bisa disimulasikan. Model akhirnya punya 61 variabel: 5 stock, 10 flow, '
    '11 auxiliary, dan 35 parameter. Lima stoknya adalah jumlah wisatawan, hotel dan akomodasi, ODTW, tenaga kerja, dan lahan terbangun, dengan nilai awal tahun 2025 seperti di tabel. '
    'Dua variabel kebijakan masuk di dua tempat: Insentif Kebijakan menambah investasi di loop R2, dan Konservasi Lahan mengurangi konversi lahan untuk pariwisata di loop B2. '
    'Model disimulasikan 2025 sampai 2050 dengan langkah waktu satu tahun.')
p.save('/tmp/pptwork/v5/sfd_1slide.pptx')
print('ok')
