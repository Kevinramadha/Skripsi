exec(open('/tmp/pptwork/helpers_v4.py').read())
from pptx import Presentation
from pptx.oxml.ns import qn
import sys
sys.path.insert(0, '/tmp/pptwork/v5')
from rewrite import REW
from titles import TITLES
p = Presentation('/tmp/pptwork/v5/s4.pptx')
S = list(p.slides)
U = lambda u: S[(u if u <= 30 else u + 1) - 1]   # nomor slide di file user -> slide

def find_title(slide):
    best = None
    for sh in slide.shapes:
        if sh.has_text_frame and sh.top is not None and 0.6 * E <= sh.top <= 1.3 * E and sh.width > 9 * E and sh.text_frame.text.strip():
            best = sh
    return best

# ---------- 1. tulis ulang teks
def paras_of(slide):
    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == 6:
                yield from walk(sh.shapes); continue
            if sh.has_text_frame:
                yield from sh.text_frame.paragraphs
            if getattr(sh, 'has_table', False) and sh.has_table:
                for row in sh.table.rows:
                    for c in row.cells:
                        yield from c.text_frame.paragraphs
    yield from walk(slide.shapes)

def set_para(para, new):
    runs = para.runs
    if not runs:
        return
    # pertahankan run awal yang menjadi prefiks teks baru (mis. "Takeaway  " tebal)
    keep, acc = 0, ''
    for r in runs[:-1]:
        if new.startswith(acc + r.text) and r.text:
            acc += r.text; keep += 1
        else:
            break
    if keep:
        runs[keep].text = new[len(acc):]
        for r in runs[keep + 1:]:
            r._r.getparent().remove(r._r)
    else:
        runs[0].text = new
        for r in runs[1:]:
            r._r.getparent().remove(r._r)

hits = 0
for s in S:
    for para in paras_of(s):
        t = ''.join(r.text for r in para.runs)
        if t in REW:
            set_para(para, REW[t]); hits += 1
print('rewrite hits', hits, 'of', len(REW))

# ---------- 2. judul vs resume
TITLE_C, RES_C = '0B5E6E', '4A5A63'
import copy
def style_title(t, title, resume):
    while len(t.text_frame.paragraphs) < 2:
        t.text_frame.paragraphs[0]._p.addnext(copy.deepcopy(t.text_frame.paragraphs[0]._p))
    for extra in t.text_frame.paragraphs[2:]:
        extra._p.getparent().remove(extra._p)
    for k, (txt, size, bold, col) in enumerate([(title, 32, True, TITLE_C), (resume, 19, False, RES_C)]):
        para = t.text_frame.paragraphs[k]
        pPr = para._p.find(qn('a:pPr'))
        if pPr is not None:
            ln = pPr.find(qn('a:lnSpc'))
            if ln is not None:
                pPr.remove(ln)
        if not para.runs:
            para.add_run()
        set_para(para, txt)
        r = para.runs[0]
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = False
        r.font.name = FB if bold else F
        r.font.color.rgb = rgb(col)
        para.line_spacing = 1.0
        para.space_after = Pt(0)
        para.space_before = Pt(3 if k == 1 else 0)
    t.top = Inches(0.78)
done = 0
for u, (title, resume) in TITLES.items():
    t = find_title(U(u))
    if t is None:
        print('no title', u); continue
    style_title(t, title, resume); done += 1
print('titles', done)

# ---------- 3. slide baru: validasi spasial Dynamic World (setelah slide 30)
s = S[30]
t = find_title(s)
for sh in list(s.shapes):
    keep = (sh._element is t._element) or sh.top >= 10.3 * E or (sh.left >= 18.0 * E and sh.top < 1.0 * E) or sh.top < 0.45 * E
    if not keep:
        sh._element.getparent().remove(sh._element)
TITLES_NEW = ('Estimasi Lahan Terbangun — Validasi Spasial', 'Kesalahan Dynamic World terkumpul di tepi objek, bukan tersebar acak')
style_title(t, *TITLES_NEW)
D = '/tmp/pptwork/v5/'
pic_fit(s, D + 'dw_nonbuilt.png', 0.9, 2.15, 8.9, 4.75)
pic_fit(s, D + 'dw_built.png', 10.2, 2.15, 8.9, 4.75)
tb(s, 0.9, 6.95, 8.9, 0.35, [('Contoh titik yang diklasifikasi bukan terbangun (titik 8, kawasan bervegetasi)', 12.5, True, TEAL)], align=PP_ALIGN.CENTER)
tb(s, 10.2, 6.95, 8.9, 0.35, [('Contoh titik yang diklasifikasi terbangun (titik 69, atap bangunan)', 12.5, True, TEAL)], align=PP_ALIGN.CENTER)
cards = [('Cara', 'Titik sampel uji akurasi dicek ulang pada citra resolusi tinggi (Google Earth): kelas Dynamic World dibandingkan dengan kondisi lapangan.'),
         ('Titik yang tepat', 'Kawasan yang jelas terbangun atau jelas bervegetasi: kelas prediksi sesuai dengan kondisi lapangan.'),
         ('Titik yang salah', 'Berada di tepi objek, misalnya tepi jalan yang perkerasannya hanya menutupi ±7 dari 10 m piksel (piksel campuran).')]
for k, (h, b) in enumerate(cards):
    box(s, 0.9 + k * 6.13, 7.5, 5.95, 1.6, fill=LIGHT, paras=[(h, 14, True, TEAL), (b, 12.5, False, DARK)])
takeaway(s, 'Kesalahan bersifat sistematis di tepi objek (sejalan dengan commission error 34 dari 100 titik), bukan acak di seluruh wilayah; data tetap layak dipakai.', y=9.25, h=0.85)
source(s, 'Sumber: Buku Subbab 3.5.2 dan 4.2.2, Gambar 14–15 (sampel uji akurasi, citra Google Earth).', y=10.15)
s.notes_slide.notes_text_frame.text = ('Validasi spasial Dynamic World: beberapa titik dari uji akurasi saya cek ulang di citra resolusi tinggi. '
    'Titik di kawasan yang jelas terbangun atau jelas bervegetasi sesuai dengan kondisi lapangan. Titik yang salah umumnya berada di tepi objek, '
    'misalnya tepi jalan yang hanya menutupi sekitar 7 dari 10 meter piksel. Jadi kesalahannya sistematis di tepi objek, bukan acak di seluruh wilayah.')

# ---------- 4. lampiran baru
def lamp_slide(s, title):
    t = find_title(s)
    for sh in list(s.shapes):
        keep = (sh._element is t._element) or sh.top >= 10.3 * E or (sh.left >= 18.0 * E and sh.top < 1.0 * E) or sh.top < 0.45 * E
        if not keep:
            sh._element.getparent().remove(sh._element)
    set_para(t.text_frame.paragraphs[0], title)
    for extra in t.text_frame.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    return s

s = lamp_slide(S[91], 'Lampiran 17 · Daftar istilah (1/2): kode uji dan tahapan')
rows = [['Kode', 'Arti'],
 ['P1', 'Uji parsial subsistem akomodasi (jumlah hotel dan TPK); input dari subsistem lain diganti data aktual'],
 ['P1b', 'Variasi P1 yang memakai rata-rata kamar per unit dari data aktual, untuk mengecek robustness rancangan P1'],
 ['P2 · P3 · P4', 'Uji parsial subsistem ODTW · tenaga kerja · lahan terbangun'],
 ['P5', 'Uji parsial subsistem wisatawan; memakai data wisnus yang sudah disesuaikan sejak 2016'],
 ['Uji penuh', 'Seluruh model dijalankan bersama tanpa penggantian data (2019–2025, 9 variabel)'],
 ['E01–E17', '17 uji kondisi ekstrem: satu parameter dibuat nol atau dilipatgandakan (mis. E01 LPE = 0; E04 rasio investasi/PDRB = 0; E09 laju konversi ×10)'],
 ['Tahap 0–4', 'Tahap kalibrasi: 0 koreksi LPE/LPD · 1 akomodasi · 2 lahan · 3 wisatawan (hanya pemeriksaan) · 4 uji kecocokan pada model penuh'],
 ['Aturan R1–R4', 'Aturan keputusan kalibrasi (berbeda dengan loop R1–R2): R1 nilai dari data termasuk kumpulan nilai setara → dipertahankan; R2 kumpulan setara menyentuh kedua ujung rentang → tidak bisa ditentukan; R3 optimum di ujung rentang → cek leave-one-year-out; R4 lainnya → dipakai bila lolos Tahap 4'],
 ['M1 · M2 · S · T', 'Syarat penolakan Tahap 4: NRMSE variabel target memburuk > 5% · rata-rata NRMSE 9 variabel memburuk > 5% · uji ekstrem gagal · simulasi tidak stabil pada horizon diperpanjang'],
 ['D-A · D-B · D-C', 'Varian pemeriksaan Tahap 3: D-A periode 2016–2019 (LPE 0,440) · D-B 2016–2019 + 2022–2025 (0,490) · D-C periode D-A dengan faktor penyesuaian 2,22 atau 2,38 (0,410 / 0,470)'],
 ['C0–C6', 'Kondisi uji robustness skenario: C0 acuan (nilai final); C1–C6 menggeser rasio LPD/LPE (0,65; 0,90), bobot daya tarik (2 kombinasi lain), dan laju konversi dasar (batas bawah/atas selang 95%)']]
table(s, 0.9, 2.0, 18.2, [2.3, 15.9], rows, size=12.5, rowh=0.62, bold_first_col=True)
source(s, 'Sumber: Buku Subbab 3.7.4, 3.8.4, 4.4.5, 4.5, 4.6 (Tabel 54), 4.8.5.', y=10.15)
s.notes_slide.notes_text_frame.text = 'Lampiran istilah: kode uji dan tahapan yang muncul di slide pengujian, kalibrasi, dan skenario.'

s = lamp_slide(S[92], 'Lampiran 17 · Daftar istilah (2/2): ukuran dan singkatan')
rows = [['Istilah', 'Arti'],
 ['R1 · R2', 'Reinforcing loop (saling memperkuat): R1 pertumbuhan wisatawan · R2 ekonomi–atraksi'],
 ['B1–B4', 'Balancing loop (menahan pertumbuhan): B1 kepadatan · B2 daya dukung lahan · B3 kapasitas kamar · B4 tenaga kerja'],
 ['E1', 'Error rata-rata = |rata-rata simulasi − rata-rata data| ÷ rata-rata data; baik bila < 0,05'],
 ['E2', 'Error variasi = |simpangan baku simulasi − simpangan baku data| ÷ simpangan baku data; baik bila < 0,30'],
 ['DC', 'Discrepancy coefficient: 0 = sempurna; 0,4–0,7 = rata-rata sampai baik (Barlas)'],
 ['U1 · U2 · U3', 'Theil inequality: porsi error karena beda rata-rata (U1), beda variasi (U2), dan sisa acak (U3)'],
 ['MAPE · NRMSE', 'Mean absolute percentage error · RMSE dibagi rata-rata data (fungsi tujuan kalibrasi)'],
 ['LOOCV', 'Leave-one-out cross-validation: satu tahun dikeluarkan bergantian lalu diprediksi oleh model sisanya'],
 ['OA · UA · PA · F1 · κ', 'Akurasi klasifikasi: keseluruhan · user\'s accuracy · producer\'s accuracy · F1-score · Kappa'],
 ['Commission / omission error', 'Commission: terpetakan terbangun padahal bukan · Omission: terbangun tetapi tidak terpetakan'],
 ['LPE · LPD', 'Laju Pertumbuhan Eksternal · Laju Penurunan Dasar wisatawan'],
 ['RDDL', 'Rasio Daya Dukung Lahan = 1 − lahan terbangun ÷ luas wilayah DIY'],
 ['TPK · TPG', 'Tingkat Penghunian Kamar · rata-rata tamu per kamar'],
 ['ODTW · NTL', 'Objek Daya Tarik Wisata · night-time light (intensitas cahaya malam VIIRS)'],
 ['BAU · Sustainable · DP', 'Business-as-Usual (tanpa kebijakan) · insentif 0,1 + konservasi 1,0 · Development Priority (insentif 0,3 + konservasi 0)'],
 ['SUS', 'System Usability Scale, skor 0–100 (rata-rata umum ±68)']]
table(s, 0.9, 2.0, 18.2, [3.6, 14.6], rows, size=12.5, rowh=0.46, bold_first_col=True)
source(s, 'Sumber: Buku Subbab 3.5, 3.7.4, 3.9, 4.1.3; Barlas (1996); Sterman (1984).', y=10.15)
s.notes_slide.notes_text_frame.text = 'Lampiran istilah: ukuran statistik dan singkatan yang dipakai di seluruh slide.'

s = lamp_slide(S[93], 'Lampiran 18 · Kategori IRTS 2008 dan KBLI 2020')
rows = [['No', 'Aktivitas karakteristik pariwisata (IRTS 2008)', 'KBLI 2020 (kategori · golongan pokok)', 'Dipakai di model'],
 ['1', 'Akomodasi untuk pengunjung', 'I · 55 Penyediaan akomodasi', 'Ya'],
 ['2', 'Penyediaan makanan dan minuman', 'I · 56 Penyediaan makanan dan minuman', 'Ya'],
 ['3', 'Angkutan penumpang kereta api', 'H · 49', 'Tidak'],
 ['4', 'Angkutan penumpang jalan', 'H · 49', 'Tidak'],
 ['5', 'Angkutan penumpang air', 'H · 50', 'Tidak'],
 ['6', 'Angkutan penumpang udara', 'H · 51', 'Tidak'],
 ['7', 'Penyewaan alat transportasi', 'N · 77', 'Tidak'],
 ['8', 'Agen perjalanan dan jasa reservasi', 'N · 79', 'Tidak'],
 ['9', 'Aktivitas budaya', 'R · 90 Kesenian, 91 Perpustakaan, museum, budaya', 'Ya (ODTW: 91)'],
 ['10', 'Aktivitas olahraga dan rekreasi', 'R · 93 Olahraga dan rekreasi lainnya', 'Ya (ODTW: 93)'],
 ['11', 'Perdagangan eceran barang khas pariwisata', 'G · 47', 'Tidak'],
 ['12', 'Aktivitas khas pariwisata lainnya (ditetapkan tiap negara)', 'Sesuai penetapan negara', 'Tidak'],
 ['—', 'Bukan kategori IRTS: jasa lainnya', 'S, T, U', 'Ikut terhitung*']]
hl = {(r, 3): 'BDEFF3' for r in (1, 2, 9, 10)}
hl[(13, 3)] = YEL_L
table(s, 0.9, 2.0, 18.2, [0.7, 7.0, 7.0, 3.5], rows, size=13, rowh=0.5, hl=hl)
box(s, 0.9, 9.1, 18.2, 0.95, fill=YEL_L, line=YEL, paras=[[('Cakupan model  ', 13, True, TEAL),
    ('PDRB dan tenaga kerja pariwisata memakai KBLI kategori I dan R, S, T, U; jumlah ODTW dikodekan pada golongan 91 dan 93. *S, T, U ikut karena BPS menyajikan R, S, T, U dalam satu kelompok, sehingga nilai sektor pariwisata cenderung lebih besar.', 12.5, False, DARK)]], anchor=MSO_ANCHOR.MIDDLE)
source(s, 'Sumber: UN & UNWTO (2008), IRTS 2008 (daftar aktivitas karakteristik pariwisata); BPS, KBLI 2020; Buku Subbab 2.1 dan 3.4.', y=10.15)
s.notes_slide.notes_text_frame.text = ('IRTS 2008 mendefinisikan 12 aktivitas karakteristik pariwisata. Model saya memakai KBLI kategori I (akomodasi dan makan minum) '
    'serta R, S, T, U untuk PDRB dan tenaga kerja pariwisata, dan ODTW dikodekan pada golongan 91 dan 93. Angkutan, agen perjalanan, dan perdagangan '
    'eceran tidak termasuk cakupan model. S, T, U ikut terhitung karena BPS menyajikannya bersama R.')

# ---------- 5. agenda & nomor halaman
replace_in_slide(S[1], '· sensitivitas · KPI', '· sensitivitas · KPI · daftar istilah · IRTS–KBLI')
for k, s in enumerate(p.slides, 1):
    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == 6:
                yield from walk(sh.shapes)
            else:
                yield sh
    for sh in walk(s.shapes):
        if sh.has_text_frame and sh.left is not None and sh.left > 17.9 * E and sh.top > 10.4 * E and sh.text_frame.text.strip().isdigit():
            set_text(sh, str(k))
p.save('/tmp/pptwork/v5/out.pptx')
print('saved', len(p.slides))
