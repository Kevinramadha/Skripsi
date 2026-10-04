exec(open('/tmp/pptwork/helpers_v4.py').read())
from pptx import Presentation
p = Presentation('/tmp/pptwork/v5/lahan_versiAB.pptx')
# buang Versi B (slide 3) dan label versi
lst = p.slides._sldIdLst
el = list(lst)[2]; p.part.drop_rel(el.rId); lst.remove(el)
for s in p.slides:
    for sh in list(s.shapes):
        if sh.has_text_frame and sh.text_frame.text.startswith('Versi '):
            sh._element.getparent().remove(sh._element)
s = p.slides[0]
NEW = [
 ('Overall accuracy (accuracy)', 'Persentase seluruh titik yang kelasnya benar. Belum ditimbang, sehingga bisa bias karena sampel dibuat 50:50'),
 ('Overall accuracy tertimbang luas (area-weighted accuracy)', 'Accuracy yang ditimbang sesuai luas tiap kelas di peta, sehingga tidak bias. Ukuran utama (Olofsson dkk., 2014)'),
 ("User's accuracy (precision)", 'Dari titik yang dipetakan terbangun, berapa persen yang benar terbangun. Makin tinggi, makin sedikit salah tanda'),
 ("Producer's accuracy (recall)", 'Dari titik yang benar terbangun di lapangan, berapa persen yang ikut terpetakan. Makin tinggi, makin sedikit yang terlewat'),
 ('F1-score (F-measure)', 'Gabungan precision dan recall dalam satu angka. Makin mendekati 100%, makin baik'),
 ("Koefisien Kappa (Cohen's Kappa)", 'Kesesuaian peta dengan rujukan setelah dikurangi kecocokan karena kebetulan. 0,61–0,80 = kesepakatan kuat (Landis & Koch, 1977)')]
tbls = [sh for sh in s.shapes if getattr(sh, 'has_table', False) and sh.has_table]
main = max(tbls, key=lambda x: len(x.table.rows))
T = main.table
def put(cell, txt):
    para = cell.text_frame.paragraphs[0]
    para.runs[0].text = txt
    for r in para.runs[1:]: r._r.getparent().remove(r._r)
put(T.cell(0, 2), 'Cara membaca / kriteria')
for r, (m, c) in enumerate(NEW, 1):
    put(T.cell(r, 0), m); put(T.cell(r, 2), c)
    for k in (2,):
        T.cell(r, k).text_frame.paragraphs[0].runs[0].font.size = Pt(round(11 * 1.15, 1))
# lebar kolom: beri ruang untuk kalimat
tot = main.width
for k, w in enumerate([3.9, 4.3, 6.0, 1.7, 2.3]):
    T.columns[k].width = int(tot * w / 18.2)
for row in T.rows:
    row.height = Inches(0.6)
# keterangan simbol
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.startswith('n₁₁'):
        set_text(sh, "User's accuracy, producer's accuracy, dan F1-score dihitung untuk kelas terbangun · n₁₁ = terbangun benar · n₀₀ = bukan terbangun benar · n₁₀ = commission error · n₀₁ = omission error · Aᵢ = luas kelas i di peta · pₒ = kesepakatan teramati · pₑ = kesepakatan karena kebetulan")
        sh.top = Inches(7.15)
# ganti 'terbobot' di catatan pembicara
for sl in p.slides:
    nt = sl.notes_slide.notes_text_frame
    nt.text = nt.text.replace('akurasi terbobot luas', 'akurasi tertimbang luas').replace('terbobot', 'tertimbang')
p.save('/tmp/pptwork/v5/lahan_versiA.pptx')
print('ok', len(p.slides))
