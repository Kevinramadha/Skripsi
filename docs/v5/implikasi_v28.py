import sys; sys.path.insert(0, '/tmp/pptwork/v5')
from single import make_subset
make_subset('/tmp/pptwork/v5/out7.pptx', [63], '/tmp/pptwork/v5/_base_imp.pptx')
exec(open('/tmp/pptwork/helpers_v4.py').read())
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
p = Presentation('/tmp/pptwork/v5/_base_imp.pptx')
GREEN, PINK, NAVY = 'C9F0D6', 'FDE2DE', '0B5E6E'
US = '/tmp/pptwork/v5/us/'


def find_title(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.top is not None and 0.6*E <= sh.top <= 1.3*E and sh.width > 9*E and sh.text_frame.text.strip():
            return sh


def prep(s, title, resume):
    t = find_title(s)
    for sh in list(s.shapes):
        keep = (sh._element is t._element) or sh.top >= 10.3*E or (sh.left >= 18.0*E and sh.top < 1.0*E) or sh.top < 0.45*E
        if not keep:
            sh._element.getparent().remove(sh._element)
    for sh in list(s.shapes):
        if sh.top > 10.3*E and 16.5*E < sh.left < 17.9*E:
            sh._element.getparent().remove(sh._element)
    for sh in [x for x in s.shapes if x.has_text_frame and x.left > 17.9*E and x.top > 10.4*E and x.text_frame.text.strip().isdigit()]:
        set_text(sh, '')
    ps = t.text_frame.paragraphs
    ps[0].runs[0].text = title
    ps[1].runs[0].text = resume


def intinya(s, text, y=9.05, h=0.85):
    return box(s, 0.9, y, 18.2, h, fill=YEL_L, line=YEL, anchor=MSO_ANCHOR.MIDDLE,
               paras=[[('Intinya  ', 14, True, TEAL), (text, 14, False, DARK)]])


def head(s, x, y, w, text, sub=None):
    box(s, x, y, w, 0.5, fill=TEAL, line=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.15,
        paras=[[(text, 13.5, True, WHITE)] + ([('   ' + sub, 11, False, 'D7EEF2')] if sub else [])])


def qa(s, x, y, w, h, rows, size=12.5):
    """rows: list of (label, text[, color])"""
    paras = []
    for r in rows:
        paras.append([(r[0] + '  ', size, True, r[2] if len(r) > 2 else TEAL), (r[1], size, False, DARK)])
    return box(s, x, y, w, h, fill=LIGHT, line=LINE, paras=paras)


def chip(s, x, y, w, h, text, fill=LIGHT, line=LINE, size=10.5, color=DARK, bold=False):
    return box(s, x, y, w, h, fill=fill, line=line, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, margin=0.06,
               paras=[(text, size, bold, color)])


def outline(s, x, y, w, h, color=RED):
    r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.background(); r.line.color.rgb = rgb(color); r.line.width = Pt(2.5); r.shadow.inherit = False
    return r


s = p.slides[0]
prep(s, 'Implikasi Kebijakan', 'Simulasi tidak memberi satu jawaban terbaik, tetapi menunjukkan apa yang dikorbankan dari tiap pilihan')
C = [('1', 'Pemerintah harus memilih prioritas',
      'Mengejar lapangan kerja dan PDRB (Development Priority) berarti 15,64% lebih banyak lahan dibuka untuk pariwisata dan kunjungan paling padat. Sustainable tidak membuka lahan baru untuk pariwisata, dan ekonominya tetap tumbuh di atas BAU.'),
     ('2', 'Menjaga lahan tidak cukup dari sektor pariwisata saja',
      'Walaupun tidak ada lahan baru untuk pariwisata, total lahan terbangun hanya turun 0,91%, karena pembukaan lahan untuk keperluan lain tetap berjalan. Perlu pengendalian lahan lintas sektor.'),
     ('3', 'Kepadatan butuh kebijakan yang berbeda',
      'Kedua kebijakan mengatur investasi dan lahan, sedangkan kepadatan datang dari jumlah wisatawan. Perlu pengaturan sebaran kunjungan, misalnya antarwaktu dan antarlokasi.'),
     ('4', 'Dua Perda DIY perlu diselaraskan',
      'Perda 1/2012 membuka lahan baru bagi investor, sementara Perda 6/2021 melarang alih fungsi lahan pangan. Perbedaan Development Priority dan Sustainable menunjukkan tarik-menarik ini dalam angka.')]
w, h = 8.95, 2.95
for k, (n, a, b) in enumerate(C):
    x = 0.9 + (k % 2) * (w + 0.3); y = 1.95 + (k // 2) * (h + 0.25)
    box(s, x, y, w, h, fill=LIGHT, line=LINE, radius=0.04)
    num(s, x + 0.3, y + 0.3, n, d=0.65, size=16)
    tb(s, x + 1.15, y + 0.33, w - 1.4, 0.6, [[(a, 15, True, TEAL)]])
    tb(s, x + 1.15, y + 0.95, w - 1.4, 1.95, [[(b, 14, False, DARK)]])
box(s, 0.9, 8.6, 18.2, 1.15, fill=TEAL, line=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.35, paras=[
    [('Hasil simulasi ini bisa menjadi bahan diskusi bagi Pemda DIY dan Dinas Pariwisata, dengan angka yang dibaca sebagai perbandingan, bukan ramalan.', 13.5, True, WHITE)]])
source(s, 'Sumber: Buku Subbab 4.8 (Implikasi Kebijakan).', y=10.12)
s.notes_slide.notes_text_frame.text = (
    'Dari simulasi ini ada empat hal yang bisa dibawa ke pemerintah. Pertama, pemerintah harus memilih prioritas: mengejar ekonomi berarti lebih banyak lahan dibuka dan lebih padat, '
    'sedangkan Sustainable menjaga lahan tetapi ekonominya tetap tumbuh di atas BAU. Kedua, menjaga lahan tidak cukup dari sektor pariwisata saja, karena total lahan terbangun hanya turun 0,91 persen. '
    'Ketiga, kepadatan tidak bisa diatasi dengan dua kebijakan ini, karena kepadatan datang dari jumlah wisatawan; perlu pengaturan sebaran kunjungan. '
    'Keempat, dua Perda DIY perlu diselaraskan, karena yang satu membuka lahan bagi investor dan yang lain melarang alih fungsi lahan pangan. '
    'Jadi hasil ini adalah bahan diskusi kebijakan, dengan angka yang dibaca sebagai perbandingan, bukan ramalan.')
p.save('/tmp/pptwork/v5/implikasi_1slide.pptx')
print('ok')
