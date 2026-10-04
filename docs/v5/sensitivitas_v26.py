import sys; sys.path.insert(0, '/tmp/pptwork/v5')
from single import make_subset
make_subset('/tmp/pptwork/v5/out7.pptx', [55], '/tmp/pptwork/v5/_base_sens.pptx')
exec(open('/tmp/pptwork/helpers_v4.py').read())
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
p = Presentation('/tmp/pptwork/v5/_base_sens.pptx')
GREEN, PINK, NAVY = 'C9F0D6', 'FDE2DE', '0B5E6E'
US = '/tmp/pptwork/v5/us/'


def find_title(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.top is not None and 0.6*E <= sh.top <= 1.3*E and sh.width > 9*E and sh.text_frame.text.strip():
            return sh


def prep(s, title, resume):
    t = find_title(s)
    for sh in list(s.shapes):
        keep = (sh._element is t._element) or sh.top >= 10.3*E or (sh.left >= 18.0*E and sh.top < 1.0*E) or sh.top < 0.45*E
        if not keep:
            sh._element.getparent().remove(sh._element)
    for sh in list(s.shapes):
        if sh.top > 10.3*E and 16.5*E < sh.left < 17.9*E:
            sh._element.getparent().remove(sh._element)
    for sh in [x for x in s.shapes if x.has_text_frame and x.left > 17.9*E and x.top > 10.4*E and x.text_frame.text.strip().isdigit()]:
        set_text(sh, '')
    ps = t.text_frame.paragraphs
    ps[0].runs[0].text = title
    ps[1].runs[0].text = resume


def intinya(s, text, y=9.05, h=0.85):
    return box(s, 0.9, y, 18.2, h, fill=YEL_L, line=YEL, anchor=MSO_ANCHOR.MIDDLE,
               paras=[[('Intinya  ', 14, True, TEAL), (text, 14, False, DARK)]])


def head(s, x, y, w, text, sub=None):
    box(s, x, y, w, 0.5, fill=TEAL, line=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.15,
        paras=[[(text, 13.5, True, WHITE)] + ([('   ' + sub, 11, False, 'D7EEF2')] if sub else [])])


def qa(s, x, y, w, h, rows, size=12.5):
    """rows: list of (label, text[, color])"""
    paras = []
    for r in rows:
        paras.append([(r[0] + '  ', size, True, r[2] if len(r) > 2 else TEAL), (r[1], size, False, DARK)])
    return box(s, x, y, w, h, fill=LIGHT, line=LINE, paras=paras)


def chip(s, x, y, w, h, text, fill=LIGHT, line=LINE, size=10.5, color=DARK, bold=False):
    return box(s, x, y, w, h, fill=fill, line=line, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, margin=0.06,
               paras=[(text, size, bold, color)])


def outline(s, x, y, w, h, color=RED):
    r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.background(); r.line.color.rgb = rgb(color); r.line.width = Pt(2.5); r.shadow.inherit = False
    return r


from pptx.enum.chart import XL_CHART_TYPE
s = p.slides[0]
prep(s, 'Analisis Sensitivitas', 'Parameter yang paling berpengaruh justru asumsi, bukan data')
box(s, 0.9, 1.95, 18.2, 0.6, fill=LIGHT, line=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.25, paras=[[
    ('Cara pengujian:  ', 11.5, True, TEAL), ('① tiap parameter diubah ±10%, satu per satu (32 parameter)     ② parameter asumsi diuji pada seluruh rentang nilainya (17 kelompok)', 11.5, False, DARK)]])
ASU, ACU, DAT, STK = 'E8A317', '8FD3DE', TEAL, '9AA5AB'
tb(s, 0.9, 2.75, 10.3, 0.4, [[('Pengaruh terbesar terhadap jumlah wisatawan 2050 (uji ±10%, %)', 12, True, DARK)]])
rows = [('Laju penurunan dasar', 42.07, ASU), ('Kepadatan referensi', 8.61, ACU), ('Bobot kepadatan', 7.49, ASU), ('Bobot ODTW', 7.27, ASU),
        ('RDDL referensi', 5.09, ACU), ('Laju pertumbuhan eksternal', 5.06, DAT), ('ODTW referensi', 3.54, ACU), ('Sensitivitas ODTW thd investasi', 2.57, STK)]
rows = rows[::-1]
cd = CategoryChartData(); cd.categories = [r[0] for r in rows]; cd.add_series('Pengaruh maksimum (%)', [r[1] for r in rows])
ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.9), Inches(3.15), Inches(10.3), Inches(4.85), cd).chart
ch.has_title = False; ch.has_legend = False
ch.font.size = Pt(11); ch.font.name = F; ch.font.color.rgb = rgb(DARK)
pl = ch.plots[0]; pl.gap_width = 45; pl.has_data_labels = True
dl = pl.data_labels; dl.number_format = '0.00'; dl.number_format_is_linked = False; dl.show_value = True
dl.font.size = Pt(11); dl.font.bold = True; dl.position = XL_LABEL_POSITION.OUTSIDE_END
for i, r in enumerate(rows):
    pt = pl.series[0].points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(r[2])
ch.value_axis.minimum_scale = 0; ch.value_axis.maximum_scale = 50
ch.value_axis.has_major_gridlines = True; ch.value_axis.major_gridlines.format.line.color.rgb = rgb('E3EEF1')
ch.value_axis.format.line.fill.background(); ch.value_axis.tick_labels.font.size = Pt(10)
ch.category_axis.tick_labels.font.size = Pt(11); ch.category_axis.format.line.color.rgb = rgb('B8C7CC')
for k, (lab, col) in enumerate([('Asumsi peneliti', ASU), ('Nilai acuan (normalisasi)', ACU), ('Dihitung dari data', DAT), ('Persamaan stok-aliran', STK)]):
    x = 1.0 + k * 2.55
    sq = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(8.2), Inches(0.25), Inches(0.25))
    sq.fill.solid(); sq.fill.fore_color.rgb = rgb(col); sq.line.fill.background()
    tb(s, x + 0.32, 8.14, 2.2, 0.4, [[(lab, 10, False, DARK)]])

F3 = [('1', 'Yang paling berpengaruh adalah asumsi', '3 dari 4 parameter teratas berstatus asumsi. Laju penurunan dasar paling besar: 42,07%.'),
      ('2', 'Ketidakpastian terbesar: rasio laju penurunan dan pertumbuhan', 'Mengubah rasio ini (0,60–0,90) membuat hasil wisatawan 2050 bisa bergeser hingga 153,57 poin persen (69,65 poin persen bila laju pertumbuhan dihitung ulang).'),
      ('3', 'Diuji lanjut pada skenario', 'Karena asumsi sangat berpengaruh, perbandingan antarskenario nanti diuji ulang dengan menggeser asumsi-asumsi ini.')]
for k, (n, a, b) in enumerate(F3):
    y = 2.8 + k * 1.95
    fill, line = LIGHT, LINE
    box(s, 11.6, y, 7.5, 1.8, fill=fill, line=line, anchor=MSO_ANCHOR.MIDDLE, margin=0.25, paras=[
        [(a, 13, True, TEAL)], [(b, 11.5, False, DARK)]])
intinya(s, 'Angka 2050 sangat bergantung pada asumsi, jadi tidak dibaca sebagai ramalan pasti.', y=8.75, h=0.95)
source(s, 'Sumber: Buku Subbab 4.7 (Tabel 56–57, Gambar 51). Rincian rancangan pengujian di lampiran.', y=10.12)
s.notes_slide.notes_text_frame.text = (
    'Analisis sensitivitas saya lakukan dengan dua cara: mengubah tiap parameter plus-minus 10 persen satu per satu, dan menguji parameter asumsi pada seluruh rentang nilainya. '
    'Hasilnya, parameter yang paling berpengaruh justru asumsi: tiga dari empat teratas, dengan laju penurunan dasar paling besar, 42 persen. '
    'Ketidakpastian terbesar ada pada rasio laju penurunan terhadap laju pertumbuhan, yang saat kalibrasi juga tidak bisa ditentukan dari data. '
    'Karena itu angka 2050 tidak saya baca sebagai ramalan pasti, dan perbandingan antarskenario nanti saya uji ulang dengan menggeser asumsi-asumsi ini. '
    'Catatan bila ditanya: untuk lahan terbangun, yang paling berpengaruh adalah laju konversi lahan, dengan lebar rentang 115,70 poin persen.')
p.save('/tmp/pptwork/v5/sensitivitas_1slide.pptx')
print('ok')
