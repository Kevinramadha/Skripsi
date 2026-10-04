import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION

E = 914400
MED = '/tmp/pptwork/dx/word/media/'
REPO = '/home/user/Skripsi/'
TEAL, DARK, GREY, LIGHT, CYAN = '0B5E6E', '14313B', '5E6E78', 'EAF8FA', '39C0D3'
YEL_L, YEL, LINE, RED, WHITE = 'FFF1BF', 'FFD23B', 'CFE5E9', 'E4604F', 'FFFFFF'
F, FB = 'Calibri (MS)', 'Calibri (MS) Bold'

p = Presentation('/tmp/pptwork/stage1.pptx')
S = list(p.slides)


def rgb(h):
    return RGBColor.from_string(h)


SCALE = 1.15


def font(run, size, bold=False, color=DARK, italic=False):
    run.font.size = Pt(round(size * SCALE, 1))
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = FB if bold else F
    run.font.color.rgb = rgb(color)


def set_text(shape, text):
    """Replace text keeping first run formatting."""
    tf = shape.text_frame
    paras = tf.paragraphs
    first = None
    for para in paras:
        if para.runs:
            first = para
            break
    if first is None:
        tf.text = text
        return
    r0 = first.runs[0]
    for r in first.runs[1:]:
        r._r.getparent().remove(r._r)
    r0.text = text
    for para in paras:
        if para is not first:
            para._p.getparent().remove(para._p)


def all_text_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from all_text_shapes(sh.shapes)
        elif sh.has_text_frame:
            yield sh


def replace_in_slide(slide, old, new):
    n = 0
    for sh in all_text_shapes(slide.shapes):
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                if old in r.text:
                    r.text = r.text.replace(old, new)
                    n += 1
    return n


def title_shape(slide):
    best = None
    for sh in slide.shapes:
        if sh.has_text_frame and 0.8 * E < sh.top < 1.2 * E and sh.width > 9 * E:
            best = sh
    return best


def set_title(slide, text):
    set_text(title_shape(slide), text)


def strip(slide):
    t = title_shape(slide)
    for sh in list(slide.shapes):
        keep = ((t is not None and sh._element is t._element) or sh.top >= 10.3 * E or (sh.left >= 18.0 * E and sh.top < 1.0 * E)
                or sh.top < 0.4 * E)
        if not keep:
            sh._element.getparent().remove(sh._element)


def tb(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0):
    """paras: list of paragraphs; paragraph = list of (text,size,bold,color[,italic]) or a single tuple."""
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, m, Inches(margin))
    first = True
    for para in paras:
        if isinstance(para, tuple):
            para = [para]
        pp = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        pp.alignment = align
        pp.space_after = Pt(4)
        for seg in para:
            r = pp.add_run()
            r.text = seg[0]
            font(r, seg[1], seg[2], seg[3], seg[4] if len(seg) > 4 else False)
    return s


def box(slide, x, y, w, h, fill=LIGHT, line=LINE, paras=None, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, margin=0.18, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right'):
        setattr(tf, m, Inches(margin))
    for m in ('margin_top', 'margin_bottom'):
        setattr(tf, m, Inches(min(margin, 0.12)))
    if paras:
        first = True
        for para in paras:
            if isinstance(para, tuple):
                para = [para]
            pp = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            pp.alignment = align
            pp.space_after = Pt(3)
            for seg in para:
                r = pp.add_run()
                r.text = seg[0]
                font(r, seg[1], seg[2], seg[3], seg[4] if len(seg) > 4 else False)
    return s


def arrow(slide, x, y, w, h=0.35, color=CYAN):
    s = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(color)
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def num(slide, x, y, n, d=0.6, fill=TEAL, color=WHITE, size=18):
    return box(slide, x, y, d, d, fill=fill, line=None, paras=[(str(n), size, True, color)],
               align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0, shape=MSO_SHAPE.OVAL)


def table(slide, x, y, w, colw, rows, size=13, rowh=0.42, header_fill=TEAL, hl=None, bold_first_col=False):
    nr, nc = len(rows), len(rows[0])
    gs = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(rowh * nr))
    t = gs.table
    tot = sum(colw)
    for j, cw in enumerate(colw):
        t.columns[j].width = Inches(w * cw / tot)
    hl = hl or {}
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(rowh)
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.08)
            c.margin_top = c.margin_bottom = Inches(0.04)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            if i == 0:
                c.fill.fore_color.rgb = rgb(header_fill)
            elif (i, j) in hl:
                c.fill.fore_color.rgb = rgb(hl[(i, j)])
            else:
                c.fill.fore_color.rgb = rgb(WHITE if i % 2 else LIGHT)
            tf = c.text_frame
            tf.word_wrap = True
            pp = tf.paragraphs[0]
            r = pp.add_run()
            r.text = str(val)
            if i == 0:
                font(r, size, True, WHITE)
            else:
                font(r, size, bold_first_col and j == 0, DARK)
    return gs


def pic(slide, path, x, y, w=None, h=None):
    kw = {}
    if w:
        kw['width'] = Inches(w)
    if h:
        kw['height'] = Inches(h)
    return slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)


def pic_fit(slide, path, x, y, w, h):
    from PIL import Image
    iw, ih = Image.open(path).size
    sc = min(w / iw, h / ih)
    pw, ph = iw * sc, ih * sc
    return pic(slide, path, x + (w - pw) / 2, y + (h - ph) / 2, w=pw)


def takeaway(slide, text, y=9.25, h=0.95, x=0.9, w=18.2):
    return box(slide, x, y, w, h, fill=YEL_L, line=YEL,
               paras=[[('Takeaway  ', 15, True, TEAL), (text, 15, False, DARK)]], anchor=MSO_ANCHOR.MIDDLE)


def source(slide, text, y=10.12):
    tb(slide, 0.9, y, 17.0, 0.3, [(text, 10.5, False, GREY, True)])


# ------------------------------------------------------------------ navigation
SECTIONS = ['Pendahuluan', 'Metodologi', 'Tujuan 1', 'Tujuan 2', 'Tujuan 3', 'Penutup']


def section_of(i):
    if 3 <= i <= 12:
        return 0
    if 13 <= i <= 16:
        return 1
    if 17 <= i <= 23:
        return 2
    if 24 <= i <= 63:
        return 3
    if 64 <= i <= 70:
        return 4
    if 71 <= i <= 74:
        return 5
    return None


SUB = {}
for a, b, lab in [(25, 32, '2A · Integrasi citra satelit'), (33, 36, '2B · Formulasi model'),
                  (37, 39, '2C · Parameterisasi'), (40, 43, '2D · Uji struktur'),
                  (44, 52, '2E · Uji perilaku & kalibrasi'), (53, 54, '2F · Analisis sensitivitas'),
                  (55, 62, '2G · Skenario kebijakan')]:
    for k in range(a, b + 1):
        SUB[k] = lab


def update_nav(slide, idx):
    sec = section_of(idx)
    groups = [sh for sh in slide.shapes if sh.shape_type == 6 and 0.2 * E < sh.top < 0.4 * E and sh.left < 14 * E]
    if len(groups) < 6:
        return
    slots = sorted(set(round(g.left / E, 1) for g in groups))
    if len(slots) != 6:
        return
    for k, x in enumerate(slots):
        gs = [g for g in groups if round(g.left / E, 1) == x]
        active = (k == sec)
        bg = gs[0]
        for sub in bg.shapes:
            try:
                sub.fill.solid()
                sub.fill.fore_color.rgb = rgb(TEAL if active else LIGHT)
            except Exception:
                pass
        for g in gs:
            for sub in g.shapes:
                if sub.has_text_frame and sub.text_frame.text.strip():
                    set_text(sub, SECTIONS[k])
                    for para in sub.text_frame.paragraphs:
                        for r in para.runs:
                            r.font.color.rgb = rgb(WHITE if active else GREY)
                            r.font.bold = active
                            r.font.name = FB if active else F
    if idx in SUB:
        tb(slide, 13.95, 0.3, 4.3, 0.45, [(SUB[idx], 14, True, TEAL)], anchor=MSO_ANCHOR.MIDDLE)


def update_pagenum(slide, idx):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.left > 17.9 * E and sh.top > 10.4 * E and sh.text_frame.text.strip().isdigit():
            set_text(sh, str(idx))


for i, s in enumerate(S, 1):
    update_nav(s, i)
    update_pagenum(s, i)


def SL(i):
    # indeks lama (draf v1) -> posisi baru; SFD utama disisipkan di posisi 34, lampiran SFD (lama 83) dihapus
    assert i != 83
    j = i if i <= 33 else (i + 1 if i <= 82 else i)
    return S[j - 1]


# ------------------------------------------------------------------ revisions of existing slides
# 2: alur paparan
s = SL(2)
set_title(s, 'Alur paparan: tiga tujuan, tiga jawaban')
for old, new in [('Latar Belakang', 'Pendahuluan'), ('Mengapa sistem dinamis dan citra satelit', 'Masalah, celah penelitian, tujuan, ruang lingkup'),
                 ('Tujuan & Kebaruan', 'Metodologi'), ('Masalah, tujuan, ruang lingkup, kontribusi', 'Data, enam tahap penelitian, preprocessing'),
                 ('Data, tahapan, pengujian, desain skenario', 'Variabel dan struktur Causal Loop Diagram'),
                 ('Model, citra, validasi, skenario, implikasi', 'Citra satelit, model, validasi, skenario kebijakan'),
                 ('Keterbatasan, kesimpulan, saran', 'Diskusi, keterbatasan, kesimpulan, saran')]:
    replace_in_slide(s, old, new)
# card titles 3-5 (exact run text)
for sh in all_text_shapes(s.shapes):
    t = sh.text_frame.text.strip()
    if t == 'Metodologi' and sh.top > 4.4 * E and sh.left > 6 * E:
        set_text(sh, 'Tujuan 1')
    elif t == 'Hasil & Pembahasan':
        set_text(sh, 'Tujuan 2')
    elif t == 'Aplikasi' and sh.top > 4.4 * E:
        set_text(sh, 'Tujuan 3')
    elif t in ('Pendahuluan', 'Metodologi', 'Tujuan 1', 'Tujuan 2', 'Tujuan 3', 'Penutup') and sh.top > 4.4 * E:
        pass
    elif t.startswith('Lampiran:'):
        set_text(sh, 'Pola setiap tujuan: metode → hasil → makna → ringkasan hasil tujuan.  Lampiran: parameter · persamaan · data historis · citra · uji struktur · uji perilaku · kalibrasi · sensitivitas · KPI')

for sh in all_text_shapes(s.shapes):
    if sh.text_frame.text.strip() in ('Pendahuluan', 'Metodologi', 'Tujuan 1', 'Tujuan 2', 'Tujuan 3', 'Penutup') and sh.top > 4.4 * E:
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                r.font.size = Pt(19)

# 14: enam tahap + pemetaan tujuan
s = SL(14)
set_title(s, 'Enam tahap penelitian menjawab tiga tujuan')
xs = [0.9, 3.98, 7.07, 10.15, 13.24, 16.32]
maps = ['Tujuan 2', 'Tujuan 1', 'Tujuan 2', 'Tujuan 2', 'Tujuan 2', 'Tujuan 3']
cols = {'Tujuan 1': TEAL, 'Tujuan 2': CYAN, 'Tujuan 3': YEL}
for x, m in zip(xs, maps):
    box(s, x + 0.45, 5.78, 1.87, 0.4, fill=cols[m], line=None, paras=[('→ ' + m, 13, True, WHITE if m != 'Tujuan 3' else DARK)],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.05, radius=0.5)
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.startswith('Alat:'):
        set_text(sh, 'Tahap 2 → Tujuan 1  ·  Tahap 1, 3, 4, 5 → Tujuan 2 (termasuk simulasi skenario)  ·  Tahap 6 → Tujuan 3.  Alat: GEE, Vensim PLE, Python (PySD), Excel/Sheets, SDEverywhere, Next.js')

# 66: tampilan aplikasi pindah ke utama
s = SL(66)
set_title(s, 'Empat halaman aplikasi: dari input tuas ke grafik keluaran')
replace_in_slide(s, 'Lampiran', 'Tujuan 3')
tb(s, 0.9, 9.3, 18.2, 0.8, [[('Input: ', 14, True, TEAL), ('pilih skenario atau atur Insentif Kebijakan (0–0,5) dan Kebijakan Konservasi Lahan (0–1).  ', 14, False, DARK),
                              ('Proses: ', 14, True, TEAL), ('model hasil konversi SDEverywhere menghitung 2025–2050 di peramban.  ', 14, False, DARK),
                              ('Output: ', 14, True, TEAL), ('indikator numerik dan grafik deret waktu; perbandingan tiga skenario.', 14, False, DARK)]])

# 68: SUS
s = SL(68)
set_title(s, 'Usability tinggi: skor SUS rata-rata 88,50')
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.startswith('Batang merah'):
        set_text(sh, 'Batang merah = skor di bawah rata-rata normatif SUS (±68; Brooke, 2013): hanya R8. Responden: 10 mahasiswa (data kuesioner) setelah mencoba lima tugas.')
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                r.font.size = Pt(13)

# 75: lampiran definisi
replace_in_slide(SL(75), 'Latar Belakang', 'Lampiran 1')

# lampiran renumbering
LAMP = {76: 2, 77: 3, 78: 4, 80: 6, 82: 8, 84: 9, 85: 10, 86: 11, 88: 13, 89: 14, 90: 15, 91: 16}
for i, n in LAMP.items():
    t = title_shape(SL(i))
    txt = t.text_frame.text
    if txt.startswith('Lampiran') and '·' in txt:
        set_text(t, 'Lampiran %d ·%s' % (n, txt.split('·', 1)[1]))

# ------------------------------------------------------------------ dividers
DIV = {
    17: ('Tujuan 1', 'Mengidentifikasi dan merumuskan variabel-variabel utama penyusun CLD dalam model sistem dinamis kebijakan pariwisata di Provinsi DIY.',
         'Metode identifikasi  →  Evaluasi kandidat  →  Hubungan kausal  →  CLD  →  Batas model', 'Ringkasan hasil: slide 23'),
    24: ('Tujuan 2', 'Mengembangkan model sistem dinamis yang komprehensif dan memanfaatkan data citra satelit untuk memutakhirkan variabel yang mengalami jeda data, serta menguji validitas strukturnya dan kesesuaiannya terhadap data historis sebagai dasar simulasi kebijakan pariwisata.',
         '2A Citra satelit  →  2B Formulasi  →  2C Parameter  →  2D Uji struktur  →  2E Uji perilaku & kalibrasi  →  2F Sensitivitas  →  2G Skenario', 'Ringkasan hasil: slide 63'),
    63: ('Tujuan 3', 'Membangun aplikasi berbasis web yang mengimplementasikan hasil pemodelan dan simulasi skenario kebijakan sebagai alat bantu analisis dan pengambilan keputusan.',
         'Alasan & kebutuhan  →  Arsitektur  →  Fitur  →  Uji fungsional  →  Usability (SUS)', 'Ringkasan hasil: slide 70'),
}
for i, (big, quote, flow, link) in DIV.items():
    s = SL(i)
    for sh in all_text_shapes(s.shapes):
        t = sh.text_frame.text.strip()
        if t == 'Terima Kasih':
            set_text(sh, big)
        elif t.startswith('Perancangan Aplikasi'):
            set_text(sh, quote)
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    r.font.size = Pt(21 if len(quote) > 200 else 24)
        elif t.startswith('Kevin Atha'):
            set_text(sh, flow)
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    r.font.size = Pt(18)
        elif t.startswith('sistemdinamis'):
            set_text(sh, link)
    if i == 24:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.startswith('2A'):
                sh.top = Inches(7.45)
            if sh.shape_type == 6 and sh.top > 8 * E:
                sh.top = Inches(8.75)

# ------------------------------------------------------------------ new content slides
# 16 preprocessing
s = SL(16); strip(s)
set_title(s, 'Empat ketidakteraturan data ditangani sebelum pemodelan')
cards = [
    ('Patahan metode wisnus 2018→2019', 'Lonjakan 8,0 → 20,5 juta (×2,57) tanpa perubahan lapangan sepadan.',
     'Seri 2015–2018 disambung (faktor 2,2997) hanya untuk uji P5; tidak dipakai menurunkan parameter.'),
    ('Lahan terbangun 2015 kosong', 'Dynamic World baru tersedia sejak 2016.',
     'Ekstrapolasi mundur tren log-linear 2016–2025; ditandai dan tidak dipakai menurunkan parameter.'),
    ('Lonjakan akomodasi 2018', '+438 unit, hampir seluruhnya non-bintang; diduga perluasan cakupan pendataan.',
     'Dipertahankan apa adanya (tidak ada dasar koreksi); dicatat dan diperiksa pada kalibrasi Tahap 1.'),
    ('Pandemi 2020–2021', 'Guncangan eksogen yang tidak dimodelkan.',
     'Tetap disimulasikan; statistik uji dilaporkan dua versi (dengan dan tanpa 2020–2021).'),
]
for k, (h, prob, treat) in enumerate(cards):
    x = 0.9 + k * 4.62
    box(s, x, 2.55, 4.38, 6.3, fill=LIGHT)
    num(s, x + 0.25, 2.8, k + 1)
    tb(s, x + 1.0, 2.8, 3.2, 0.9, [(h, 17, True, TEAL)])
    tb(s, x + 0.3, 4.0, 3.8, 0.4, [('Masalah', 13, True, RED)])
    tb(s, x + 0.3, 4.4, 3.8, 1.6, [(prob, 15, False, DARK)])
    box(s, x + 0.25, 6.25, 3.88, 2.4, fill=WHITE, paras=[('Perlakuan', 13, True, TEAL), (treat, 15, False, DARK)])
takeaway(s, 'Perlakuan ditetapkan sebelum model dijalankan, sehingga tidak disesuaikan dengan hasil simulasi.')
source(s, 'Sumber: Buku Subbab 3.7.1, 4.5.1, 4.6.2. Nilai TK 2015–2016 dan pengeluaran 2015–2017 diisi ekstrapolasi CAGR (dokumentasi repository; Lampiran 5).')

# 18 metode CLD
s = SL(18); strip(s)
set_title(s, 'Metode: kandidat variabel disusun dari rujukan, dikontekstualisasi, lalu dievaluasi')
steps = [('Model rujukan', 'Mai & Smith (2018), Cat Ba Island: titik awal mengenali komponen dan feedback loop.'),
         ('Kontekstualisasi DIY', 'UU 10/2009, RIPPARDA DIY, Renja Dispar DIY 2025, Renstra Kemenpar 2025–2029, IPKN 2024.'),
         ('Literatur substantif', 'Dukungan mekanisme per hubungan: daya tarik, ekonomi, akomodasi, TK, kepadatan, daya dukung.'),
         ('Evaluasi struktur', 'Empat tingkat: variabel, hubungan kausal, feedback loop, batas model.')]
for k, (h, d) in enumerate(steps):
    x = 0.9 + k * 4.62
    box(s, x, 2.55, 4.1, 3.05, fill=LIGHT)
    num(s, x + 0.25, 2.75, k + 1)
    tb(s, x + 1.0, 2.8, 2.95, 0.6, [(h, 17, True, TEAL)])
    tb(s, x + 0.3, 3.6, 3.6, 1.9, [(d, 14.5, False, DARK)])
    if k < 3:
        arrow(s, x + 4.15, 3.9, 0.4)
rows = [['Tingkat', 'Aspek', 'Pertanyaan evaluasi', 'Dasar'],
        ['Variabel', 'Relevansi & definisi', 'Relevan terhadap masalah dan punya definisi operasional jelas?', 'Sterman (2000)'],
        ['Hubungan kausal', 'Mekanisme', 'Ada mekanisme yang menjelaskan perubahan variabel akibat?', 'Sterman (2000); Barlas (1996)'],
        ['Hubungan kausal', 'Polaritas', 'Tanda (+/−) konsisten dengan arah pengaruh ceteris paribus?', 'Richardson (1997)'],
        ['Hubungan kausal', 'Dukungan sumber', 'Didukung teori, bukti empiris, atau identitas model yang jelas?', 'Barlas (1996)'],
        ['Feedback loop', 'Konsistensi', 'Rangkaian menghasilkan mekanisme reinforcing/balancing yang logis?', 'Sterman (2000)'],
        ['Batas model', 'Kecukupan', 'Cukup menjelaskan masalah tanpa aspek yang tidak diperlukan?', 'Mai & Smith (2018)']]
table(s, 0.9, 5.85, 18.2, [2.3, 2.3, 8.6, 3.2], rows, size=13, rowh=0.44)
takeaway(s, 'CLD adalah hasil adaptasi + kontekstualisasi + evaluasi, bukan salinan model rujukan; ketersediaan data dipertimbangkan saat operasionalisasi.', y=9.15, h=0.85)
source(s, 'Sumber: Buku Subbab 3.7.2.1–3.7.2.2, Tabel 10–11.', y=10.08)

# 20 matriks hubungan kausal
s = SL(20); strip(s)
set_title(s, 'Hasil: 18 hubungan kausal dengan dasar yang dibedakan')
R = [['Loop', 'Hubungan kausal', 'Pol.', 'Jenis dasar'],
     ['R1', 'Jumlah Wisatawan → Laju Kedatangan', '+', 'Struktur pertumbuhan (Sterman; Mai & Smith)'],
     ['R1,R2,B1,B2', 'Daya Tarik Destinasi → Laju Kedatangan', '+', 'Literatur (Hu & Ritchie; Leiper)'],
     ['R2', 'Jumlah Wisatawan → Pengeluaran Wisatawan', '+', 'Identitas agregasi'],
     ['R2', 'Pengeluaran → PDRB Sektor Pariwisata', '+', 'Literatur TSA (Frechtling; Munjal)'],
     ['R2', 'PDRB → Investasi Sektor Pariwisata', '+', 'Formulasi berbasis data'],
     ['R2', 'Investasi → Laju Pembangunan ODTW', '+', 'Identitas stok-aliran; Fatina dkk.'],
     ['R2', 'Jumlah ODTW → Daya Tarik', '+', 'Literatur (Leiper; Hu & Ritchie)'],
     ['B1', 'Jumlah Wisatawan → Kepadatan', '+', 'Identitas model'],
     ['B1', 'Kepadatan → Daya Tarik', '−', 'Literatur (Saveriades; Mai & Smith)']]
R2_ = [['Loop', 'Hubungan kausal', 'Pol.', 'Jenis dasar'],
       ['B2', 'Konstruksi hotel & pembangunan ODTW → Konversi lahan pariwisata', '+', 'Formulasi model'],
       ['B2', 'Konversi lahan pariwisata → Lahan Terbangun', '+', 'Identitas stok-aliran'],
       ['B2', 'Lahan Terbangun → Rasio Daya Dukung Lahan', '−', 'Identitas model'],
       ['B2', 'Rasio Daya Dukung Lahan → Daya Tarik', '+', 'Literatur + operasionalisasi'],
       ['B3', 'Jumlah Hotel → Rasio Permintaan thd Kapasitas Kamar', '−', 'Identitas kapasitas'],
       ['B3', 'Rasio Permintaan → Laju Konstruksi Hotel', '+', 'Literatur (Wheaton & Rossoff)'],
       ['B4', 'PDRB → Tenaga Kerja Dibutuhkan', '+', 'Literatur (Munjal; Sánchez López)'],
       ['B4', 'Tenaga Kerja Pariwisata → Selisih TK Dibutuhkan', '−', 'Identitas penyesuaian'],
       ['B4', 'Selisih TK Dibutuhkan → Laju Penyerapan TK', '+', 'Struktur goal-seeking']]
table(s, 0.9, 2.5, 9.0, [1.3, 4.3, 0.6, 3.2], R, size=12, rowh=0.62)
table(s, 10.1, 2.5, 9.0, [0.8, 4.8, 0.6, 3.2], R2_, size=12, rowh=0.62)
takeaway(s, 'Setiap panah terlacak ke mekanisme dan sumber; dukungan konseptual tidak dipakai sebagai dasar nilai parameter. Hubungan TK → pembangunan ODTW tidak dipertahankan.', y=8.95, h=1.0)
source(s, 'Sumber: Buku Subbab 4.1.3, Tabel 25 (versi lengkap dengan mekanisme di buku).', y=10.08)

# 22 batas model
s = SL(22); strip(s)
set_title(s, 'Hasil: batas model menetapkan ranah tafsir hasil')
colsd = [('Endogen · 26 variabel', TEAL, WHITE, ['Wisatawan, kedatangan, penurunan', 'Akomodasi, konstruksi, demolisi', 'ODTW, pembangunan, penutupan', 'PDRB, pengeluaran, investasi', 'TK, penyerapan, keluar', 'Lahan terbangun & konversinya', 'Daya tarik, kepadatan, daya dukung', 'Okupansi & tekanan permintaan kamar']),
         ('Eksogen · 35 parameter', CYAN, WHITE, ['Pendorong permintaan eksternal (LPE, LPD)', 'Perilaku & karakteristik wisatawan', 'Karakteristik struktural akomodasi', 'Koefisien hasil derivasi data historis', 'Koefisien & batas kebutuhan lahan', 'Parameter & nilai referensi daya tarik', 'Pembangunan ODTW non-investasi', 'Dua tuas kebijakan']),
         ('Dikeluarkan · 15 aspek', YEL, DARK, ['Pemisahan wisman–wisnus; harga', 'Aksesibilitas; promosi & citra', 'Kualitas SDM; musim/bulanan', 'Guncangan (pandemi, erupsi, gempa)', 'Persaingan antardestinasi; zonasi RTRW', 'Dampak lingkungan non-lahan; sosial budaya', 'Distribusi manfaat; investasi APBN/APBD', 'Perpindahan TK antarsektor; harga lahan'])]
for k, (h, f, c, items) in enumerate(colsd):
    x = 0.9 + k * 6.15
    box(s, x, 2.5, 5.9, 0.7, fill=f, line=None, paras=[(h, 18, True, c)], anchor=MSO_ANCHOR.MIDDLE)
    box(s, x, 3.3, 5.9, 4.9, fill=LIGHT, paras=[('• ' + it, 14.5, False, DARK) for it in items])
box(s, 0.9, 8.35, 18.2, 0.75, fill=WHITE, paras=[[('Konsekuensi utama:  ', 14, True, RED), ('mekanisme penyeimbang lewat harga tidak tercakup; promosi & aksesibilitas tersirat di LPE; lintasan tanpa guncangan; daya dukung lahan cenderung longgar (luas administratif).', 14, False, DARK)]], anchor=MSO_ANCHOR.MIDDLE)
takeaway(s, 'Hasil model tidak boleh ditafsirkan pada ranah harga, promosi, musim, maupun guncangan eksternal.', y=9.22, h=0.8)
source(s, 'Sumber: Buku Subbab 3.7.4 (uji kecukupan batas), 4.1.2, Tabel 13–15; rincian di Lampiran 9.', y=10.08)

# 23 ringkasan T1
def ringkasan(s, title, q, items, nxt):
    strip(s)
    set_title(s, title)
    box(s, 0.9, 2.5, 18.2, 0.85, fill=TEAL, line=None, paras=[(q, 18, True, WHITE)], anchor=MSO_ANCHOR.MIDDLE)
    n = len(items)
    hh = (5.5 - 0.2 * (n - 1)) / n
    for k, (h, d) in enumerate(items):
        y = 3.6 + k * (hh + 0.2)
        box(s, 0.9, y, 18.2, hh, fill=LIGHT)
        num(s, 1.15, y + (hh - 0.6) / 2, k + 1)
        tb(s, 2.0, y + 0.1, 4.6, hh - 0.2, [(h, 16.5, True, TEAL)], anchor=MSO_ANCHOR.MIDDLE)
        tb(s, 6.7, y + 0.1, 12.2, hh - 0.2, [(d, 15, False, DARK)], anchor=MSO_ANCHOR.MIDDLE)
    takeaway(s, nxt, y=9.3, h=0.8)


ringkasan(SL(23), 'Ringkasan Hasil Tujuan 1', 'Apa hasil yang diperoleh untuk Tujuan 1?', [
    ('15 variabel kunci', 'Teridentifikasi dari 6 subsistem (permintaan & daya tarik, atraksi, ekonomi & investasi, akomodasi, lahan, tenaga kerja); seluruhnya terpetakan ke statistik resmi.'),
    ('18 hubungan kausal', 'Dievaluasi arah, mekanisme, dan jenis dukungannya; 1 hubungan (TK → pembangunan ODTW) tidak dipertahankan.'),
    ('6 feedback loop', 'R1 pertumbuhan, R2 ekonomi–atraksi, B1 kepadatan, B2 daya dukung lahan, B3 okupansi, B4 penyesuaian TK — hipotesis dinamis penelitian.'),
    ('Batas model eksplisit', '26 endogen, 35 eksogen, 15 aspek dikeluarkan beserta konsekuensinya terhadap tafsir hasil.')],
    'CLD final menjadi masukan langsung Tujuan 2: dioperasionalkan menjadi SFD 61 variabel (5 stock, 10 flow, 11 auxiliary, 35 parameter).')

# 29 NTL estimasi + validasi spasial
s = SL(29); strip(s)
set_title(s, 'NTL mengisi ODTW 2015–2017 dan 2025; validasi spasial konsisten')
rows = [['Tahun', 'NTL', 'Estimasi', 'Selang prediksi 95%'],
        ['2015', '1,262', '165,2', '138,5 – 191,9'], ['2016', '1,054', '158,1', '128,9 – 187,3'],
        ['2017', '1,438', '171,4', '146,4 – 196,3'], ['2025', '2,288', '200,7', '176,7 – 224,8']]
tb(s, 0.9, 2.45, 8.0, 0.5, [('Estimasi pada tahun tanpa data BPS', 16, True, TEAL)])
table(s, 0.9, 3.0, 8.0, [1.3, 1.3, 1.6, 3.0], rows, size=14, rowh=0.5, hl={(4, 2): YEL_L})
box(s, 0.9, 5.7, 8.0, 3.35, fill=LIGHT, paras=[
    [('• 2025 = 201 unit ', 14.5, True, TEAL), ('→ stok awal & ODTW Referensi; turun dari 218 (2024) sehingga diuji sensitivitas 189–218.', 14.5, False, DARK)],
    [('• 2016 selang terlebar (≈58 unit) ', 14.5, True, TEAL), ('karena NTL 1,054 di luar rentang data latih 1,483–2,512.', 14.5, False, DARK)],
    [('• Validasi spasial (NTL 2025, buffer 300 m): ', 14.5, True, TEAL), ('Malioboro–Keraton 37,376 · Kotagede 12,631 · Pantai Baron 1,692 · Sentolo 1,161 · TN Merapi 0,959.', 14.5, False, DARK)]])
pic_fit(s, MED + 'image14.png', 9.3, 2.45, 9.8, 6.6)
takeaway(s, 'NTL peka membedakan kawasan padat/aktif dari kawasan minim aktivitas, tetapi kurang tajam untuk aktivitas skala kecil (Sentolo vs Pantai Baron).', y=9.2, h=0.85)
source(s, 'Sumber: Buku Subbab 4.2.1, Tabel 30, Gambar 12; repository GEE.rtf (buffer 300 m, skala 500 m).', y=10.1)

# 31 DW diagnosis
s = SL(31); strip(s)
set_title(s, 'Deret Dynamic World berderau karena intensitas observasi, bukan perubahan lahan')
cd = CategoryChartData()
yrs = ['2016\n4,3', '2017\n7,0', '2018\n16,2', '2019\n22,4', '2020\n13,5', '2021\n8,5', '2022\n7,1', '2023\n16,7', '2024\n15,2', '2025\n9,3']
cd.categories = yrs
cd.add_series('Lahan terbangun DW (ha)', (44561, 50224, 56624, 63480, 58962, 59888, 57709, 75624, 73512, 55030))
gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.9), Inches(2.95), Inches(9.6), Inches(5.6), cd)
ch = gf.chart
ch.has_legend = False
ch.has_title = False
ser = ch.plots[0].series[0]
ser.format.line.color.rgb = rgb(TEAL)
ser.format.line.width = Pt(2.5)
ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = rgb(TEAL)
pl = ch.plots[0]
pl.has_data_labels = True
pl.data_labels.number_format = '#,##0'
pl.data_labels.number_format_is_linked = False
pl.data_labels.position = XL_LABEL_POSITION.ABOVE
pl.data_labels.font.size = Pt(10); pl.data_labels.font.color.rgb = rgb(DARK)
ch.category_axis.tick_labels.font.size = Pt(11); ch.category_axis.tick_labels.font.color.rgb = rgb(GREY)
ch.value_axis.tick_labels.font.size = Pt(11); ch.value_axis.tick_labels.font.color.rgb = rgb(GREY)
ch.value_axis.major_gridlines.format.line.color.rgb = rgb(LINE)
ch.value_axis.minimum_scale = 40000
tb(s, 0.9, 2.45, 9.6, 0.5, [('Luas lahan terbangun (ha) · label sumbu: tahun dan observasi per piksel', 15, True, TEAL)])
box(s, 10.8, 2.45, 8.3, 2.0, fill=YEL_L, line=YEL, paras=[('−27% dalam dua tahun', 20, True, RED), ('75.624 ha (2023) → 55.030 ha (2025): tidak mungkin secara fisik → artefak klasifikasi. Tahun dengan observasi tinggi (2019, 2023) juga tahun dengan luas tinggi.', 14, False, DARK)])
box(s, 10.8, 4.6, 8.3, 2.25, fill=LIGHT, paras=[('Diagnosis gaya ANCOVA', 16, True, TEAL), ('Luasₜ = a + b·Obsₜ;  Luas_terkoreksi = Luasₜ − b(Obsₜ − Obs_acuan)', 14, False, DARK, True),
    ('Syarat: Obs tidak berkorelasi dengan tahun (r = 0,174; p = 0,631) → terpenuhi. Luas vs Obs: r = 0,788 (p = 0,0067). Sesudah koreksi r ≈ 0; SD perubahan tahunan 12,3% → 6,6%.', 13.5, False, DARK)])
box(s, 10.8, 7.0, 8.3, 1.55, fill=LIGHT, paras=[('Validasi temporal (2015–2019)', 16, True, TEAL), ('BPS naik 40,2% lalu turun 38,8%; DLHK turun 12,9%; DW naik konsisten ±12–13%/tahun → DW tetap dipilih.', 13.5, False, DARK)])
takeaway(s, 'Nilai mentah tetap dipakai (tahun dasar 2025 menjadi acuan seluruh parameter dan kalibrasi); koreksi hanya alat diagnostik. Derau ini dibawa sebagai keterbatasan ke kalibrasi lahan.', y=8.75, h=1.0)
source(s, 'Sumber: Buku Subbab 3.5.2, 4.2.2, Gambar 13. Statistik korelasi dari notebook repository "Data Citra untuk Lahan Terbangun" (dihitung pada luas probabilitas lunak).', y=10.05)

# 32 simpulan citra
s = SL(32); strip(s)
set_title(s, 'Simpulan integrasi citra: dua variabel diperbarui, nilainya kondisional')
rows = [['Variabel', 'Citra & platform', 'Metode', 'Bukti kelayakan', 'Masuk ke model', 'Keterbatasan'],
        ['Jumlah ODTW', 'VIIRS DNB (NOAA), GEE, 500 m', 'Regresi ODTW = 121,61 + 34,59·NTL', 'R² selisih 0,405 (≥ 0,30); RMSE LOOCV 5,99% (≤ 10%)', '2015–2017: 165/158/171; 2025: 201 (stok awal & referensi)', 'Proksi tidak langsung; n = 7; selang prediksi lebar'],
        ['Lahan terbangun', 'Dynamic World (Sentinel-2), GEE, 10 m', 'Klasifikasi langsung kelas built (prob ≥ 0,5)', 'OA terbobot 89,95%; Kappa 0,61; 200 titik', '2016–2025; stok awal 55.029,54 ha; laju konversi dasar', 'Derau temporal (−27% 2023→2025); komisi dominan']]
table(s, 0.9, 2.55, 18.2, [2.0, 2.8, 3.0, 3.4, 3.6, 3.4], rows, size=13.5, rowh=1.15)
box(s, 0.9, 6.2, 8.9, 2.8, fill=LIGHT, paras=[('Prinsip yang dipegang', 17, True, TEAL), ('Citra memperbarui variabel yang sudah punya data resmi; tidak menjadi variabel endogen dan tidak menggantikan statistik resmi. Statistik resmi tetap acuan utama.', 17, False, DARK)])
box(s, 10.2, 6.2, 8.9, 2.8, fill=LIGHT, paras=[('Pelajaran metodologis', 17, True, TEAL), ('Nilai citra kondisional: variabel yang dapat diklasifikasi langsung (lahan) lebih kuat daripada variabel yang memerlukan regresi proksi (ODTW).', 17, False, DARK)])
takeaway(s, 'Unsur "memanfaatkan citra satelit untuk memutakhirkan variabel berjeda data" pada Tujuan 2 terpenuhi, dengan kelayakan yang terukur.', y=9.2, h=0.85)
source(s, 'Sumber: Buku Subbab 4.2, 4.10.2, Tabel 29–31.', y=10.1)

# 35 persamaan subsistem lain
s = SL(35); strip(s)
set_title(s, 'Persamaan akomodasi, ODTW, tenaga kerja, lahan, dan letak tuas kebijakan')
eqs = [('Akomodasi (B3)', ['Konstruksi = Investasi × 2,31385 × MAX(0; Rasio Permintaan − 0,275)', 'Rasio Permintaan = (Total Malam ÷ TPG) ÷ (Hotel × 21,4168 × 329,03)', 'TPK = MIN(1; Rasio Permintaan) · Demolisi = Hotel × 0,05']),
       ('ODTW (R2)', ['Pembangunan ODTW = Investasi × 0,01052 + 5,427', 'Penutupan ODTW = ODTW × 0,04993', 'Investasi = PDRB × 0,0488 × (1 + Insentif Kebijakan)']),
       ('Tenaga kerja (B4)', ['TK Dibutuhkan = PDRB × Intensitas TK', 'Intensitas TK = 21,5366 × e^(−0,0110(t − 2025))', 'Penyerapan = MAX(0; Keluar + (TK Dibutuhkan − TK) ÷ 1)']),
       ('Lahan (B2)', ['Konversi non-pariwisata = 0,0436052 × Lahan × RDDL', 'Konversi pariwisata = (Konstruksi × 0,1 + Pemb. ODTW × 0,5) × (1 − Konservasi Lahan) × RDDL ÷ 0,826425', 'RDDL = MAX(0; 1 − Lahan ÷ 317.036)'])]
for k, (h, ls) in enumerate(eqs):
    x = 0.9 + (k % 2) * 9.2
    y = 2.5 + (k // 2) * 3.25
    box(s, x, y, 9.0, 3.05, fill=LIGHT, paras=[(h, 18, True, TEAL)] + [('• ' + l, 16, False, DARK) for l in ls])
takeaway(s, 'Dua tuas menempel di dua jalur: Insentif Kebijakan pada investasi (R2), Kebijakan Konservasi Lahan pada konversi lahan pariwisata (B2). Pembatas MIN/MAX hasil perbaikan uji kondisi ekstrem.', y=9.15, h=0.95)
source(s, 'Sumber: Buku Subbab 4.3.1, Lampiran 6; model Vensim final (repository [04] Simulasi Skenario).', y=10.13)

# 36 kategori parameter
s = SL(36); strip(s)
set_title(s, '35 parameter dari lima kategori sumber; hanya yang tidak pasti boleh dikalibrasi')
cats = [('Data statistik resmi', '11', TEAL, WHITE, ['Pengeluaran/kunjungan 0,00272', 'Rasio nilai tambah 0,1531', 'Rasio investasi/PDRB 0,0488', 'Malam tersedia/kamar 329,03'], 'Tidak dikalibrasi'),
        ('Identitas stok-aliran', '8', CYAN, WHITE, ['LPE 0,27726 · LPD 0,207945', 'TPK ambang 0,275', 'Sens. konstruksi 2,31385', 'Laju konversi dasar 0,0436'], 'Tidak (kecuali yang tidak pasti: ambang, konversi)'),
        ('Ketentuan regulasi', '2', '3AC1D1', WHITE, ['Laju demolisi 0,05 (UU PPh)', 'Laju keluar TK 0,03 (PP 45/2015)'], 'Boleh (proksi)'),
        ('Asumsi pemodelan', '8', YEL, DARK, ['Bobot 0,40 / 0,35 / 0,25', 'Elastisitas 0,3 · batas 1,5', 'Lahan/hotel 0,1 · /ODTW 0,5', 'Waktu penyesuaian TK 1'], 'Diuji sensitivitas'),
        ('Normalisasi & tuas', '6', '8A99A1', WHITE, ['ODTW ref 201', 'Kepadatan ref 128,363', 'RDDL ref 0,826425', 'Insentif & Konservasi = 0'], 'Tidak dikalibrasi')]
for k, (h, n, f, c, items, stat) in enumerate(cats):
    x = 0.9 + k * 3.68
    box(s, x, 2.5, 3.45, 1.45, fill=f, line=None, paras=[(n, 34, True, c), (h, 15, True, c)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    box(s, x, 4.05, 3.45, 3.3, fill=LIGHT, paras=[('• ' + it, 13.5, False, DARK) for it in items])
    box(s, x, 7.45, 3.45, 1.2, fill=WHITE, paras=[('Kalibrasi?', 12.5, True, TEAL), (stat, 13, False, DARK)])
takeaway(s, 'Status sumber ditetapkan sejak awal dan menentukan parameter mana yang boleh diubah; ketidakpastian dibawa sebagai rentang ke kalibrasi dan sensitivitas.', y=8.85, h=1.0)
source(s, 'Sumber: Buku Subbab 3.4.2, 3.7.3, Tabel 12, 33–37; daftar lengkap di Lampiran 2.', y=10.05)

# 37 contoh penurunan
s = SL(37); strip(s)
set_title(s, 'Contoh penurunan parameter: dari data dan identitas stok-aliran, bukan ditebak')
der = [('Laju Pertumbuhan Eksternal & Penurunan Dasar', ['LPE = g(2024→2025) ÷ (Daya Tarik 2024 − 0,75)', '= 0,06716 ÷ (0,99223 − 0,75) = 0,27726', 'LPD = 0,75 × LPE = 0,207945', 'Euler: perubahan 2024→2025 ditentukan Daya Tarik 2024']),
       ('Sensitivitas konstruksi terhadap TPK', ['Konstruksi tersirat = ΔHotel + 0,05 × Hotel(t−1)', 'Σ 2016–2025 = 1.126 + 806,5 = 1.932,5 unit', 'Σ Investasi × MAX(0; TPK − 0,275) = 835,19', 'S = 1.932,5 ÷ 835,19 = 2,31385']),
       ('Laju penutupan ODTW (populasi stabil)', ['Porsi umur > T = e^(−(d+g)T); T = 5, 10, 20 tahun', '(d+g) = 0,0617; 0,0785; 0,0694 → rata-rata 0,0699', 'g ODTW 2015–2025 = 0,0199', 'd = 0,0699 − 0,0199 = 0,04993']),
       ('Laju konversi dasar (logistik)', ['ln Lahanₜ = a + βt (DW 2016–2025): β = 0,0348', 'k = (e^β − 1) ÷ rata-rata RDDL = (e^0,0348 − 1) ÷ 0,8121', 'k = 0,0436052', 'CI 95%: 0,0034118 – 0,0851069 (rentang kalibrasi)'])]
for k, (h, ls) in enumerate(der):
    x = 0.9 + (k % 2) * 9.2
    y = 2.5 + (k // 2) * 3.25
    box(s, x, y, 9.0, 3.05, fill=LIGHT, paras=[(h, 18, True, TEAL)] + [(l, 16, False, DARK) for l in ls])
takeaway(s, 'LPE diperoleh dari koreksi Tahap 0 (sebelumnya 0,26864 karena mengasumsikan Daya Tarik 2024 = 1); rasio 0,75 adalah asumsi struktural yang diuji pada sensitivitas.', y=9.15, h=0.95)
source(s, 'Sumber: Buku Subbab 3.4.2.2, 4.3.2.2, 4.6.1; repository Parameter Hotel, Parameter ODTW, Tahap 0 - Kalibrasi.', y=10.13)

# 45 hasil uji parsial
s = SL(45); strip(s)
set_title(s, 'Hasil uji parsial: struktur subsistem memadai, wisatawan gagal uji tren')
rows = [['Uji · variabel', 'Tren', 'E1', 'E2', 'DC', 'MAPE', 'Dominan'],
        ['P1 · Jumlah hotel', 'Sama', '0,0402', '0,0773', '0,2011', '6,05%', 'U3'],
        ['P1b · Jumlah hotel', 'Sama', '0,0182', '0,1223', '0,2319', '7,86%', 'U3'],
        ['P1 · TPK', 'Sama', '0,0643', '2,5684', '0,6692', '17,37%', 'U2'],
        ['P1b · TPK', 'Sama', '0,0610', '2,9435', '0,7053', '19,56%', 'U2'],
        ['P2 · Jumlah ODTW', 'Sama', '0,0880', '0,3303', '0,3367', '8,48%', 'U1'],
        ['P3 · Tenaga kerja', 'Sama', '0,1370', '0,3027', '0,4945', '12,88%', 'U1'],
        ['P4 · Lahan terbangun', 'Sama', '0,1125', '0,3918', '0,4813', '12,94%', 'U1, U3'],
        ['P5 · Jumlah wisatawan', 'Berbeda nyata', '—', '—', '0,4057*', '—', 'U1+U2 ≈ 0,96']]
table(s, 0.9, 2.5, 9.6, [2.8, 1.6, 0.9, 0.9, 1.0, 1.0, 1.4], rows, size=12.5, rowh=0.55,
      hl={(8, 1): 'F4D3CF', (1, 4): YEL_L})
tb(s, 0.9, 7.55, 9.6, 1.1, [('Versi tanpa 2020–2021 (2016–2025). "Sama" = tren tidak berbeda nyata. E1/E2 P5 tidak ditafsirkan karena tren berbeda (Barlas, 1989). *DC P5 dari notebook repository (tidak ditabelkan di buku).', 11.5, False, GREY, True)])
pic_fit(s, MED + 'image28.png', 10.8, 2.45, 4.1, 2.45)
pic_fit(s, MED + 'image34.png', 15.0, 2.45, 4.1, 2.45)
pic_fit(s, MED + 'image35.png', 10.8, 5.05, 4.1, 2.45)
box(s, 15.0, 5.05, 4.1, 2.45, fill=LIGHT, paras=[('Wisatawan (P5)', 15, True, RED), ('Simulasi ±6,4%/th vs data sambungan ±11,7%/th; galat sistematis.', 13.5, False, DARK)])
tb(s, 10.8, 7.55, 8.3, 0.4, [('Kiri atas: P1 hotel · kanan atas: P4 lahan · kiri bawah: P5 wisatawan', 11.5, False, GREY, True)])
takeaway(s, 'Ketika masukan dari subsistem lain diganti data aktual, tiap subsistem mereproduksi pola historisnya (MAPE seluruhnya < 20%); masalah terlokalisasi pada subsistem wisatawan.', y=8.75, h=1.0)
source(s, 'Sumber: Buku Subbab 4.5.2, Tabel 46, Gambar 25, 31, 32; repository uji_perilaku_baseline_v2.ipynb.', y=10.05)

# 50 diagnostik tahap 3
s = SL(50); strip(s)
set_title(s, 'Diagnostik Tahap 3: kesenjangan wisatawan berasal dari rezim pemulihan pascapandemi')
pic_fit(s, MED + 'image52.png', 0.9, 2.45, 8.9, 4.2)
pic_fit(s, MED + 'image53.png', 10.2, 2.45, 8.9, 4.2)
tb(s, 0.9, 6.65, 8.9, 0.4, [('Punggung keteridentifikasian: LPE 0,26–0,80 berpasangan dengan rasio 0,60–0,86', 12, False, GREY, True)])
tb(s, 10.2, 6.65, 8.9, 0.4, [('Data sambungan vs simulasi LPE 0,27726 dan LPE diagnostik 0,440 (D-A)', 12, False, GREY, True)])
rows = [['Tahun', '2017', '2018', '2019', '2020', '2021', '2022', '2023', '2024', '2025'],
        ['Data', '1,42%', '20,36%', '11,58%', '−4,43%', '16,52%', '12,72%', '18,59%', '24,86%', '6,72%'],
        ['Simulasi', '6,44%', '6,53%', '6,47%', '6,51%', '6,47%', '6,26%', '6,55%', '6,39%', '6,72%']]
table(s, 0.9, 7.1, 11.2, [1.6] + [1] * 9, rows, size=12.5, rowh=0.45, bold_first_col=True,
      hl={(1, 6): YEL_L, (1, 7): YEL_L, (1, 8): YEL_L})
box(s, 12.4, 7.1, 6.7, 1.35, fill=LIGHT, paras=[('LPE yang dituntut data: 0,410–0,490', 14.5, True, TEAL), ('Bergantung periode (D-A 0,440; D-B 0,490) dan faktor sambung (2,22–2,38).', 13, False, DARK)])
takeaway(s, 'Data hanya mengidentifikasi selisih bersih LPE dan LPD; lonjakan 2022–2024 (puncak 24,86%) adalah rezim di luar batas model (wisnus nasional 2024 +21,61%). LPE/LPD dipertahankan.', y=8.75, h=1.0)
source(s, 'Sumber: Buku Subbab 4.6.4, Tabel 53–54, Gambar 49–50.', y=10.05)

# 51 model layak
s = SL(51); strip(s)
set_title(s, 'Apakah model layak untuk simulasi skenario? Bukti dan batasnya')
ev = [('Struktur', '6 uji struktur lolos: satuan konsisten, materi kekal, 6/6 loop berfungsi, 17 uji ekstrem 0 gagal, langkah waktu memadai.', 'LOLOS', TEAL),
      ('Perilaku subsistem', 'Uji parsial: MAPE seluruh variabel < 20%; akomodasi DC 0,20; tren sama kecuali wisatawan.', 'MEMADAI', TEAL),
      ('Perilaku model penuh', 'Tidak ada variabel berkategori MAPE buruk; DC melemah karena galat merambat dari satu sumber.', 'MEMADAI, DENGAN CATATAN', '3AC1D1'),
      ('Sumber galat', 'Terlokalisasi pada pertumbuhan wisatawan → rezim pemulihan pascapandemi di luar batas model, bukan kesalahan parameter.', 'TERJELASKAN', '3AC1D1'),
      ('Nilai parameter', 'Kalibrasi 3 parameter → 0 berubah; nilai turunan data terkonfirmasi (Oliva, 2003).', 'TERKONFIRMASI', TEAL)]
for k, (h, d, st, c) in enumerate(ev):
    y = 2.5 + k * 1.1
    box(s, 0.9, y, 12.0, 0.95, fill=LIGHT, paras=[[(h + '  ', 15, True, TEAL), (d, 13.5, False, DARK)]], anchor=MSO_ANCHOR.MIDDLE)
    box(s, 13.1, y, 3.0, 0.95, fill=c, line=None, paras=[(st, 13, True, WHITE)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.05)
box(s, 16.35, 2.5, 2.75, 5.35, fill=YEL_L, line=YEL, paras=[('Batas pemakaian', 15, True, RED), ('Bukan alat ramal angka absolut', 13, False, DARK), ('Evaluasi pascakalibrasi ≠ validasi independen', 13, False, DARK), ('Tanpa guncangan eksternal', 13, False, DARK), ('Rasio LPD/LPE tidak teridentifikasi', 13, False, DARK)])
box(s, 0.9, 8.15, 18.2, 0.95, fill=TEAL, line=None, paras=[('Keputusan: model layak dipakai untuk membandingkan arah dan besaran relatif antarskenario kebijakan 2025–2050.', 17, True, WHITE)], anchor=MSO_ANCHOR.MIDDLE)
source(s, 'Sintesis dari Buku Subbab 4.4–4.6, 4.10.1, 5.1 (tidak menambah angka baru).', y=9.35)

# 62 ringkasan T2
ringkasan(SL(62), 'Ringkasan Hasil Tujuan 2', 'Apa hasil yang diperoleh untuk Tujuan 2?', [
    ('Citra satelit terintegrasi', 'VIIRS mengisi ODTW 2015–2017 & 2025 (R² selisih 0,405; LOOCV 5,99%); Dynamic World menyediakan lahan terbangun 2016–2025 (OA 89,95%; κ 0,61).'),
    ('Model SD 61 variabel', '5 stock, 10 flow, 11 auxiliary, 35 parameter dari 5 kategori sumber; nilai awal 2025 terlacak ke data resmi dan citra.'),
    ('Valid secara struktur', '6 uji struktur lolos (17 uji ekstrem tanpa gagal); perilaku memadai per subsistem; kalibrasi tidak mengubah parameter.'),
    ('Dasar simulasi kebijakan', 'Tidak ada skenario unggul di ketiga dimensi; DP unggul ekonomi, Sustainable unggul lingkungan, BAU unggul kepadatan — urutan kokoh pada 7 kondisi.')],
    'Model terintegrasi citra, tervalidasi, dan mampu membandingkan skenario kebijakan secara kuantitatif beserta kekokohannya.')

# 64 kebutuhan
s = SL(64); strip(s)
set_title(s, 'Dari model ke alat bantu keputusan: alasan dan kebutuhan pengguna')
box(s, 0.9, 2.5, 8.6, 2.6, fill=LIGHT, paras=[('Masalah', 17, True, RED), ('Hasil model biasanya berhenti di grafik statis dan laporan; menjalankan model butuh Vensim dan keahlian pemodelan, sehingga sulit dipakai pemangku kepentingan non-teknis.', 15, False, DARK)])
box(s, 0.9, 5.3, 8.6, 2.65, fill=LIGHT, paras=[('Prinsip rancangan', 17, True, TEAL), ('Model sistem dinamis tetap menjadi sumber logika perhitungan; antarmuka hanya lapisan interaksi (menerima masukan, meneruskan ke model, menampilkan keluaran). Konsep decision support system (Power, 2002).', 15, False, DARK)])
rows = [['Kebutuhan pengguna', 'Fitur yang menjawab'],
        ['Memahami dasar model dan keabsahannya', 'Halaman Model: struktur & evaluasi'],
        ['Membandingkan pilihan kebijakan', 'Halaman Skenario: BAU, Sustainable, DP + bandingkan'],
        ['Mencoba nilai kebijakan sendiri', 'Halaman Simulasi: atur 2 tuas, pilih variabel'],
        ['Akses tanpa instalasi', 'Aplikasi web (Vercel), model berjalan di peramban'],
        ['Evaluasi kualitas', 'Black-box (ISO/IEC 25010) + SUS (ISO 9241-11)']]
table(s, 9.9, 2.5, 9.2, [4.4, 4.8], rows, size=13.5, rowh=0.9)
takeaway(s, 'Aplikasi menjembatani model ilmiah dengan kebutuhan Pemda DIY, Dinas Pariwisata, dan pemangku kepentingan tanpa keahlian pemodelan.', y=8.2, h=0.9)
source(s, 'Sumber: Buku Subbab 1.1, 2.2, 3.9, 4.9.1. Kebutuhan diturunkan dari rancangan penelitian, bukan dari survei kebutuhan formal.', y=9.3)

# 67 black-box
s = SL(67); strip(s)
set_title(s, 'Pengujian fungsional black-box: 9 dari 9 kasus Pass')
rows = [['No', 'Fungsi yang diuji', 'Tindakan / masukan', 'Hasil yang diharapkan', 'Status'],
        ['1', 'Navigasi aplikasi', 'Pilih menu Beranda, Model, Skenario, Simulasi', 'Halaman sesuai menu tampil', 'Pass'],
        ['2', 'Pergantian tampilan Model', 'Pilih tab Struktur / Evaluasi Model', 'Konten sesuai tab', 'Pass'],
        ['3', 'Pemilihan skenario', 'Pilih BAU, Sustainable, atau DP', 'Skenario menjadi aktif', 'Pass'],
        ['4', 'Perbandingan skenario', 'Aktifkan Bandingkan Semua Skenario', 'Ketiga skenario tampil berdampingan', 'Pass'],
        ['5', 'Menjalankan simulasi', 'Tekan Jalankan Simulasi', 'Hasil 2025–2050 tampil', 'Pass'],
        ['6', 'Pemilihan variabel keluaran', 'Pilih / batalkan variabel', 'Grafik dan data menyesuaikan', 'Pass'],
        ['7', 'Perubahan Insentif Kebijakan', 'Ubah nilai lewat kontrol parameter', 'Nilai berubah dan dipakai simulasi', 'Pass'],
        ['8', 'Perubahan Konservasi Lahan', 'Ubah nilai lewat kontrol parameter', 'Nilai berubah dan dipakai simulasi', 'Pass'],
        ['9', 'Reset parameter', 'Tekan Reset ke BAU', 'Parameter kembali ke BAU', 'Pass']]
hl = {(r, 4): 'BDEFF3' for r in range(1, 10)}
table(s, 0.9, 2.5, 18.2, [0.6, 3.6, 5.2, 5.0, 1.2], rows, size=13.5, rowh=0.6, hl=hl)
takeaway(s, 'Fungsi utama berjalan sesuai spesifikasi (100% kasus uji). Pengujian terbatas pada kasus yang ditetapkan dan tidak menilai kemudahan penggunaan — itu dinilai lewat SUS.', y=8.75, h=1.0)
source(s, 'Sumber: Buku Subbab 3.9.2, 4.9.2, Tabel 22 dan 64; repository Evaluasi SUS/BlackBox.xlsx. Metode: black-box testing (ISTQB, 2019).', y=10.05)

# 69 ringkasan T3
ringkasan(SL(69), 'Ringkasan Hasil Tujuan 3', 'Apa hasil yang diperoleh untuk Tujuan 3?', [
    ('Aplikasi web DST', 'Model Vensim final dikonversi SDEverywhere dan dijalankan di peramban (Next.js, TypeScript, Tailwind): sistemdinamispariwisata.vercel.app.'),
    ('Empat halaman', 'Beranda, Model (struktur & evaluasi), Skenario (tiga skenario + perbandingan), Simulasi (dua tuas, pilihan variabel, 2025–2050).'),
    ('Fungsional', '9 dari 9 kasus black-box Pass.'),
    ('Usability', 'SUS rata-rata 88,50 (median 92,50; n = 10) — antara Excellent dan Best Imaginable; masukan: perlu bantuan interpretasi hasil.')],
    'Hasil pemodelan dan simulasi skenario kebijakan kini dapat dieksplorasi tanpa perangkat lunak pemodelan.')

# ------------------------------------------------------------------ new lampiran
def lamp(i, n, title):
    s = SL(i)
    for sh in list(s.shapes):
        if sh.shape_type == 19:
            sh._element.getparent().remove(sh._element)
    set_title(s, 'Lampiran %d · %s' % (n, title))
    return s


s = lamp(79, 5, 'Data pendukung dari dokumentasi repository')
rows = [['Tahun', 'TK pariwisata (jiwa)', 'Status TK', 'Pengeluaran/kunjungan (Rp)', 'Status pengeluaran'],
        ['2015', '254.537', 'Ekstrapolasi CAGR 3,67%', '1.023.749', 'Ekstrapolasi CAGR 10,27%'],
        ['2016', '263.878', 'Ekstrapolasi CAGR 3,67%', '1.128.845', 'Ekstrapolasi CAGR 10,27%'],
        ['2017', '273.563', 'Sakernas', '1.244.730', 'Ekstrapolasi CAGR 10,27%'],
        ['2018', '354.684', 'Sakernas', '1.372.512', 'Observasi (tertimbang)'],
        ['2025', '364.994', 'Sakernas', '2.720.209', 'Observasi (tertimbang)']]
table(s, 0.9, 2.5, 11.0, [1.1, 2.3, 3.0, 2.8, 3.2], rows, size=13, rowh=0.55)
box(s, 0.9, 6.0, 11.0, 3.0, fill=LIGHT, paras=[('Catatan', 15, True, TEAL),
    ('• TK 2015–2016: NTL (R² selisih 0,070) dan PDRB (R² selisih 0,093) ditolak sebagai estimator; dipakai CAGR 2017–2025.', 13, False, DARK),
    ('• Pengeluaran per kunjungan = rata-rata tertimbang wisnus (DIY, ribu Rp) dan wisman (nasional, US$ × kurs tengah BI).', 13, False, DARK),
    ('• Faktor sambung wisnus: g_pra 7,63%; g_pasca 15,52%; g_wajar 11,58%; lompatan 2,566 → 2,2997.', 13, False, DARK)])
box(s, 12.2, 2.5, 6.9, 6.5, fill=LIGHT, paras=[('Spesifikasi Google Earth Engine', 15, True, TEAL),
    ('NTL: NOAA/VIIRS/DNB/MONTHLY_V1/VCMSLCFG, band avg_rad, rata-rata bulanan → tahunan, reduceRegions mean, skala 500 m.', 13, False, DARK),
    ('Batas wilayah: FAO/GAUL/2015/level1 (DIY).', 13, False, DARK),
    ('Dynamic World: GOOGLE/DYNAMICWORLD/V1, band built dirata-ratakan per tahun, ambang ≥ 0,5, pixelArea, skala 10 m.', 13, False, DARK),
    ('Uji akurasi: stratifiedSample 100 titik/kelas, seed 2024, tahun 2025.', 13, False, DARK),
    ('Validasi spasial NTL: titik sampel dengan buffer 300 m.', 13, False, DARK)])
source(s, 'Sumber: repository (Ekstrapolasi Tenaga Kerja Pariwisata, Perhitungan Pengeluaran wisatawan, Master Data, GEE.rtf). Tidak mengubah nilai final di buku.', y=9.4)

s = lamp(81, 7, 'KPI pelengkap per skenario tahun 2050')
rows = [['Indikator 2050', 'BAU', 'Sustainable', 'Development Priority', 'Sumber'],
        ['Tingkat penghunian kamar (deskriptif)', '0,3584', '0,3537', '0,3458', 'Repository'],
        ['Jumlah hotel dan akomodasi (unit)', '5.404', '5.631', '6.033', 'Repository'],
        ['Jumlah ODTW (unit)', '367', '394', '449', 'Repository'],
        ['Investasi kumulatif 2025–2050 (miliar Rp)', '37.318', '41.550', '50.133', 'Buku Tabel 60'],
        ['Tahun ambang kepadatan 2,0 terlampaui', '2043', '2042', '2041', 'Buku 4.8.3'],
        ['Rasio daya dukung lahan (ambang > 0,331)', '0,6145', '0,6180', '0,6140', 'Buku Tabel 60']]
table(s, 0.9, 2.5, 18.2, [6.0, 2.4, 2.6, 3.4, 3.0], rows, size=14, rowh=0.62)
tb(s, 0.9, 7.2, 18.2, 1.0, [('TPK tercantum sebagai KPI deskriptif pada Tabel 21 buku tetapi tidak muncul di Tabel 60; nilainya diambil dari hasil_simulasi_skenario.xlsx (23 run: 3 skenario × 7 kondisi + 2 dekomposisi).', 13, False, GREY, True)])

s = lamp(87, 12, 'Horizon diperpanjang sampai 2150 (uji kestabilan, BAU)')
rows = [['Tahun', 'Wisatawan (juta)', 'Daya tarik', 'RDDL', 'Lahan (ha)', 'Hotel (unit)', 'ODTW', 'TK (ribu)'],
        ['2050', '96,20', '0,813', '0,615', '122.208', '5.404', '367', '650,4'],
        ['2075', '122,23', '0,758', '0,346', '207.210', '7.191', '534', '637,1'],
        ['2100', '113,79', '0,725', '0,149', '269.864', '6.910', '597', '454,6'],
        ['2125', '91,30', '0,715', '0,054', '299.837', '5.608', '556', '277,7'],
        ['2150', '72,65', '0,721', '0,018', '311.190', '4.444', '480', '167,5']]
table(s, 0.9, 2.5, 18.2, [1.3, 2.3, 1.7, 1.4, 2.0, 1.9, 1.4, 1.7], rows, size=14, rowh=0.6)
box(s, 0.9, 6.4, 18.2, 1.9, fill=LIGHT, paras=[('Bacaan', 15, True, TEAL), ('Wisatawan memuncak ±122,9 juta pada 2080 lalu turun ketika RDDL mendekati nol: pola overshoot–penurunan khas struktur daya dukung (limits to growth). Simulasi tetap stabil (TPK maks 0,388; daya tarik min 0,715). Uji ini dirujuk di Buku 4.3.2.4 sebagai dasar pemilihan bobot dan elastisitas; tabelnya hanya ada di repository.', 13.5, False, DARK)])
source(s, 'Sumber: repository hasil_uji_kondisi_ekstrem_v2.xlsx (sheet Horizon 2150). Bukan hasil utama penelitian; untuk tanya jawab.', y=8.6)

p.save('/tmp/pptwork/stage2.pptx')
print('saved', len(p.slides))

# ------------------------------------------------------------------ speaker notes
NOTES = {
 2: 'Paparan mengikuti tiga tujuan penelitian, bukan urutan bab. Setelah pendahuluan dan fondasi metodologi (data, enam tahap, preprocessing), setiap tujuan dibahas dengan pola metode, hasil, makna, lalu ditutup slide ringkasan hasil tujuan. Lampiran disiapkan untuk tanya jawab.',
 16: 'Sebelum pemodelan ada empat ketidakteraturan data. Patahan metode wisnus 2018 ke 2019 disambung dengan faktor 2,2997 hanya untuk uji P5. Lahan 2015 diekstrapolasi dan tidak dipakai menurunkan parameter. Lonjakan akomodasi 2018 dipertahankan karena tidak ada dasar koreksi. Tahun pandemi tetap disimulasikan dan statistik dilaporkan dua versi. Semua perlakuan ditetapkan sebelum model dijalankan.',
 17: 'Tujuan 1: mengidentifikasi dan merumuskan variabel penyusun CLD. Alurnya metode identifikasi, evaluasi kandidat, hubungan kausal, CLD, dan batas model.',
 18: 'Variabel tidak diambil mentah dari Mai dan Smith. Model rujukan hanya titik awal, lalu dikontekstualisasi dengan UU 10/2009, RIPPARDA, Renja Dispar, Renstra Kemenpar, dan IPKN, didukung literatur per mekanisme, lalu dievaluasi dengan enam kriteria pada empat tingkat.',
 20: 'Delapan belas hubungan kausal dievaluasi. Yang penting: dasarnya dibedakan. Ada yang didukung literatur substantif, ada yang berupa identitas atau operasionalisasi model. Dukungan konseptual tidak dipakai sebagai dasar nilai parameter. Hubungan tenaga kerja ke pembangunan ODTW dibuang.',
 22: 'Batas model menetapkan apa yang boleh dan tidak boleh ditafsirkan: 26 variabel endogen, 35 parameter eksogen, dan 15 aspek yang dikeluarkan, misalnya harga, promosi, musim, dan guncangan. Konsekuensinya dicatat agar hasil tidak dibaca di luar ranahnya.',
 23: 'Jawaban Tujuan 1: 15 variabel kunci dari enam subsistem, 18 hubungan kausal, enam loop, dan batas model eksplisit. CLD ini menjadi masukan langsung Tujuan 2.',
 24: 'Tujuan 2 adalah bagian terbesar: integrasi citra satelit, formulasi, parameterisasi, uji struktur, uji perilaku dan kalibrasi, sensitivitas, lalu simulasi skenario kebijakan sebagai pemakaian model.',
 29: 'Model regresi dipakai mengisi tahun tanpa data BPS, disertai selang prediksi 95 persen. Nilai 2025 sebesar 201 unit menjadi stok awal, dan karena turun dari 218 nilainya diuji sensitivitas 189 sampai 218. Validasi spasial menunjukkan NTL tinggi di Malioboro dan Kotagede, rendah di TN Merapi.',
 31: 'Deret Dynamic World turun 27 persen dari 2023 ke 2025, yang tidak mungkin secara fisik. Diagnosis menunjukkan luas mengikuti jumlah observasi per piksel, bukan tahun. Nilai mentah tetap dipakai karena tahun dasar 2025 sudah menjadi acuan parameter dan kalibrasi; koreksi hanya alat diagnostik.',
 32: 'Simpulan bagian citra: dua variabel diperbarui dengan bukti kelayakan terukur. Pelajarannya kondisional: klasifikasi langsung seperti lahan lebih kuat daripada proksi regresi seperti ODTW.',
 35: 'Persamaan subsistem lain. Perhatikan letak dua tuas: insentif pada investasi di jalur R2, konservasi pada konversi lahan pariwisata di jalur B2. Pembatas MIN dan MAX ditambahkan setelah uji kondisi ekstrem.',
 36: 'Tiga puluh lima parameter dikelompokkan menurut sumber. Pengelompokan ini menentukan apa yang boleh dikalibrasi: data resmi, identitas stok-aliran, dan nilai referensi tidak dikalibrasi.',
 37: 'Contoh penurunan. LPE diturunkan dari pertumbuhan 2024 ke 2025 dengan indeks daya tarik 2024 sesuai metode Euler. Sensitivitas konstruksi dari identitas stok-aliran, laju penutupan ODTW dari teori populasi stabil, laju konversi dari tren log-linear beserta selang kepercayaannya.',
 45: 'Uji parsial mengganti masukan dari subsistem lain dengan data aktual. Hasilnya, akomodasi paling baik, semua MAPE di bawah 20 persen. Satu-satunya yang gagal uji tren adalah subsistem wisatawan: simulasi sekitar 6,4 persen per tahun, data sambungan sekitar 11,7 persen.',
 50: 'Tahap 3 bukan mencari nilai baru, tetapi mendiagnosis. Data hanya mengidentifikasi selisih bersih LPE dan LPD, terlihat dari punggung himpunan indiferen. Kesenjangan berasal dari lonjakan pemulihan 2022 sampai 2024, rezim di luar batas model.',
 51: 'Kesimpulan bagian pengujian: model lolos uji struktur, memadai per subsistem, galat terlokalisasi dan terjelaskan, parameter terkonfirmasi kalibrasi. Karena itu model layak untuk membandingkan skenario, bukan untuk meramal angka absolut.',
 62: 'Jawaban Tujuan 2: citra satelit terintegrasi, model 61 variabel terparameterisasi, valid secara struktur, dan sebagai dasar simulasi kebijakan menunjukkan tidak ada skenario unggul di semua dimensi dengan urutan yang kokoh.',
 63: 'Tujuan 3: membangun aplikasi web yang mengimplementasikan model dan skenario sebagai alat bantu keputusan.',
 64: 'Alasan aplikasi: hasil model biasanya statis dan butuh Vensim. Aplikasi menjaga model sebagai sumber logika, sementara antarmuka hanya lapisan interaksi. Kebutuhan pengguna dipetakan ke fitur.',
 67: 'Sembilan fungsi utama diuji black-box dan semuanya Pass. Ini membuktikan fungsi berjalan sesuai spesifikasi, belum menilai kemudahan penggunaan.',
 69: 'Jawaban Tujuan 3: aplikasi web berjalan di peramban, empat halaman, 9 dari 9 fungsi Pass, dan SUS 88,50.',
 79: 'Cadangan: data yang dilengkapi ekstrapolasi CAGR dan spesifikasi teknis Google Earth Engine, dari dokumentasi repository.',
 81: 'Cadangan: KPI pelengkap per skenario, termasuk TPK yang tidak ada di Tabel 60 buku.',
 87: 'Cadangan: horizon diperpanjang sampai 2150 menunjukkan pola puncak lalu turun ketika daya dukung lahan menipis, dan model tetap stabil.',
}
for i, t in NOTES.items():
    ns = SL(i).notes_slide
    ns.notes_text_frame.text = t
p.save('/tmp/pptwork/stage2.pptx')
print('notes done')

# ------------------------------------------------------------------ revisi v2
# slide 21: CLD + tabel loop
s = SL(21); strip(s)
set_title(s, 'Hasil: CLD final dengan dua loop penguat dan empat loop penyeimbang')
pic_fit(s, MED + 'image11.png', 0.9, 2.45, 7.6, 6.55)
rows = [['Loop', 'Jenis', 'Rantai kausal ringkas', 'Peran dalam model'],
        ['R1', 'Penguat', 'Wisatawan (+) → Laju Kedatangan (+) → Wisatawan', 'Pertumbuhan kunjungan bersifat akumulatif'],
        ['R2', 'Penguat', 'Wisatawan → Pengeluaran → PDRB → Investasi → Pembangunan ODTW → ODTW → Daya Tarik → Kedatangan', 'Jalur ekonomi–atraksi; tempat tuas Insentif Kebijakan bekerja'],
        ['B1', 'Penyeimbang', 'Wisatawan (+) → Kepadatan (−) → Daya Tarik → Kedatangan', 'Rem kepadatan (crowding)'],
        ['B2', 'Penyeimbang', 'Konstruksi hotel & ODTW (+) → Konversi lahan → Lahan Terbangun (−) → RDDL (+) → Daya Tarik', 'Rem daya dukung lahan; tempat tuas Konservasi Lahan bekerja'],
        ['B3', 'Penyeimbang', 'Hotel (−) → Rasio Permintaan thd Kapasitas Kamar (+) → Laju Konstruksi → Hotel', 'Penyesuaian kapasitas akomodasi'],
        ['B4', 'Penyeimbang', 'TK Pariwisata (−) → Selisih TK Dibutuhkan (+) → Laju Penyerapan → TK', 'Goal-seeking tenaga kerja; loop lokal']]
hl = {}
for r in range(1, 7):
    hl[(r, 0)] = 'BDEFF3' if r <= 2 else YEL_L
    hl[(r, 1)] = 'BDEFF3' if r <= 2 else YEL_L
table(s, 8.8, 2.45, 10.3, [0.9, 1.7, 4.6, 3.1], rows, size=13, rowh=0.93, hl=hl, bold_first_col=True)
takeaway(s, 'Pertumbuhan didorong R1–R2 dan ditahan B1–B2; dua tuas kebijakan menempel pada R2 (insentif) dan B2 (konservasi). Hubungan TK → pembangunan ODTW tidak dipertahankan.', y=9.15, h=0.9)
source(s, 'Sumber: Buku Subbab 4.1.3, Tabel 26, Gambar CLD.', y=10.12)

# slide 34: SFD utama
s = S[33]; strip(s)
set_title(s, 'Hasil: stock-flow diagram pariwisata DIY')
pic_fit(s, MED + 'image12.png', 0.9, 2.4, 13.6, 6.65)
box(s, 14.75, 2.4, 4.35, 3.0, fill=LIGHT, paras=[('5 stock (kotak)', 16, True, TEAL),
    ('Jumlah Wisatawan · Hotel & Akomodasi · ODTW · Tenaga Kerja · Lahan Terbangun', 13.5, False, DARK)])
box(s, 14.75, 5.55, 4.35, 1.75, fill=LIGHT, paras=[('10 flow · 11 auxiliary · 35 parameter', 15, True, TEAL),
    ('26 endogen + 35 eksogen = 61 variabel', 13.5, False, DARK)])
box(s, 14.75, 7.45, 4.35, 1.6, fill=YEL_L, line=YEL, paras=[('Dua tuas kebijakan', 15, True, TEAL),
    ('Insentif → investasi (R2); Konservasi → konversi lahan pariwisata (B2)', 13.5, False, DARK)])
takeaway(s, 'SFD adalah operasionalisasi CLD menjadi model kuantitatif yang disimulasikan 2025–2050 (Δt 1 tahun, Euler). Persamaan lengkap di Lampiran 15.', y=9.2, h=0.85)
source(s, 'Sumber: Buku Subbab 4.1.4, Gambar 10; model Vensim final.', y=10.12)
SPK = s.notes_slide
SPK.notes_text_frame.text = 'Ini hasil utama formulasi Tujuan 2: CLD dioperasionalkan menjadi stock-flow diagram. Lima stok sebagai kotak, aliran sebagai katup, dan variabel bantu yang menghubungkannya. Tunjukkan letak dua tuas: insentif pada jalur investasi, konservasi pada konversi lahan pariwisata.'

# penanda prioritas W/S
W = {6, 8, 9, 14, 18, 19, 21, 23, 25, 28, 30, 32, 34, 35, 37, 42, 43, 46, 47, 48, 50, 52, 54, 56, 57, 59, 61, 62, 63, 66, 67, 69, 70, 73}
for i in range(2, 75):
    pr = 'W' if i in W else 'S'
    sl = S[i - 1]
    dark = i in (17, 24, 64)
    x, y = (18.55, 10.45) if dark else (17.55, 10.55)
    c = box(sl, x, y, 0.46, 0.46, fill=RED if pr == 'W' else CYAN, line=None,
            paras=[(pr, 14, True, WHITE)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0, shape=MSO_SHAPE.OVAL)
    c.name = 'Penanda prioritas ' + pr

# legenda di slide 2
s = S[1]
box(s, 0.9, 9.35, 0.46, 0.46, fill=RED, line=None, paras=[('W', 14, True, WHITE)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0, shape=MSO_SHAPE.OVAL)
tb(s, 1.5, 9.38, 6.5, 0.45, [('Wajib dijelaskan — inti argumen', 15, False, DARK)])
box(s, 7.2, 9.35, 0.46, 0.46, fill=CYAN, line=None, paras=[('S', 14, True, WHITE)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0, shape=MSO_SHAPE.OVAL)
tb(s, 7.8, 9.38, 6.5, 0.45, [('Sebut singkat — cukup takeaway (15–30 detik)', 15, False, DARK)])
tb(s, 13.6, 9.38, 5.5, 0.45, [('Penanda ada di pojok kanan bawah setiap slide', 13, False, GREY, True)])

p.save('/tmp/pptwork/stage2.pptx')
print('v2 done', len(p.slides))

# ------------------------------------------------------------------ v3: label tahapan / subbab di atas judul
import copy as _copy
from pptx.oxml.ns import qn
STAGE = {
 3: '1.1 · Latar Belakang', 4: '1.1 · Latar Belakang', 5: '1.1 · Latar Belakang',
 6: '1.1 · Latar Belakang — Permasalahan', 7: '1.1 · Latar Belakang — Pendekatan Sistem Dinamis',
 8: '1.1 · Latar Belakang — Celah Penelitian', 9: '1.2–1.3 · Identifikasi Masalah dan Tujuan Penelitian',
 10: '1.2 · Batasan Masalah', 11: '2.3 · Kontribusi Penelitian Terkait', 12: '3.3 · Kerangka Pikir Penelitian',
 13: '3.2 & 3.4 · Lokasi Penelitian, Data dan Sumber Data', 14: '3.1 & 3.7 · Tahapan Penelitian',
 15: '3.6 · Alat dan Perangkat Lunak', 16: '3.7.1 · Pengumpulan dan Preprocessing Data',
 18: '3.7.2 · Identifikasi dan Perumusan Variabel CLD', 19: '4.1.1 · Hasil Identifikasi dan Evaluasi Variabel CLD',
 20: '4.1.3 · Perumusan dan Evaluasi Hubungan Kausal', 21: '4.1.3 · Causal Loop Diagram dan Struktur Feedback Loop',
 22: '4.1.2 · Penetapan Batas Model',
 25: '3.5 & 4.2 · Integrasi Data Citra Satelit', 26: '3.5.1 · Estimasi Jumlah ODTW dengan Night-Time Light',
 27: '3.5 · Validasi dan Evaluasi Kelayakan Estimator', 28: '4.2.1 · Estimasi Jumlah ODTW',
 29: '4.2.1 · Estimasi ODTW — Validasi Temporal dan Spasial', 30: '4.2.2 · Estimasi Lahan Terbangun — Uji Akurasi',
 31: '4.2.2 · Diagnosis Ketidakstabilan Temporal', 32: '4.2 · Simpulan Integrasi Data Citra Satelit',
 33: '4.1.4 · Konversi ke Stock-Flow Diagram', 34: '4.1.4 · Stock-Flow Diagram',
 35: '4.3.1 · Persamaan Stock — Subsistem Wisatawan', 36: '4.3.1 · Persamaan Stock — Subsistem Lainnya',
 37: '3.4.2 & 4.3.2 · Penetapan Parameter', 38: '4.3.2 · Penurunan Nilai Parameter', 39: '4.3.2.4 · Parameter Asumsi Pemodelan',
 40: '3.7.4 · Rancangan Uji Struktur', 41: '4.4.1–4.4.3 · Konsistensi Dimensi, Batas Model, Kekekalan Materi',
 42: '4.4.4 · Uji Loop Umpan Balik', 43: '4.4.5–4.4.6 · Uji Kondisi Ekstrem dan Galat Integrasi',
 44: '3.7.4 · Rancangan Uji Perilaku', 45: '4.5.1 · Rancangan Uji Perilaku — Penyambungan Data',
 46: '4.5.2 · Hasil Uji Parsial', 47: '4.5.3 · Hasil Uji Penuh', 48: '4.5.3 · Hasil Uji Penuh — Sumber Galat',
 49: '3.7.4 · Protokol Kalibrasi', 50: '4.6.1–4.6.5 · Hasil Kalibrasi Tahap 0–4', 51: '4.6.4 · Kalibrasi Diagnostik Subsistem Wisatawan',
 52: '4.4–4.6 · Sintesis Kelayakan Model', 53: '3.7.5 · Rancangan Analisis Sensitivitas', 54: '4.7 · Hasil Analisis Sensitivitas',
 55: '3.8.1 · Kriteria Penetapan Tuas Kebijakan', 56: '4.8.1–4.8.2 · Rancangan Skenario dan Indikator Kinerja',
 57: '4.8.3 · Hasil Perbandingan Skenario', 58: '4.8.3 · Hasil Perbandingan Skenario — Komposisi Daya Tarik',
 59: '4.8.3 · Hasil Perbandingan Skenario — Ambang Kepadatan', 60: '4.8.4 · Dekomposisi Kontribusi Tuas',
 61: '4.8.5 · Kekokohan Peringkat Skenario', 62: '4.8.6 · Implikasi Kebijakan',
 65: '3.9 & 4.9.1 · Implementasi Aplikasi Berbasis Web', 66: '3.9.1 & 4.9.1 · Arsitektur dan Implementasi Aplikasi',
 67: '4.9.1 · Implementasi Aplikasi — Halaman Aplikasi', 68: '4.9.2 · Hasil Pengujian Fungsional (Black-box)',
 69: '4.9.3 · Hasil Evaluasi Usability (SUS)',
 71: '4.10.2 · Diskusi', 72: '4.10.1 · Catatan Keterbatasan', 73: '5.1 · Kesimpulan', 74: '5.2 · Saran',
}
STAGE_COLOR = '1C8FA0'
for i, label in STAGE.items():
    sl = S[i - 1]
    t = title_shape(sl)
    txBody = t.text_frame._txBody
    first = t.text_frame.paragraphs[0]._p
    for para in t.text_frame.paragraphs:
        pPr = para._p.find(qn('a:pPr'))
        if pPr is not None:
            ln = pPr.find(qn('a:lnSpc'))
            if ln is not None:
                ln.getparent().remove(ln)
        para.line_spacing = 1.0
        for r in para.runs:
            r.font.size = Pt(35)
    newp = _copy.deepcopy(first)
    for r in list(newp.findall(qn('a:r')))[1:]:
        newp.remove(r)
    first.addprevious(newp)
    pp = t.text_frame.paragraphs[0]
    pp.runs[0].text = label
    pp.runs[0].font.size = Pt(19)
    pp.runs[0].font.bold = True
    pp.runs[0].font.name = FB
    pp.runs[0].font.color.rgb = rgb(STAGE_COLOR)
    pp.space_after = Pt(2)
    t.top = Inches(0.8)

p.save('/tmp/pptwork/stage2.pptx')
print('v3 stage labels', len(STAGE))
