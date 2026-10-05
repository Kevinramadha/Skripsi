import copy, re
exec(open('/tmp/pptwork/helpers_v4.py').read())
SCALE = 1.0
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

MED = '/tmp/pptwork/dx/word/media/'
REPO = '/home/user/Skripsi/'
GREEN, PINK, GRY = 'C9F0D6', 'FDE2DE', '8A979E'
KEEP_IDS = {2, 4, 6, 8, 10, 13, 14, 16, 18, 20, 22, 25}
R_NS = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'


def fmt(x, d=2):
    """Angka gaya Indonesia: titik ribuan, koma desimal."""
    if x is None or x == '':
        return '—'
    if isinstance(x, str):
        return x
    s = f'{x:,.{d}f}'
    return s.replace(',', '§').replace('.', ',').replace('§', '.')


def pct(x, d=2, sign=False):
    s = fmt(x, d) + '%'
    if sign and x > 0:
        s = '+' + s
    return s.replace('-', '−')


class Deck:
    def __init__(self, prs, tmpl):
        self.prs, self.tmpl, self.new = prs, tmpl, []

    def slide(self, title, size=36):
        prs, tmpl = self.prs, self.tmpl
        s = prs.slides.add_slide(tmpl.slide_layout)
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        tree = s.shapes._spTree
        for sh in tmpl.shapes:
            if sh.shape_id not in KEEP_IDS:
                continue
            el = copy.deepcopy(sh._element)
            for node in el.iter():
                for att in ('embed', 'link', 'id'):
                    k = R_NS + att
                    if k in node.attrib:
                        rel = tmpl.part.rels[node.attrib[k]]
                        node.attrib[k] = s.part.relate_to(rel._target, rel.reltype)
            tree.append(el)
        for sh in s.shapes:
            if sh.shape_id == 13 or (sh.has_text_frame and sh.text_frame.text.startswith('Lampiran 18')):
                r = sh.text_frame.paragraphs[0].runs[0]
                r.text = title
                r.font.size = Pt(size if len(title) <= 52 else (30 if len(title) <= 64 else 26))
            if sh.has_text_frame and sh.text_frame.text.strip() == '94':
                sh.text_frame.paragraphs[0].runs[0].text = ''
        self.new.append(s)
        return s


# ---------- komponen ----------
def T(s, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    return tb(s, x, y, w, h, paras, align=align, anchor=anchor)


def sub(s, text, y=1.55):
    """Subjudul/resume di bawah judul."""
    return tb(s, 0.9, y, 18.2, 0.45, [[(text, 15, False, '4A5A63')]])


def src(s, text, y=10.12):
    tb(s, 0.9, y, 17.0, 0.3, [(text, 10.5, False, GREY, True)])


def hdr(s, x, y, w, text, sub_=None, fill=TEAL, h=0.48):
    box(s, x, y, w, h, fill=fill, line=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.15,
        paras=[[(text, 13, True, WHITE)] + ([('   ' + sub_, 10.5, False, 'D7EEF2')] if sub_ else [])])


def card(s, x, y, w, h, title, lines, fill=LIGHT, line=LINE, tcol=TEAL, size=11.5, tsize=13):
    paras = []
    if title:
        paras.append([(title, tsize, True, tcol)])
    for ln in lines:
        if isinstance(ln, tuple) and len(ln) == 2 and isinstance(ln[0], str) and isinstance(ln[1], str):
            paras.append([(ln[0] + '  ', size, True, DARK), (ln[1], size, False, DARK)])
        elif isinstance(ln, list):
            paras.append(ln)
        else:
            paras.append([(ln, size, False, DARK)])
    return box(s, x, y, w, h, fill=fill, line=line, paras=paras, margin=0.18)


def formula(s, x, y, w, h, lines, title=None, size=13):
    paras = []
    if title:
        paras.append([(title, 11, True, TEAL)])
    for ln in lines:
        if isinstance(ln, tuple):
            paras.append([(ln[0], size, False, DARK, True), (ln[1], size - 2.5, False, GREY)])
        else:
            paras.append([(ln, size, False, DARK, True)])
    return box(s, x, y, w, h, fill=YEL_L, line=YEL, paras=paras, margin=0.2, anchor=MSO_ANCHOR.MIDDLE)


def tbl(s, x, y, w, colw, rows, size=10.5, rowh=0.38, hl=None, center=(), bold_first=False, hdr_fill=TEAL, wrap_h=None):
    gf = table(s, x, y, w, colw, rows, size=size, rowh=rowh, header_fill=hdr_fill, hl=hl, bold_first_col=bold_first)
    T_ = gf.table
    for r in range(len(rows)):
        for c in center:
            T_.cell(r, c).text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    return gf


def pic_box(s, path, x, y, w, h, border=True):
    iw, ih = Image.open(path).size
    sc = min(w / iw, h / ih)
    pw, ph = iw * sc, ih * sc
    pc = s.shapes.add_picture(path, Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))
    if border:
        pc.line.color.rgb = rgb(LINE); pc.line.width = Pt(1)
    return pc


def bullets(s, x, y, w, h, items, size=12, gap=4, col=DARK, mark='•'):
    paras = []
    for it in items:
        if isinstance(it, tuple):
            paras.append([(mark + '  ', size, True, TEAL), (it[0] + ' ', size, True, col), (it[1], size, False, col)])
        else:
            paras.append([(mark + '  ', size, True, TEAL), (it, size, False, col)])
    return tb(s, x, y, w, h, paras)


def steps(s, x, y, w, items, h=0.9, gap=0.12, size=11.5, numfill=TEAL):
    for k, (a, b) in enumerate(items):
        yy = y + k * (h + gap)
        num(s, x, yy + (h - 0.5) / 2, k + 1, d=0.5, fill=numfill, size=13)
        box(s, x + 0.65, yy, w - 0.65, h, fill=LIGHT, line=LINE, anchor=MSO_ANCHOR.MIDDLE, margin=0.15,
            paras=[[(a, size + 0.5, True, TEAL)], [(b, size, False, DARK)]])


def flow(s, x, y, w, h, items, fills=None, size=11):
    n = len(items); g = 0.3
    bw = (w - g * (n - 1)) / n
    for k, it in enumerate(items):
        xx = x + k * (bw + g)
        f = (fills or [LIGHT] * n)[k]
        a, b = it if isinstance(it, tuple) else (it, '')
        box(s, xx, y, bw, h, fill=f, line=LINE if f != TEAL else None, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, margin=0.08,
            paras=[[(a, size + 1, True, WHITE if f == TEAL else TEAL)]] + ([[(b, size - 1, False, 'D7EEF2' if f == TEAL else DARK)]] if b else []))
        if k < n - 1:
            arrow(s, xx + bw + 0.02, y + h / 2 - 0.13, g - 0.04, 0.26)


def note(s, text, y=9.25, h=0.75, title='Catatan'):
    return box(s, 0.9, y, 18.2, h, fill=YEL_L, line=YEL, anchor=MSO_ANCHOR.MIDDLE, margin=0.2,
               paras=[[(title + '  ', 12, True, TEAL), (text, 12, False, DARK)]])


def line_chart(s, x, y, w, h, cats, series, colors, fmt_='#,##0', legend=True, fs=10, markers=True, dash=None):
    cd = CategoryChartData(); cd.categories = cats
    for nm, vals in series:
        cd.add_series(nm, vals)
    ch = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS if markers else XL_CHART_TYPE.LINE, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    ch.has_title = False; ch.has_legend = legend
    ch.font.size = Pt(fs); ch.font.name = F; ch.font.color.rgb = rgb(DARK)
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(fs)
    for k, sr in enumerate(ch.plots[0].series):
        sr.format.line.color.rgb = rgb(colors[k]); sr.format.line.width = Pt(2.25); sr.smooth = False
        if markers:
            sr.marker.format.fill.solid(); sr.marker.format.fill.fore_color.rgb = rgb(colors[k]); sr.marker.format.line.color.rgb = rgb(colors[k])
        if dash and dash[k]:
            sr.format.line.dash_style = 4
    ch.value_axis.has_major_gridlines = True; ch.value_axis.major_gridlines.format.line.color.rgb = rgb('E3EEF1')
    ch.value_axis.tick_labels.number_format = fmt_; ch.value_axis.tick_labels.number_format_is_linked = False
    ch.value_axis.tick_labels.font.size = Pt(fs - 1); ch.category_axis.tick_labels.font.size = Pt(fs - 1)
    ch.value_axis.format.line.fill.background()
    return ch


def bar_chart(s, x, y, w, h, cats, series, colors, fmt_='0.00', horizontal=False, legend=False, fs=10, labels=True, point_colors=None):
    cd = CategoryChartData(); cd.categories = cats
    for nm, vals in series:
        cd.add_series(nm, vals)
    t = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    ch = s.shapes.add_chart(t, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    ch.has_title = False; ch.has_legend = legend
    ch.font.size = Pt(fs); ch.font.name = F; ch.font.color.rgb = rgb(DARK)
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(fs)
    pl = ch.plots[0]; pl.gap_width = 45
    if len(series) > 1:
        pl.overlap = 0
    for k, sr in enumerate(pl.series):
        sr.format.fill.solid(); sr.format.fill.fore_color.rgb = rgb(colors[k])
    if point_colors:
        for i, c in enumerate(point_colors):
            pt = pl.series[0].points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(c)
    if labels:
        pl.has_data_labels = True
        dl = pl.data_labels; dl.number_format = fmt_; dl.number_format_is_linked = False; dl.show_value = True
        dl.font.size = Pt(fs - 1); dl.position = XL_LABEL_POSITION.OUTSIDE_END
    ch.value_axis.has_major_gridlines = True; ch.value_axis.major_gridlines.format.line.color.rgb = rgb('E3EEF1')
    ch.value_axis.tick_labels.font.size = Pt(fs - 1); ch.category_axis.tick_labels.font.size = Pt(fs - 1)
    ch.value_axis.format.line.fill.background()
    return ch
