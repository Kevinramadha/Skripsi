# Naskah presentasi sidang (15 menit) -> .docx dan .md
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

# (no slide, judul slide, detik, naskah)
S = [
 ('Pembuka', None),
 (1, 'Judul', 20,
  "Assalamu'alaikum warahmatullahi wabarakatuh. Selamat pagi Bapak pembimbing dan Bapak-bapak penguji. "
  "Saya Kevin Atha Fathoni Ramadha, NIM 222212691, kelas 4SI1. Saya akan mempresentasikan skripsi saya yang berjudul "
  "\"Perancangan Aplikasi Simulasi Kebijakan Pariwisata Menggunakan Sistem Dinamis dengan Integrasi Data Citra Satelit, studi kasus Provinsi DIY.\""),
 ('Pendahuluan dan latar belakang', None),
 (2, 'Peran strategis pariwisata (dunia)', 20,
  "Pariwisata adalah sektor yang strategis. Di tingkat dunia, sektor perjalanan dan pariwisata diproyeksikan tumbuh paling cepat dalam lima tahun ke depan, "
  "sekitar 3,6 persen per tahun. Pada 2025 ada 1,53 miliar wisatawan internasional, dan sektor ini menyumbang sekitar 10 persen PDB dunia."),
 (3, 'Peran strategis pariwisata (Indonesia)', 15,
  "Di Indonesia, kontribusi pariwisata terhadap PDB naik dari 2,2 persen pada 2020 menjadi 4,9 persen pada 2025. "
  "Kunjungan wisman mencapai 15,4 juta, dan perjalanan wisnus sekitar 1,2 miliar."),
 (4, 'DIY sebagai lokasi penelitian', 25,
  "Saya memilih DIY karena pada 2025 ada lebih dari 40 juta perjalanan wisnus ke DIY. "
  "Sebanyak 59,69 persen di antaranya berasal dari luar provinsi, tertinggi kedua setelah DKI Jakarta. "
  "DIY juga peringkat pertama nasional pada sub-indeks Travel and Tourism Demand Drivers IPKN 2024, dan pariwisata adalah sektor prioritas kedua Pemda DIY."),
 (5, 'Sistem dinamis sebagai solusi', 35,
  "Namun pertumbuhan ini punya sisi lain. Kunjungan wisatawan mendorong akomodasi, objek wisata, nilai tambah, dan tenaga kerja, tetapi juga menekan ruang dan lahan. "
  "Kepadatan dan berkurangnya daya dukung pada akhirnya melemahkan daya tarik dan menahan kunjungan. Hubungan seperti ini saling memengaruhi, tidak linear, dan ada jeda waktunya. "
  "Metode peramalan seperti ARIMA atau regresi baik untuk memproyeksikan satu variabel, tetapi tidak menangkap umpan balik seperti ini. "
  "Karena itu saya memakai sistem dinamis, yang memodelkan stok, aliran, dan umpan balik, sehingga kebijakan bisa diuji secara virtual sebelum diterapkan. "
  "Pendekatan ini sudah dipakai di Cat Ba Island, Labuan Bajo, dan Baduy."),
 (6, 'Tujuan penelitian', 20,
  "Penelitian ini punya tiga tujuan. Pertama, mengidentifikasi variabel utama penyusun causal loop diagram. "
  "Kedua, mengembangkan model sistem dinamis yang memanfaatkan citra satelit untuk mengisi data yang belum tersedia, lalu menguji strukturnya dan kecocokannya dengan data historis. "
  "Ketiga, membangun aplikasi web sebagai alat bantu pengambilan keputusan."),
 (7, 'Batasan penelitian', 25,
  "Penelitian ini dibatasi pada DIY, dengan data 2015 sampai 2025, tahun dasar 2025, dan simulasi sampai 2050. "
  "Model memuat lima stok utama dan tidak memasukkan guncangan seperti pandemi, juga harga, promosi, dan musim. "
  "Citra satelit hanya dipakai untuk memperbarui data lahan terbangun dan jumlah ODTW. "
  "Hasilnya dipakai untuk membandingkan arah dan besar perubahan antarskenario, bukan untuk meramalkan angka sebenarnya."),
 (8, 'Kontribusi penelitian', 15,
  "Penelitian sistem dinamis pariwisata sebelumnya hanya memakai data statistik. "
  "Penelitian ini menambahkan citra satelit untuk mengisi data yang belum tersedia, dan menyajikan modelnya dalam aplikasi web yang bisa dipakai langsung."),
 ('Metodologi', None),
 (9, 'Kerangka pikir', 20,
  "Kerangka pikirnya: tiga masalah dijawab dengan tiga tujuan. Solusinya menggabungkan data statistik dan citra satelit dengan metode sistem dinamis, dari CLD sampai stock and flow diagram. "
  "Hasilnya dievaluasi lewat uji struktur, uji perilaku, simulasi skenario, dan untuk aplikasinya lewat SUS."),
 (10, 'Data dan sumber data', 15,
  "Data yang dipakai adalah data tahunan tingkat provinsi 2015 sampai 2025, sebagian besar dari BPS DIY, data investasi dari BKPM, "
  "dan dua data citra dari Google Earth Engine: Dynamic World untuk lahan terbangun dan cahaya malam VIIRS untuk ODTW."),
 (11, 'Alat dan perangkat lunak', 12,
  "Citra diolah di Google Earth Engine, model dibangun di Vensim, basis data di Excel, pengujian berulang dengan Python dan PySD, "
  "dan aplikasi dibuat dengan SDEverywhere dan Next.js."),
 (12, 'Tahapan penelitian', 15,
  "Tahapannya mengikuti Mai dan Smith 2018: merumuskan masalah, menyusun hipotesis dinamis dalam CLD, membangun model simulasi, menguji model, "
  "lalu merancang dan mengevaluasi kebijakan."),
 (13, 'Pengumpulan dan preprocessing data', 25,
  "Ada dua masalah data. Pertama, metode pencatatan wisnus berubah pada 2018 ke 2019, sehingga datanya melonjak dari 8,0 menjadi 20,5 juta. "
  "Seri 2015 sampai 2018 saya sambung dengan faktor 2,2997, dan hanya dipakai untuk uji P5, tidak untuk menghitung parameter. "
  "Kedua, Dynamic World baru tersedia sejak 2016, jadi lahan 2015 diisi dengan ekstrapolasi mundur."),
 ('Tujuan 1', None),
 (14, 'Pembatas Tujuan 1', 3, "Masuk ke tujuan pertama."),
 (15, 'Identifikasi variabel CLD', 25,
  "Variabel disusun dalam empat langkah: dimulai dari model Mai dan Smith, disesuaikan dengan kondisi DIY lewat UU 10 Tahun 2009, RIPPARDA, Renja Dinas Pariwisata, Renstra Kemenpar, dan IPKN, "
  "didukung literatur, lalu dievaluasi pada empat tingkat: variabel, hubungan kausal, feedback loop, dan batas model."),
 (16, 'Hasil identifikasi variabel', 15,
  "Hasilnya, semua kandidat variabel dipertahankan. Hanya tenaga kerja yang perannya direvisi: menjadi variabel keluaran dan tidak lagi memengaruhi pembangunan ODTW, karena dasar kausalnya kurang kuat."),
 (17, 'Causal loop diagram', 30,
  "CLD yang terbentuk punya enam loop. R1 menggambarkan pertumbuhan kunjungan yang menumpuk. "
  "R2 adalah jalur ekonomi: wisatawan, pengeluaran, PDRB, investasi, pembangunan ODTW, lalu kembali menaikkan daya tarik. "
  "Lalu ada empat loop penahan: B1 kepadatan, B2 lahan, B3 hotel yang mengikuti permintaan kamar, dan B4 tenaga kerja yang menyesuaikan kebutuhan."),
 ('Tujuan 2', None),
 (18, 'Pembatas Tujuan 2', 3, "Selanjutnya tujuan kedua."),
 (19, 'Citra satelit sebagai estimator', 25,
  "Citra satelit dipakai untuk mengisi dua variabel. Untuk ODTW, data BPS hanya tersedia 2018 sampai 2024, jadi cahaya malam VIIRS diregresikan ke ODTW untuk mengisi 2015 sampai 2017 dan 2025; "
  "stok awal ODTW 2025 menjadi 201 unit. Untuk lahan terbangun, luasnya dihitung langsung dari kelas built Dynamic World tanpa regresi, dengan stok awal 55.029,54 hektar."),
 (20, 'Evaluasi estimasi ODTW', 25,
  "Korelasi data aslinya tinggi, R kuadrat 0,785, tetapi hanya pembanding karena keduanya sama-sama naik. "
  "Yang dipakai adalah korelasi perubahan tahunan, R kuadrat 0,405, di atas batas 0,30. "
  "Dengan LOOCV, prediksi rata-rata meleset 11,29 unit atau 5,99 persen, di bawah batas 10 persen, jadi layak dipakai."),
 (21, 'Validasi temporal ODTW', 12,
  "Secara arah, 4 dari 6 periode searah. Yang tidak searah adalah 2022 dan 2024, sama dengan tahun yang errornya paling besar."),
 (22, 'Validasi spasial ODTW', 12,
  "Secara spasial, dari lima titik, kawasan Malioboro paling terang dan Taman Nasional Gunung Merapi paling redup, sesuai kondisi lapangan."),
 (23, 'Validasi temporal lahan terbangun', 30,
  "Untuk lahan terbangun, data resmi BPS dan DLHK tidak konsisten. Misalnya data BPS naik 40 persen lalu turun 39 persen, hal yang tidak mungkin terjadi pada lahan terbangun dalam setahun. "
  "Dynamic World naik stabil sekitar 12 sampai 13 persen per tahun sampai 2019, jadi itulah yang dipakai. "
  "Setelah 2019, nilainya naik-turun mengikuti jumlah citra yang terekam, bukan perubahan lahan, sehingga penurunan 27 persen pada 2023 sampai 2025 saya catat sebagai keterbatasan."),
 (24, 'Uji akurasi lahan terbangun', 25,
  "Akurasinya diuji dengan 200 titik acak. Metrik utamanya akurasi tertimbang luas, sebesar 89,95 persen, dan Kappa 0,61 yang berarti kesepakatan kuat. "
  "Kesalahan utamanya, 34 titik dipetakan terbangun padahal bukan, sehingga luas lahan terbangun kemungkinan sedikit lebih besar dari kondisi sebenarnya."),
 (25, 'Contoh titik sampel', 8,
  "Ini contohnya: titik 8 berupa vegetasi, dan titik 69 berupa atap bangunan."),
 (26, 'Konversi CLD ke SFD', 15,
  "CLD lalu diterjemahkan menjadi stock and flow diagram dengan 5 stok, 10 aliran, 11 variabel bantu, dan 35 parameter."),
 (27, 'Uji struktur: kesesuaian dan dimensi', 12,
  "Uji struktur pertama: setiap hubungan punya dasar teori atau bukti empiris, dan pemeriksaan satuan di Vensim menyatakan Units are OK dan Model is OK."),
 (28, 'Uji kekekalan materi dan feedback loop', 30,
  "Uji kekekalan materi memastikan stok hanya berubah karena yang masuk dan yang keluar. Contohnya hotel: 2.291 ditambah 158,18 dikurangi 114,55 sama dengan 2.334,63, persis sama dengan hasil simulasi. "
  "Pada uji feedback loop, satu loop dimatikan bergantian. B1 paling kuat: tanpa B1, wisatawan 2050 naik 160 persen. "
  "Setiap balancing loop yang dimatikan membuat pertumbuhan lebih cepat, jadi keenam loop bekerja sesuai CLD."),
 (29, 'Uji kondisi ekstrem dan integration error', 35,
  "Uji kondisi ekstrem terdiri dari 17 uji dengan enam syarat logis. Hasilnya 15 lolos, 2 lolos dengan catatan, dan tidak ada yang gagal. "
  "Misalnya saat investasi dinolkan, TPK berhenti tepat di 1 walaupun permintaan kamar 2,27 kali kapasitas. "
  "Pada E15, lahan sempat melewati luas wilayah sebesar 57,7 hektar, tetapi selisih itu hilang saat time step diperkecil ke 0,25 tahun. "
  "Jadi penyebabnya cara hitung, bukan struktur model. Pada simulasi dasar, time step 1 tahun sudah cukup."),
 (30, 'Rancangan uji perilaku', 25,
  "Uji perilaku dilakukan dua kali. Uji parsial menguji tiap subsistem secara terpisah, dengan masukan dari subsistem lain diganti data aktual. "
  "Uji penuh menjalankan seluruh model 2019 sampai 2025. Penilaiannya berurutan: tren, E1, E2, lalu DC, dengan MAPE dan U1 sampai U3 sebagai pelengkap."),
 (31, 'Hasil uji parsial', 30,
  "Pada uji parsial, semua subsistem lolos uji tren kecuali wisatawan. Hotel paling baik, dengan DC 0,20 dan MAPE 6 persen. "
  "P1b, yang memakai jumlah kamar per unit aktual, polanya sama, jadi hasil P1 tidak bergantung pada asumsi itu. "
  "Wisatawan bermasalah: simulasi tumbuh sekitar 6,4 persen per tahun, sedangkan data 11,7 persen. Selisihnya berpola, jadi ditelusuri lewat kalibrasi."),
 (32, 'Hasil uji penuh', 20,
  "Pada uji penuh, MAPE semua variabel di bawah 50 persen: ODTW sangat baik, lima variabel baik, dan tiga variabel ekonomi layak. "
  "Namun DC hotel dan ODTW lebih buruk dibanding saat diuji sendiri-sendiri."),
 (33, 'Kenapa hasil uji penuh menurun', 30,
  "Penyebabnya satu. Wisatawan tumbuh lebih lambat dari data, sehingga pengeluaran, PDRB, dan investasi ikut rendah, dan hotel serta ODTW yang dibangun lebih sedikit. "
  "Buktinya, pada 6 dari 9 variabel errornya didominasi U1, yaitu selisih rata-rata dari satu sumber. "
  "Sumbernya adalah lonjakan wisatawan setelah pandemi, 12,7 sampai 24,9 persen per tahun pada 2022 sampai 2024, sementara model hanya 6,3 sampai 6,7 persen. Lonjakan ini di luar cakupan model."),
 (34, 'Kalibrasi: aturan', 20,
  "Kalibrasi saya pakai sebagai alat uji: apakah selisih tadi karena nilai parameter atau karena hal di luar model. "
  "Hanya parameter yang belum pasti yang dikalibrasi, dengan rentang yang berdasar, ukuran NRMSE, dan nilai baru harus tetap lolos di model penuh."),
 (35, 'Hasil kalibrasi', 40,
  "Tahap 0 hanya mengoreksi rumus LPE dan LPD, dampaknya 1,6 persen. "
  "Pada Tahap 1 akomodasi, NRMSE turun 11,64 persen, tetapi nilai terbaiknya ada di ujung batas dan hanya karena lonjakan data hotel 2018 sampai 2019; tanpa 2019, perbaikannya tinggal 5,74 persen. "
  "Tahap 2 lahan turun 14,76 persen dan diterima sementara, tetapi di Tahap 4, saat dipakai di model penuh, NRMSE lahan justru naik 32,4 persen, jadi ditolak. "
  "Tahap 3 menunjukkan LPE dan LPD tidak bisa ditentukan terpisah dari data. Hasil akhirnya, tidak ada parameter yang diganti: nilai dari data sudah layak."),
 (36, 'Ringkasan pengujian model', 12,
  "Jadi, struktur model lolos, perilakunya cukup baik, dan kalibrasi tidak mengubah nilai parameter. Model layak dipakai untuk membandingkan skenario kebijakan sampai 2050."),
 (37, 'Analisis sensitivitas', 30,
  "Analisis sensitivitas dilakukan dengan mengubah 32 parameter sebesar 10 persen, dan menguji 17 kelompok parameter asumsi pada rentang penuhnya. "
  "Laju penurunan dasar paling berpengaruh, 42 persen, dan 3 dari 4 parameter teratas adalah asumsi. "
  "Rasio laju penurunan terhadap laju pertumbuhan, 0,75, tidak bisa dipastikan dari data. Bila diubah ke 0,60 sampai 0,90, jumlah wisatawan 2050 bisa bergeser dari minus 42 sampai plus 28 persen. "
  "Karena itu angka absolutnya bukan ramalan, dan perbandingan skenario diuji ulang dengan menggeser asumsi ini."),
 (38, 'Skenario kebijakan', 25,
  "Dua tuas kebijakan lolos tiga kriteria: berpengaruh pada model, punya dasar perda, dan bisa dikendalikan pemerintah. "
  "Yaitu insentif investasi, berdasarkan Perda DIY 1 Tahun 2012, dan konservasi lahan, berdasarkan Perda DIY 6 Tahun 2021. "
  "Tiga skenarionya: Business-as-Usual tanpa keduanya, Sustainable dengan insentif kecil dan konservasi penuh, dan Development Priority dengan insentif besar tanpa pembatasan lahan."),
 (39, 'Perbandingan skenario', 40,
  "Hasilnya pada 2050: Development Priority unggul di ekonomi, lapangan kerja naik 7,18 persen dan PDRB 7,71 persen dibanding BAU. "
  "Sustainable unggul di lingkungan: tidak ada lahan baru yang dibuka untuk pariwisata, dibanding 947 hektar pada BAU, dan ekonominya tetap naik. "
  "BAU punya kepadatan paling rendah, karena wisatawan tumbuh paling lambat. Namun ambang kepadatan 2,0 tetap terlampaui di semua skenario, antara 2041 dan 2043. "
  "Jadi tidak ada skenario yang unggul di semua dimensi, dan urutan ini tetap sama pada tujuh kondisi uji kekokohan."),
 ('Tujuan 3', None),
 (40, 'Pembatas Tujuan 3', 3, "Terakhir, tujuan ketiga."),
 (41, 'Implementasi ke aplikasi web', 30,
  "Hasil pemodelan biasanya berhenti di laporan, dan untuk memakainya orang harus menjalankan software pemodelan. "
  "Karena itu model Vensim yang sudah diuji saya konversi dengan SDEverywhere dan dijalankan di aplikasi web. "
  "Penggunanya adalah Pemda dan Dinas Pariwisata DIY, yaitu pengambil kebijakan yang bukan ahli pemodelan, jadi cukup memilih skenario atau menggeser tuas. Aplikasinya bisa dicoba lewat QR code ini."),
 (42, 'Halaman sistem', 12,
  "Aplikasinya punya empat halaman: Beranda, Model yang menampilkan struktur dan hasil uji, Skenario untuk membandingkan tiga skenario, dan Simulasi untuk mencoba nilai tuas sendiri."),
 (43, 'Hasil evaluasi sistem', 25,
  "Dari uji black-box, sembilan dari sembilan fungsi berjalan sesuai rancangan. "
  "Dari SUS dengan 10 responden, rata-ratanya 88,50, di atas rata-rata 68, dan berada di antara Excellent dan Best Imaginable. "
  "Respondennya masih mahasiswa, jadi uji ke pegawai dinas menjadi saran ke depan."),
 ('Penutup', None),
 (44, 'Kesimpulan dan saran', 35,
  "Kesimpulannya. Pertama, model tersusun dari 5 stok, 10 aliran, 11 variabel bantu, dan 35 parameter dengan enam feedback loop, diadaptasi dari Mai dan Smith sesuai data Indonesia. "
  "Kedua, uji struktur terpenuhi, simulasinya cukup cocok dengan pola historis, dan setiap skenario unggul di dimensi yang berbeda dengan urutan yang tetap konsisten. "
  "Ketiga, aplikasi berjalan pada sembilan dari sembilan fungsi dengan skor SUS 88,50. "
  "Saran saya: perkuat data, tambahkan guncangan seperti pandemi ke dalam model, dan tentukan nilai kebijakan bersama Pemda dan Dinas Pariwisata."),
 (45, 'Terima kasih', 10,
  "Demikian presentasi dari saya. Terima kasih, dan saya siap menerima masukan dari Bapak-bapak. Wassalamu'alaikum warahmatullahi wabarakatuh."),
]


def mmss(s):
    return f'{s // 60}:{s % 60:02d}'


WPM = 145
S = [x if (x[1] is None) else (x[0], x[1], round(len(x[3].split()) * 60 / WPM) + 2, x[3]) for x in S]
rows = [x for x in S if x[1] is not None and isinstance(x[0], int)]
total = sum(r[2] for r in rows)
words = sum(len(r[3].split()) for r in rows)

doc = Document()
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(11.5)
for sec in doc.sections:
    sec.left_margin = sec.right_margin = Cm(2.2); sec.top_margin = sec.bottom_margin = Cm(2)
h = doc.add_heading('Naskah Presentasi Sidang Skripsi', 0)
p = doc.add_paragraph(); p.add_run('Kevin Atha Fathoni Ramadha · 222212691 · durasi target 15 menit').italic = True
p = doc.add_paragraph()
p.add_run(f'Total waktu rencana: {mmss(total)} · ±{words} kata (tempo ±145 kata per menit, ditambah ±2 detik pindah slide). ').bold = True
p.add_run('Kolom waktu di kanan judul slide adalah waktu kumulatif saat slide itu selesai; pakai sebagai patokan saat latihan.')

md = ['# Naskah Presentasi Sidang Skripsi', '', f'Kevin Atha Fathoni Ramadha · 222212691 · target 15 menit · total rencana {mmss(total)} · ±{words} kata', '']
cum = 0
for x in S:
    if x[1] is None:
        doc.add_heading(x[0], 1); md += [f'## {x[0]}', '']
        continue
    n, t, sec, txt = x
    cum += sec
    hp = doc.add_paragraph()
    r = hp.add_run(f'Slide {n} · {t}'); r.bold = True; r.font.color.rgb = RGBColor(0x0B, 0x5E, 0x6E)
    r2 = hp.add_run(f'    ({sec} dtk · s.d. {mmss(cum)})'); r2.font.size = Pt(9.5); r2.font.color.rgb = RGBColor(0x5E, 0x6E, 0x78)
    doc.add_paragraph(txt)
    md += [f'**Slide {n} · {t}** _({sec} dtk · s.d. {mmss(cum)})_', '', txt, '']

doc.add_heading('Tips membawakan', 1)
for t in ['Slide pembatas tujuan (14, 18, 40) cukup diklik sambil mengucapkan satu kalimat transisi.',
          'Slide paling padat (29, 35, 39) adalah titik rawan lewat waktu; bila terlambat lebih dari 30 detik di slide 36, persingkat slide 37 menjadi kalimat pertama dan terakhir saja.',
          'Angka yang sebaiknya hafal: 89,95% (akurasi lahan), 5,99% (LOOCV ODTW), 15/2/0 (uji ekstrem), 6,4% vs 11,7% (pertumbuhan wisatawan), 32,4% (Tahap 4 ditolak), 88,50 (SUS).']:
    doc.add_paragraph(t, style='List Bullet')
md += ['## Tips membawakan', '', '- Slide pembatas tujuan (14, 18, 40) cukup diklik sambil mengucapkan satu kalimat transisi.',
       '- Slide paling padat (29, 35, 39) adalah titik rawan lewat waktu; bila terlambat lebih dari 30 detik di slide 36, persingkat slide 37.',
       '- Angka yang sebaiknya hafal: 89,95%, 5,99%, 15/2/0, 6,4% vs 11,7%, 32,4%, 88,50.']
out = '/home/user/Skripsi/Naskah Presentasi Sidang (15 menit)'
doc.save(out + '.docx')
open(out + '.md', 'w').write('\n'.join(md))
print('total', mmss(total), 'kata', words)
