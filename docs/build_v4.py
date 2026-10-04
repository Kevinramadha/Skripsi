exec(open('/tmp/pptwork/helpers_v4.py').read())
from pptx import Presentation
p = Presentation('/tmp/pptwork/stage2.pptx')
S = list(p.slides)
old = {i: s for i, s in enumerate(S, 1)}

MAIN = [1, 6, 8, 9, 14, 17, 19, 21, 24, 25, 32, 34, 41, 50, 52, 56, 57, 59, 61, 64, 66, 69, 73, 74, 75]
INDEX = 2
DROP = [23, 63, 70]
CAD = [i for i in range(3, 75) if i not in MAIN and i not in DROP]
LAMP = list(range(76, 92))
ORDER = MAIN + [INDEX] + CAD + LAMP
NEW = {o: k for k, o in enumerate(ORDER, 1)}
assert len(set(ORDER)) == len(ORDER) == 91 - len(DROP)

TIME = {6: '0:45', 8: '1:00', 9: '0:45', 14: '0:45', 17: '0:05', 19: '0:45', 21: '1:15', 24: '0:05',
        25: '0:45', 32: '0:45', 34: '1:00', 41: '0:30', 50: '1:00', 52: '0:45', 56: '0:40', 57: '1:00',
        59: '0:30', 61: '0:30', 64: '0:05', 66: '0:45', 69: '0:45', 73: '1:00', 74: '0:20'}
DARK_SL = {17, 24, 64}

def remove_markers(s):
    for sh in list(s.shapes):
        if sh.name.startswith('Penanda prioritas'):
            sh._element.getparent().remove(sh._element)

for i, s in old.items():
    remove_markers(s)
for i, t in TIME.items():
    s = old[i]
    x, y = (18.2, 10.45) if i in DARK_SL else (17.2, 10.55)
    c = box(s, x, y, 0.8, 0.42, fill=RED, line=None, paras=[(t, 12, True, WHITE)],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0, radius=0.5)
    c.name = 'Target waktu ' + t
for i in CAD:
    s = old[i]
    c = box(s, 16.75, 10.55, 1.25, 0.42, fill=LIGHT, line=LINE, paras=[('Cadangan', 11, True, GREY)],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0, radius=0.5)
    c.name = 'Penanda cadangan'

# --- tambahan isi
s = old[6]
box(s, 0.9, 8.55, 10.8, 1.4, fill=YEL_L, line=YEL, paras=[('Mengapa DIY?', 15, True, TEAL),
    ('40 juta perjalanan wisnus 2025 (peringkat ke-7 nasional) · peringkat 1 sub-indeks Demand Drivers IPKN 2024 · sektor prioritas kedua RPJMD DIY setelah pendidikan.', 13, False, DARK)])
s = old[14]
tb(s, 0.9, 8.85, 18.2, 0.4, [('Data: statistik resmi tahunan 2015–2025 (BPS, BKPM) + citra satelit (GEE) · tahun dasar 2025 · horizon 2050 · Δt 1 tahun · DIY 317.036 ha', 13, True, TEAL)])

replace_in_slide(old[17], 'Ringkasan hasil: slide 23', 'Hasil kunci: slide %d' % NEW[21])
replace_in_slide(old[24], 'Ringkasan hasil: slide 63', 'Hasil kunci: slide %d' % NEW[57])
replace_in_slide(old[64], 'Ringkasan hasil: slide 70', 'Hasil kunci: slide %d' % NEW[69])

# --- indeks slide cadangan
s = old[INDEX]
strip(s)
remove_markers(s)
set_title(s, 'Indeks slide cadangan dan lampiran (untuk tanya jawab)')
def rng(lst):
    n = sorted(NEW[i] for i in lst)
    return ', '.join(str(x) for x in n) if len(n) <= 2 or n[-1] - n[0] != len(n) - 1 else '%d–%d' % (n[0], n[-1])
left = [['Topik cadangan', 'Slide'],
        ['Latar belakang & data pariwisata', rng([3, 4, 5])],
        ['Pendekatan sistem dinamis', rng([7])],
        ['Batasan masalah · penelitian terkait', rng([10, 11])],
        ['Kerangka pikir · data · alat', rng([12, 13, 15])],
        ['Preprocessing data', rng([16])],
        ['Metode CLD · hubungan kausal · batas model', rng([18, 20, 22])],
        ['Citra: tahapan, metrik, NTL, Dynamic World', rng([26, 27, 28, 29, 30, 31])],
        ['Konversi SFD · persamaan', rng([33, 35, 36])],
        ['Parameter: sumber, penurunan, asumsi', rng([37, 38, 39])],
        ['Uji struktur: rancangan, loop, ekstrem', rng([40, 42, 43])],
        ['Uji perilaku: rancangan, parsial, penuh, galat', rng([44, 45, 46, 47, 48])],
        ['Kalibrasi: protokol · diagnostik wisatawan', rng([49, 51])],
        ['Sensitivitas · kriteria tuas', rng([53, 54, 55])],
        ['Skenario: daya tarik, dekomposisi, implikasi', rng([58, 60, 62])],
        ['Aplikasi: kebutuhan, halaman, black-box', rng([65, 67, 68])],
        ['Diskusi · keterbatasan', rng([71, 72])]]
L = lambda k: NEW[75 + k]
right = [['Lampiran', 'Slide'],
         ['1 · Definisi pariwisata', str(L(1))], ['2 · Nilai 35 parameter', str(L(2))],
         ['3 · Validasi LOOCV ODTW', str(L(3))], ['4 · Akurasi Dynamic World', str(L(4))],
         ['5 · Data pendukung repository', str(L(5))], ['6 · Sepuluh indikator kinerja', str(L(6))],
         ['7 · KPI pelengkap 2050', str(L(7))], ['8 · Uji perilaku lengkap', str(L(8))],
         ['9 · Aspek dikeluarkan dari model', str(L(9))], ['10 · 17 uji kondisi ekstrem', str(L(10))],
         ['11 · Uji loop umpan balik', str(L(11))], ['12 · Horizon 2150', str(L(12))],
         ['13 · Kalibrasi per tahap', str(L(13))], ['14 · Sensitivitas keluaran lain', str(L(14))],
         ['15 · Persamaan model', str(L(15))], ['16 · Data historis', str(L(16))]]
table(s, 0.9, 2.15, 10.6, [8.7, 1.9], left, size=11.5, rowh=0.47)
table(s, 12.0, 2.15, 7.1, [5.4, 1.7], right, size=11.5, rowh=0.47)
s.notes_slide.notes_text_frame.text = 'Slide indeks, tidak dipresentasikan. Saat tanya jawab, ketik nomor slide lalu Enter untuk melompat langsung (mode Slide Show).'

# --- reorder & drop
sldIdLst = p.slides._sldIdLst
ids = list(sldIdLst)
for i in DROP:
    el = ids[i - 1]
    p.part.drop_rel(el.rId)
    sldIdLst.remove(el)
for o in ORDER:
    el = ids[o - 1]
    sldIdLst.remove(el)
    sldIdLst.append(el)

# --- nav & nomor halaman
SECTIONS = ['Pendahuluan', 'Metodologi', 'Tujuan 1', 'Tujuan 2', 'Tujuan 3', 'Penutup']
def set_nav(slide, sec):
    groups = [sh for sh in slide.shapes if sh.shape_type == 6 and 0.2 * E < sh.top < 0.4 * E and sh.left < 14 * E]
    slots = sorted(set(round(g.left / E, 1) for g in groups))
    if len(slots) != 6:
        return
    for k, x in enumerate(slots):
        gs = [g for g in groups if round(g.left / E, 1) == x]
        active = (k == sec)
        for sub in gs[0].shapes:
            try:
                sub.fill.solid(); sub.fill.fore_color.rgb = rgb(TEAL if active else LIGHT)
            except Exception:
                pass
        for g in gs:
            for sub in g.shapes:
                if sub.has_text_frame and sub.text_frame.text.strip():
                    for para in sub.text_frame.paragraphs:
                        for r in para.runs:
                            r.font.color.rgb = rgb(WHITE if active else GREY)
                            r.font.bold = active
                            r.font.name = FB if active else F
set_nav(old[INDEX], None)
for k, s in enumerate(p.slides, 1):
    for sh in s.shapes:
        if sh.has_text_frame and sh.left > 17.9 * E and sh.top > 10.4 * E and sh.text_frame.text.strip().isdigit():
            set_text(sh, str(k))

exec(open('/tmp/pptwork/notes_v4.py').read())
for i, txt in NOTES16.items():
    old[i].notes_slide.notes_text_frame.text = txt

p.save('/tmp/pptwork/v4.pptx')
print('saved', len(p.slides), 'main', len(MAIN), 'index', NEW[INDEX], 'cad', NEW[CAD[0]], '-', NEW[CAD[-1]], 'lamp', NEW[76], '-', NEW[91])
print({o: NEW[o] for o in MAIN})
